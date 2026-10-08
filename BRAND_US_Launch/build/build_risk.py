#!/usr/bin/env python3
"""Build output/08_Risk_Register.xlsx.
Tabs: Cover, Risk Register, Kill Criteria, Heat Map.
Every model number comes from build/figures.json through build/fmt.py; static facts from build/facts.json.
Score = Probability times Impact (live formula). Kill Criteria carries the Day 90 and Month 12 tables from
figures.json kill_criteria with the Low, Medium and High model values beside each threshold and a live PASS,
REVIEW or KILL formula per metric. Heat Map counts risks by probability and impact with COUNTIFS.
No hyphens, en dashes or em dashes in any text (official form and registration numbers keep their format).
Usage: python3 build_risk.py            builds the workbook
       python3 build_risk.py --patch-fonts   after $RECALC: force every font back to Arial (LibreOffice stores a fallback font for Hangul cells)
       python3 build_risk.py --check         after $RECALC: read values back and assert the gate formulas match figures.json
"""
import json, math, os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import fmt as F
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

ROOT = os.environ.get("ROOT", "/home/user/lakehouse-tracker/BRAND_US_Launch")
BUILD = os.path.join(ROOT, "build"); OUT = os.path.join(ROOT, "output")
PATH = os.path.join(OUT, "08_Risk_Register.xlsx")
facts = json.load(open(os.path.join(BUILD, "facts.json")))
fig = json.load(open(os.path.join(BUILD, "figures.json")))

# ---------------------------------------------------------------- helpers
def patch_fonts(path):
    import zipfile, shutil, re
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                txt = data.decode("utf8")
                names = set(re.findall(r'<name val="([^"]+)"', txt))
                txt = re.sub(r'<name val="(?!Arial")[^"]+"', '<name val="Arial"', txt)
                print("styles.xml font names before patch:", sorted(names))
                data = txt.encode("utf8")
            zout.writestr(item, data)
    shutil.move(tmp, path)

if "--patch-fonts" in sys.argv:
    patch_fonts(PATH); sys.exit(0)

EB, ME = "Evidence based", "Management estimate"
FONT, SZ = "Arial", 10
FILL_IN = PatternFill("solid", fgColor="FFFF00"); BLUE = "0000FF"; GREEN = "008000"
FILL_HDR = PatternFill("solid", fgColor="D9E1F2"); FILL_SEC = PatternFill("solid", fgColor="EDEDED")
FILL_TOT = PatternFill("solid", fgColor="E2EFDA")
FILL_RED = PatternFill("solid", fgColor="FFC7CE"); FILL_AMB = PatternFill("solid", fgColor="FFEB9C"); FILL_GRN = PatternFill("solid", fgColor="C6EFCE")
FMT_MONEY = '$#,##0;($#,##0);0'; FMT_MONEY2 = '$#,##0.00;($#,##0.00);0'; FMT_PCT = '0.0%'; FMT_NUM1 = '0.0'
THIN = Side(style="thin", color="BFBFBF"); BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BUILT = fig["meta"]["built"]
PH = lambda s: "[PLACEHOLDER: " + s + "]"

def W(ws, r, c, v, bold=False, color="000000", fill=None, fmt=None, wrap=False, size=None, italic=False, align=None, border=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = Font(name=FONT, size=size or SZ, bold=bold, italic=italic, color=color)
    if fill: cell.fill = fill
    if fmt: cell.number_format = fmt
    cell.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
    if border: cell.border = BOX
    return cell

# ---------------------------------------------------------------- numbers used in the text (all from figures.json or facts.json)
S = fig["scenarios"]; M, L, H = S["Medium"], S["Low"], S["High"]
GT = fig["gate_thresholds"]; gt = lambda k: GT[k]["value"]
duty = fig["duty"]; ue = fig["unit_economics"]["hero"]; cb = fig["unit_economics"]["cost_bases"]
torn18 = {t["driver"]: t for t in fig["sensitivity"]["tornado_m18"]}
torn5 = {t["driver"]: t for t in fig["sensitivity"]["tornado"]}
lev = {l["n"]: l for l in fig["levers"]}
OPEN_CAP = M["cash"]["opening_capital"]
e, k, am, lg, cp, ut, tm = facts["entity"], facts["korea"], facts["amazon"], facts["logistics"], facts["compliance"], facts["us_tax"], facts["trademark"]
ALI_LO, ALI_HI = facts["august"]["alibaba_benchmark"]["fob_low_usd"], facts["august"]["alibaba_benchmark"]["fob_high_usd"]

money, money_k, pct = F.money, F.money_k, F.pct
m18_net_M = M["m18"]["net"]; m18_rev_M = M["m18"]["revenue"]
m18_net_L = L["m18"]["net"]; m18_rev_L = L["m18"]["revenue"]
y5_net_M = M["annual"]["Y5"]["net"]; y5_rev_M = M["annual"]["Y5"]["revenue"]; y5_rev_L = L["annual"]["Y5"]["revenue"]; y5_net_L = L["annual"]["Y5"]["net"]
gap_M, gap_month_M = M["cash"]["funding_gap"], M["cash"]["funding_needed_by_month"]
peak_M, peak_month_M = M["cash"]["peak_need"], M["cash"]["peak_need_month"]
peak_L, peak_month_L = L["cash"]["peak_need"], L["cash"]["peak_need_month"]
gap_L = L["cash"]["funding_gap"]
runway_M = M["cash"]["runway_months_day90"]; runway_H = H["cash"]["runway_months_day90"]
m12_rev_M = M["gates"]["month12"]["metrics"]["rev"]["value"]; m12_rev_L = L["gates"]["month12"]["metrics"]["rev"]["value"]
cum12_L = L["gates"]["month12"]["metrics"]["cum"]["value"]
tacos3_L = L["gates"]["day90"]["metrics"]["tacos"]["value"]; tacos3_M = M["gates"]["day90"]["metrics"]["tacos"]["value"]
cpu_L = L["gates"]["day90"]["metrics"]["cpu"]["value"]; cpu_M = M["gates"]["day90"]["metrics"]["cpu"]["value"]
months_units_M = M["m18"]["units"]

# ---------------------------------------------------------------- the risks
# [ID, Risk, Category, Probability, Impact, Early warning signal, Mitigation, Owner, Trigger for action, Kill criterion, Basis]
R = []
def add(risk, cat, p, i, ews, mit, owner, trig, kill, basis):
    assert 1 <= p <= 5 and 1 <= i <= 5 and kill in ("Yes", "No") and basis in (EB, ME)
    R.append(["R%02d" % (len(R) + 1), risk, cat, p, i, ews, mit, owner, trig, kill, basis])

C_TAR, C_PROD, C_AMZ, C_CASH, C_OPS, C_LEGAL, C_MKT, C_PEOPLE = "Tariff and customs", "Product and compliance", "Amazon account", "Cash and FX", "Supply chain and operations", "Legal, tax and structure", "Market and demand", "People"

add("Pending second Section 301 action adds 10 points of duty on China origin goods (stress case " + pct(duty["stress"]["rate"]) + " against the copper base " + pct(duty["copper"]["rate"]) + ")",
    C_TAR, 4, 3,
    "USTR publishes a determination or a Federal Register notice on the investigation opened in March 2026; broker alert",
    "Stress case is already modeled: Medium 18 month net falls from " + money_k(m18_net_M) + " to " + money_k(fig["full_recalc"]["duty"]["stress"]["m18_net"]) + ", Year 5 net from " + money_k(y5_net_M) + " to " + money_k(fig["full_recalc"]["duty"]["stress"]["y5_net"]) + " (full model recalc at the stress rate). Hold the price ladder test (lever 10) in reserve, push the fiber ruling (lever 1), buy no more than the 2.5 month cover before a rate change date, and ask the broker about an earlier entry date for goods already on the water",
    "Broker, Mark", "Notice published with an effective date: rerun the model with DutyScenario 3 the same week and reprice within 30 days", "No", EB)
add("CBP binding ruling keeps the copper classification 8544.42.90 (" + pct(duty["copper"]["rate"]) + ") or the fiber heading also sits on List 3 (" + pct(duty["fiber_list3_check"]["rate"]) + " check row), so the 38% to 33.6% landed lever never arrives",
    C_TAR, 3, 3,
    "Ruling letter received with 8544.42; broker reports that 8544.70.00 is covered by List 3",
    "The base case already carries the copper rate, so the downside is a lost upside of " + money_k(lev[1]["model_effect"]) + " over 18 months (lever 1). File the ruling in week 1 with the factory's construction drawings showing the fiber data path; keep the landed cost at 38% in every plan until the letter arrives",
    "Broker, Mark", "Ruling letter received: update DutyScenario on Assumptions and the Levers tab the same day", "No", EB)
add("Factory cost gap: the Alibaba benchmark for a comparable 30 ft AOC is $" + str(ALI_LO) + " to $" + str(ALI_HI) + " FOB against the H2 base of " + money(ue["fob_usd"], 2) + " FOB for the hero; at the benchmark the hero loses " + money(cb["alibaba_low"]["net"], 2) + " to " + money(cb["alibaba_high"]["net"], 2) + " per unit at $44.99 and Medium 18 month net turns to " + money_k(torn18["Factory cost (FOB)"]["low_value"]),
    C_PROD, 3, 5,
    "First USD quote from the factory for the LLC (not the Korea wholesale price) comes in above the H2 base; MOQ or payment terms differ from the Korea deal",
    "Get the factory's written USD FOB quote for the five launch SKUs before the LLC is formed. The single largest unverified assumption in the plan: if the quote is more than 25% above the H2 base, rerun SKU Unit Economics and reprice or stop before the first PO",
    "Lucy", "Written factory quote in USD above the H2 base: rerun the model before any PO; above the Alibaba low benchmark: stop and renegotiate or switch factory", "Yes", EB)
add("HDMI licensing: the factory is not an HDMI Adopter or has no ATC test report for the shipped construction, and HDMI LA seizes goods or forces a listing removal",
    C_PROD, 2, 5,
    "Factory cannot show the adopter certificate or the ATC report number within 10 days of the request; HDMI LA letter; Amazon compliance request",
    "No PO before the HDMI Adopter certificate " + PH("certificate number from factory") + " and the ATC report " + PH("report numbers") + " are on file. Packaging and promotion use the category name (" + cp["category_8k"] + "), never a version number. Titles may say 48Gbps and 8K",
    "Lucy, factory", "Any HDMI LA contact or a missing certificate at the PO date: hold the PO and the listing until the evidence is on file", "Yes", EB)
add("Amazon account suspension or failed identity verification (legal name mismatch across state, IRS, bank and Seller Central; account health drop; related account flag)",
    C_AMZ, 2, 5,
    "Verification request for more documents; account health score drops below Healthy; Seller Performance notice; payouts held",
    "Exact legal name match before any registration; Lucy as the beneficial owner on the account; VA checks account health every day; appeal pack template ready (invoices in the LLC's name, certificates). A suspension longer than " + str(gt("d90_ip_days")) + " days in the gate window is a Day 90 kill trigger",
    "Mark, VA", "Any performance notification: Mark answers within 48 hours; suspension open more than " + str(gt("d90_ip_days")) + " days: gate review", "Yes", EB)
add("Intellectual property complaint against the listing (trademark, HDMI logo use, utility or design patent claim on the connector) that Amazon treats as a rights owner complaint",
    C_AMZ, 2, 4,
    "Rights owner notice in Seller Central; listing goes down for a single ASIN; competitor files a report",
    "Brand Registry on a pending application (" + tm["brand_registry"] + "); logo use confirmed in writing by the factory under its license terms; no competitor brand names in backend keywords; attorney on call for a counter notice. A complaint open more than " + str(gt("d90_ip_days")) + " days is a Day 90 kill trigger",
    "Mark, attorney", "Complaint received: attorney response within 5 business days; open more than " + str(gt("d90_ip_days")) + " days: gate review", "Yes", EB)
add("Working capital runs out before the ramp pays back: Medium peak cash need " + money_k(peak_M) + " in " + F.month_label(peak_month_M) + " against " + money_k(OPEN_CAP) + " opening capital, funding gap " + money_k(gap_M) + " needed by " + F.month_label(gap_month_M) + "; Low peak " + money_k(peak_L) + " in " + F.month_label(peak_month_L),
    C_CASH, 4, 5,
    "Cash balance tracks below the model's monthly line; next PO deposit cannot be paid without delaying it; runway at the Day 90 gate below " + str(gt("d90_run_pass")) + " months",
    "Commit the funding line (retained earnings first, then the line of credit or investor amount) before " + F.month_label(gap_month_M) + ", not when the balance turns negative. Hold inventory cover at 2.5 months, order in smaller lots during the ramp, and run the monthly cash check in the close. Runway below " + str(gt("d90_run_kill")) + " months is a Day 90 kill trigger",
    "Lucy, Mark", "Cash balance more than 10% below the model line for two months, or runway below " + str(gt("d90_run_pass")) + " months: stop the next PO until funding is in place", "Yes", EB)
add("FX: a stronger won raises the USD cost of KRW priced goods; at 1,300 KRW per USD Medium 18 month net is " + money_k(fig["full_recalc"]["fx"]["1300"]["m18_net"]) + " against " + money_k(m18_net_M) + " at 1,430 (1,550 gives " + money_k(fig["full_recalc"]["fx"]["1550"]["m18_net"]) + ")",
    C_CASH, 3, 2,
    "KRW per USD moves below 1,350 for a month; factory invoices in KRW rather than USD",
    "Ask the factory for a USD quote, which removes the effect from the cost side; keep capital in USD once transferred; FX is an input cell (FXRate) so the model reprices in minutes",
    "Lucy", "Rate below 1,350 at a PO date: rerun the model at the spot rate before signing the PO", "No", EB)
add("Return rate above plan: Medium carries 3%, Low 4%; at 5% Medium 18 month net falls to " + money_k(torn18["Return rate"]["low_value"]) + "; above " + pct(gt("d90_ret_kill"), 0) + " at Day 90 is a kill trigger",
    C_PROD, 3, 3,
    "Return reasons show wrong direction, no signal or defective; return rate above " + pct(gt("d90_ret_pass"), 0) + " in any month; negative reviews mention no picture",
    "Direction proof design (color coded SOURCE and DISPLAY ends, printed arrows, QR setup video) is lever 3 (3% to 2%); FAQ in the listing on directionality and in wall installs; 5 Year Signal Guarantee with the replacement shipped before the return arrives keeps the review clean; third party QC before departure",
    "Lucy, factory", "Return rate above " + pct(gt("d90_ret_pass"), 0) + " for two months: pull the return reasons and change the insert card or packaging within 30 days", "Yes", EB)
add("Counterfeits and listing hijackers sell on the ASIN at a lower price, take the buy box and generate bad reviews on the brand",
    C_AMZ, 3, 3,
    "A second seller appears on the ASIN; buy box share drops; reviews mention a different box or no guarantee card",
    "Brand Registry (hijacker protection) on the pending application; Transparency codes or a serialized guarantee card " + PH("program fee") + "; test buy and report through Brand Registry; keep the FNSKU labels applied at the factory",
    "Mark", "Any unauthorized seller on a BRAND ASIN: test buy and report within 7 days", "No", ME)
add("Amazon fee changes: referral, FBA fulfillment, storage, inbound placement or the Professional fee rise faster than the 3% per year FBA inflation in the model",
    C_AMZ, 4, 2,
    "Amazon fee change notice (usually announced for the next calendar year); FBA fee preview shows a higher size tier",
    "Model carries 3% FBA inflation per year from Year 2; packaging to a smaller size tier (lever 8, $5.50 to $4.75 on 7M and 10M); 3PL feeding FBA to avoid aged inventory surcharges (lever 9); reprice through the list and sale price ladder (lever 10)",
    "Mark", "Fee notice received: update the SKU Unit Economics tab within 14 days and reprice if net per unit on the hero falls below " + money(gt("d90_cpu_pass")) + " before advertising", "No", ME)
add("Key person: Lucy (100% owner) is the only signer for the LLC, the bank, the factory and the Korean filings; illness or absence stops payments and filings",
    C_PEOPLE, 2, 5,
    "Payment or filing deadlines slip; Lucy unreachable for more than 5 business days during the launch window",
    "Operating agreement names a successor or a limited power of attorney for payments " + PH("attorney wording") + "; bank user with view and payment approval limits; shared password manager; compliance calendar visible to Mark. Any grant of authority to Mark must keep the treaty position (no contract signing on the LLC's behalf)",
    "Lucy, attorney", "Lucy unavailable for more than 5 business days: Mark runs the compliance calendar and defers POs", "No", ME)
add("Key person: Mark (US side advisor, Amazon operations, broker and bank contact) steps back and the US operating knowledge leaves with him",
    C_PEOPLE, 2, 4,
    "Consulting agreement not renewed; response times stretch; Seller Central tasks pile up with the VA",
    "Written consulting or commission agreement with notice terms; SOPs in 05_Operating_Processes.docx so the VA or a replacement can run the daily checks; Mark's retainer is modeled from Year 3 (Medium) " + PH("fee agreed") + "; keep the broker, bank and 3PL contacts in the LLC's name, not his",
    "Lucy, Mark", "Notice given under the agreement: 60 day handover with the SOPs and a replacement named", "No", ME)
add("Korea filings: the overseas direct investment report (ODI) is not accepted by the foreign exchange bank before capital moves, or the remittance, securities acquisition and annual business performance reports are missed; enforcement against individual online sellers is rising",
    C_LEGAL, 3, 4,
    "Bank asks for documents the LLC does not have yet; capital transfer attempted before acceptance; annual report date passes without the LLC's year end figures",
    "Hard stop: no capital transfer before the ODI report is accepted. The 세무사 files the report and the three follow up reports (" + ", ".join(k["odi_follow_ups"]) + "); the compliance calendar carries each date; the overseas financial account report threshold (" + F.krw(k["overseas_financial_account_report_threshold_krw"]) + " month end aggregate) is checked monthly",
    "Lucy, 세무사", "Any bank query on the report: answer before the transfer; a missed follow up: file late with the 세무사 within 30 days", "No", EB)
add("US tax filings: Form 5472 with the pro forma Form 1120 is missed or late; the penalty is " + money(e["form_5472_penalty_usd"]) + " per year even with no income",
    C_LEGAL, 2, 4,
    "No CPA engaged by Month 1; reportable transaction log not kept; April 15 passes without an extension",
    "CPA engaged in week 4 (" + money(fig["scenarios"]["Medium"]["annual"]["Y2"]["compliance"]) + " per year of compliance cost in the model from Year 2); quarterly reportable transaction log (capital, owner draws, fees between Lucy, H2 Company and BRAND LLC); calendar entry for the filing and the extension",
    "CPA, Lucy", "No engagement letter by Month 1: Mark escalates; April 1 without the return drafted: file the extension", "No", EB)
add("California asserts income tax nexus from FBA inventory stored in California warehouses, adding a state return, fees and possible back tax",
    C_LEGAL, 3, 2,
    "California Franchise Tax Board letter; FBA inventory reports show most stock in California fulfillment centers",
    "On the CPA list from Month 1 (" + ut["california_nexus"] + "); Amazon collects marketplace sales tax, so the exposure is income tax only; CPA quote for a California filing " + PH("CPA quote"),
    "CPA", "Any California letter: CPA answers within the letter's deadline; do not ignore it", "No", EB)
add("Product liability claim (heat, fire or injury from the active heads, or property damage from an in wall install) without the right insurance in place",
    C_PROD, 1, 5,
    "Customer message or review reports heat, smoke or a burn; Amazon safety complaint; monthly sales pass " + money(am["insurance_trigger_monthly_sales_usd"]) + " without a bound policy",
    "Product liability insurance of " + money(am["product_liability_insurance_usd"]) + " is required once sales exceed " + money(am["insurance_trigger_monthly_sales_usd"]) + " in a month; quote in hand before launch and bind the month the threshold is crossed; FCC test report " + PH("report number") + " and factory declarations on file; no in wall claim before CL3 documentation; incident plan: Lucy and the insurer notified the same day",
    "Lucy, Mark", "Any safety report: notify the insurer the same day, pull the lot number, and hold the SKU until the cause is known", "No", EB)
add("Shipping damage or loss on air lots and the first ocean shipments (crushed cartons, lost pallets, FBA receiving discrepancies)",
    C_OPS, 3, 2,
    "FBA receiving shows fewer units than shipped; carrier damage report; QC photos show weak master cartons",
    "Third party QC before departure (" + money(lg["qc_label_per_unit_usd"], 2) + " per unit in the model with labeling); carton drop test in the QC checklist; cargo insurance per shipment " + PH("forwarder quote") + "; reconcile every shipment within 14 days and file reimbursement claims",
    "Lucy, 3PL", "Discrepancy above 2% of a shipment: claim within 14 days and change the carton spec on the next PO", "No", ME)
add("Demand shortfall: conversion below plan so units track the Low case (60% of Medium); Low 18 month revenue " + money_k(m18_rev_L) + " and net " + money_k(m18_net_L) + ", Month 12 revenue " + money(m12_rev_L) + " against the " + money(gt("m12_rev_kill")) + " kill line and " + money(m12_rev_M) + " in Medium",
    C_MKT, 3, 4,
    "Unit session conversion below " + pct(gt("d90_conv_pass"), 0) + " on the hero after the first 30 days of ads; sessions grow but orders do not; units more than 20% below the ramp for two months",
    "Enter narrow and own long length high spec fiber where copper fails past 3 meters at 48Gbps; length variations under one parent ASIN (lever 6); review velocity through Vine and the insert card (lever 5); Spec Match finder and FAQ to lift conversion; keyword pruning (lever 4). The Day 90 gate decides: hero conversion below " + pct(gt("d90_conv_kill"), 0) + " is a kill trigger",
    "Mark", "Units below 80% of the ramp for two months: gate review; Month 12 revenue below " + money(gt("m12_rev_kill")) + ": kill", "Yes", EB)
add("Competition and price compression: $8 copper cables anchor the category and other AOC brands cut price; ASP minus 10% takes Medium 18 month net to " + money_k(torn18["ASP"]["low_value"]) + " and the model already compresses ASP 2% per year",
    C_MKT, 4, 3,
    "Hero sale price forced below $39.99 to hold the buy box; competing AOC listings at the same length under $35; conversion holds only with coupons",
    "Sell the signal, not the cable: do not compete with $8 copper; differentiators (direction proof design, 5 Year Signal Guarantee, CL3 when documented, detachable head); price ladder and coupons (lever 10, list vs sale price); the model carries at least 1 to 2 points of margin compression per year",
    "Mark, Lucy", "ASP below 90% of plan for a quarter: rerun the model and the sensitivity tab; decide between margin and volume at the monthly review", "No", EB)
add("Review manipulation accusation: insert card wording, Vine use, trade up registration or follow up emails read as incentivized reviews and Amazon removes reviews or suspends the account",
    C_AMZ, 2, 4,
    "Review removal notice; policy warning in account health; reviews disappear from the listing",
    "Insert card reviewed against the current Amazon policy " + PH("policy text check") + ": no incentives, no steering by rating; Vine only through the official program; the email list is built only from people who register on the brand site; one policy owner (Mark) signs off every customer facing text",
    "Mark", "Any policy warning: pull the insert card from the next PO and answer Amazon within 48 hours", "No", ME)
add("Listing suppression or compliance block (HDMI naming rule, Prop 65 attribute, FCC or safety document request, pesticide or hazmat misflag) takes the listing down during the ramp; the Low case loses one quarter (Months 19 to 21 at 10% of trend)",
    C_AMZ, 3, 4,
    "Listing status shows suppressed or search suppressed; compliance document request in Seller Central; buy box missing on the hero",
    "Listing text checked against 09_Amazon_Listing_Kit.docx (category name, no version numbers, Prop 65 wording matches the box); FCC report, factory declarations and safety data ready to upload the same day; VA daily listing check. A suppressed listing open more than " + str(gt("d90_ip_days")) + " days is a Day 90 kill trigger",
    "Mark, VA", "Suppression found: documents uploaded within 24 hours; open more than " + str(gt("d90_ip_days")) + " days: gate review", "Yes", EB)
add("Inventory aging: slow lengths (25M, 50M, the 4K line) and over ordering during the ramp leave stock past 180 days, triggering aged inventory surcharges and tying up cash",
    C_OPS, 3, 3,
    "Sell through below 50% of forecast on any SKU for two months; FBA inventory age report shows units past 90 days; cover above 4 months",
    "Launch only the five hero SKUs (100 to 300 units each on air); inventory cover target 2.5 months of forward forecast (Medium); reorder points in the SOPs; 3PL holds bulk and feeds FBA (lever 9); delist losers at the Day 90 gate; ocean freight only on proven SKUs",
    "Lucy, 3PL", "Any SKU above 4 months of cover: stop reordering it, run a coupon, and move stock from FBA to the 3PL before the surcharge date", "No", ME)
add("3PL failure: lost or miscounted inventory, slow FBA replenishment, a closure or a billing dispute that freezes the stock",
    C_OPS, 2, 3,
    "Cycle count variance above 1%; replenishment orders take more than 5 business days; invoices rise without volume",
    "Written 3PL agreement with service levels and insurance " + PH("3PL quote") + "; monthly cycle count reconciled to the model's on hand inventory; keep at least 6 weeks of hero stock inside FBA; second 3PL identified as a fallback",
    "Lucy, 3PL", "Two missed service levels in a month: formal notice and quote the fallback 3PL", "No", ME)
add("Bank account closure: the fintech bank (Mercury or Relay) closes or freezes the foreign owned LLC's account, so Amazon payouts and factory payments stop",
    C_CASH, 3, 4,
    "Bank requests more ownership documents; payout fails in Seller Central; card declined",
    "Open the account with the full document set (certificate, EIN letter, operating agreement, passport); keep a second USD account at the fallback (Payoneer or WorldFirst) ready; no unexplained transfers between personal and LLC accounts; every transfer labeled (capital, owner draw, consulting fee) for Form 5472",
    "Lucy, Mark", "Any bank document request: answer within 3 business days; freeze: switch the Amazon deposit method to the fallback account the same week", "No", ME)
add("Treaty position lost: Mark or another US person negotiates or signs contracts for the LLC, a US office or employee appears, and the IRS treats the LLC as having a permanent establishment, so US income tax applies to the profits",
    C_LEGAL, 2, 4,
    "Any contract signed in the US on the LLC's behalf; Mark's address used as the LLC's place of business; a US hire discussed",
    ut["rule"] + " (" + ut["treaty"] + "); " + ut["fba_no_pe"] + ". Conditions: " + ", ".join(ut["conditions"]) + ". The consulting agreement states that Mark has no authority to negotiate or sign; the LLC uses a registered agent or mailbox address; quarterly confirmation line in the compliance review; CPA memo on file",
    "Lucy, CPA", "Any proposal that gives a US person signing authority or creates a US office: CPA review before it happens", "No", EB)
add("Single factory dependency: the Chinese factory raises MOQ, changes the connector or fiber construction without notice, drifts on quality, or stops supplying the LLC directly",
    C_OPS, 3, 4,
    "Lead time stretches past 45 days; QC defect rate above 2%; factory declines the direction proof design change; a competing distributor gets exclusivity",
    "Written supply agreement with the LLC as buyer, change notice clause and the ATC and FCC reports per construction; third party QC before departure; a second qualified factory sampled in Year 1 " + PH("second factory") + "; factory volume tiers (lever 7, minus 5% at 5,000 units per quarter) negotiated with the agreement",
    "Lucy", "Lead time above 45 days or defect rate above 2% on any PO: hold the next PO until the corrective action report is accepted", "No", ME)
add("Advertising stays inefficient: TACOS does not fall below " + pct(gt("d90_tacos_pass"), 0) + " by Month 3 (Low case " + pct(tacos3_L, 0) + ", Medium " + pct(tacos3_M, 0) + "); 3 points of extra ad spend cost " + money_k(-torn18["Ad spend"]["low_minus_base"]) + " of Medium 18 month net; above " + pct(gt("d90_tacos_kill"), 0) + " in Month 3 is a kill trigger",
    C_MKT, 3, 3,
    "Weekly TACOS not falling for three weeks; ACOS above 40% on the hero campaign; search term report shows spend on irrelevant copper cable terms",
    "Keyword pruning and negative keywords every week (lever 4); phase the ad budget as modeled (30%, 22%, 15%, 11% of revenue); one parent ASIN so ads pool reviews and sessions (lever 6); the ad floor of 8% of revenue is never crossed, so the lever is efficiency, not cuts",
    "Mark, VA", "TACOS above " + pct(gt("d90_tacos_pass"), 0) + " in Month 3: gate review; above " + pct(gt("d90_tacos_kill"), 0) + ": kill", "Yes", EB)
add("Ocean freight transition: the switch from air to ocean after the Day 90 gate (lever 2) adds a 2 month lead time, and a port delay or a customs exam leaves the hero out of stock during the ramp",
    C_OPS, 3, 3,
    "Ocean transit above 45 days; CBP exam notice; FBA stock on the hero below 4 weeks of cover",
    "Switch only proven SKUs to ocean; keep an air top up option with the forwarder; inventory cover 2.5 months with the in transit stock funded by the business and shown separately in the model; ISF and entry documents handed to the broker the same day the goods ship",
    "Lucy, broker, 3PL", "Hero cover below 4 weeks: book an air top up lot within 48 hours", "No", ME)
add("Ownership structure change: Mark takes equity in the LLC, which becomes a partnership with Form 1065 and " + pct(0.37, 0) + " withholding on Lucy's share of profit, and the Korean ODI report has to be amended",
    C_LEGAL, 2, 3,
    "Compensation discussion moves from a fee or commission to a share of the company",
    e["partnership_warning"] + ". Default: " + e["ownership_default"] + ". Alternative to price before any change: " + e["alternative"] + ". Decision 2 in 11_Decision_Log.md",
    "Lucy, Mark, CPA, 세무사", "Any equity proposal: CPA and 세무사 opinions in writing before a term sheet", "No", EB)
add("Product compliance documents missing at launch: FCC test report for the active heads, Prop 65 label, CL3 documentation for the in wall claim, and the Premium High Speed certification for the 4K line are not in hand when the first lot ships",
    C_PROD, 3, 4,
    "Factory has not sent the report numbers 2 weeks before the PO; artwork proof lacks the Prop 65 label; listing draft claims in wall",
    "Hard stops in 06_Compliance_Checklist.xlsx: FCC report " + PH("report number and test lab") + " before the PO; Prop 65 label on every artwork proof; no in wall claim before the CL3 documentation " + PH("document and listing number") + "; the 4K line carries no category claim until the factory shows the certification",
    "Lucy, factory", "Any document missing 2 weeks before the PO date: delay the PO rather than ship without it", "No", EB)
add("H2 Company structure: H2 is a Korean sole proprietorship (registration number " + k["company_registration_number"] + ", middle digits 64), so Lucy invests personally; the 세무사 may advise incorporating H2 first, which changes the investor, the reports and the Korean tax on the LLC's profit",
    C_LEGAL, 3, 2,
    "세무사 advice pending past week 2; bank questions who the investor is",
    "Ask the 세무사 the incorporate H2 first question in week 1 and record the answer in 11_Decision_Log.md; the ODI report is filed by whichever entity the answer names; the model assumes Lucy owns 100% personally",
    "Lucy, 세무사", "Advice to incorporate first: pause the capital transfer, incorporate, then file the ODI report", "No", ME)

assert len(R) >= 25, len(R)
CATS = [C_TAR, C_PROD, C_AMZ, C_CASH, C_OPS, C_LEGAL, C_MKT, C_PEOPLE]
assert set(r[2] for r in R) == set(CATS)

# ---------------------------------------------------------------- workbook
wb = Workbook()

# ---- Cover ----
cv = wb.active; cv.title = "Cover"
cv.column_dimensions["A"].width = 30; cv.column_dimensions["B"].width = 112
r = 1
W(cv, r, 1, "08 Risk Register: BRAND LLC US launch", bold=True, size=14); r += 1
W(cv, r, 1, "Built " + BUILT + ". Generator build/build_risk.py. Model numbers from build/figures.json (pulled from output/03_Financial_Model_5Y.xlsx), facts from build/facts.json.", italic=True, color="595959"); r += 2
rows = [
    ("Purpose", "Every risk that can stop the launch or the 5 year plan, scored for probability and impact, with the early warning signal to watch, the mitigation, the owner, the trigger that forces action, and whether the risk is tied to one of the explicit kill criteria that stop further capital."),
    ("How to use", "Work in the Risk Register tab. Change only the yellow cells (Probability, Impact, Kill criterion if the gate rules change). Score, rank, the Heat Map and the Kill Criteria status cells recalculate. Mark reviews the register in week 4 of every month (05_Operating_Processes.docx section 7) and at the Day 90 and Month 12 gates; Lucy decides on any trigger that fires."),
    ("Readers", "Lucy (decides), Mark (keeps the register and the gate review), 세무사 (Korea filings and structure rows), US CPA (US tax and treaty rows), attorney (IP and ownership rows), customs broker (tariff and classification rows)."),
    ("Probability scale (1 to 5)", "1 = under 10% chance in the next 12 months. 2 = 10% to 25%. 3 = 25% to 50%. 4 = 50% to 75%. 5 = above 75%. Management estimates in every row; revisit at each gate."),
    ("Impact scale (1 to 5)", "1 = under $5,000 or a few days of delay. 2 = $5,000 to $25,000 or a month of delay. 3 = $25,000 to $100,000, about the size of the Medium funding gap (" + money_k(gap_M) + "), or a lost month of sales. 4 = $100,000 to $300,000, a lost quarter, or a filing penalty with account consequences. 5 = ends the Amazon plan, the funding plan or the legal structure. Management estimates; the dollar anchors come from the model."),
    ("Score", "Score = Probability times Impact (live formula, 1 to 25). Red 15 to 25 = act now, named owner reports monthly. Amber 10 to 14 = watch, mitigation in place. Green 1 to 9 = accept and monitor. Rank breaks ties by register order."),
    ("Kill criterion", "Yes means the risk maps to a kill trigger in the Day 90 or Month 12 gate (Kill Criteria tab). " + fig["kill_criteria"]["kill_means"] + "."),
    ("Basis", "Evidence based = the risk fact and its numbers come from the brief, the August model, a verified regulatory fact or the recalculated model. Management estimate = the build team's assumption, to be confirmed by the named owner. The probability and impact scores are management estimates in every row."),
    ("Model context (Medium case)", "18 month revenue " + F.money_m(m18_rev_M) + ", net " + money_k(m18_net_M) + " (" + pct(M["m18"]["net_margin"]) + "), cumulative breakeven " + F.month_label(M["m18"]["breakeven_month"]) + ". Year 5 revenue " + F.money_m(y5_rev_M) + ", net " + money_k(y5_net_M) + ". Peak cash need " + money_k(peak_M) + " in " + F.month_label(peak_month_M) + "; funding gap " + money_k(gap_M) + " needed by " + F.month_label(gap_month_M) + ". Low case: 18 month net " + money_k(m18_net_L) + ", Year 5 revenue " + F.money_m(y5_rev_L) + ", Month 12 gate " + fig["kill_criteria"]["results"]["Low"]["month12"] + "."),
    ("Legend", "Yellow fill with blue text = cell you edit. Black = fixed text or formula. Green = pulled from another sheet. [PLACEHOLDER: ...] = a fact that does not exist yet; never a guess."),
    ("Hard stops that sit above this register", "No capital transfer before the ODI report is accepted. No PO before the HDMI Adopter certificate, ATC report and FCC report are on file. No in wall claim before the CL3 documentation. No HDMI version number on any box or promotion. Funding line committed before " + F.month_label(gap_month_M) + "."),
]
for a, b in rows:
    W(cv, r, 1, a, bold=True); W(cv, r, 2, b, wrap=True)
    cv.row_dimensions[r].height = 14.5 * max(1, math.ceil(len(b) / 118)) + 2; r += 1
r += 1
W(cv, r, 1, "Glossary", bold=True); r += 1
GLOSS = [("AOC", "active optical cable"), ("HDMI", "High Definition Multimedia Interface"), ("HDMI LA", "HDMI Licensing Administrator"), ("ATC", "HDMI Authorized Test Center"), ("ASP", "average selling price"), ("COGS", "cost of goods sold"), ("FOB", "free on board (the price at the factory's port)"), ("FBA", "Fulfillment by Amazon"), ("FNSKU", "Fulfillment Network Stock Keeping Unit (Amazon's barcode)"), ("ASIN", "Amazon Standard Identification Number"), ("ACOS", "advertising cost of sales (ad spend divided by ad attributed sales)"), ("TACOS", "total advertising cost of sales (ad spend divided by total sales)"), ("PPC", "pay per click advertising"), ("3PL", "third party logistics warehouse"), ("HTS", "Harmonized Tariff Schedule"), ("CBP", "US Customs and Border Protection"), ("MFN", "most favored nation duty rate"), ("Section 301", "US trade law tariffs on China origin goods"), ("USTR", "Office of the United States Trade Representative"), ("EIN", "Employer Identification Number"), ("LLC", "limited liability company"), ("IRS", "Internal Revenue Service"), ("CPA", "US certified public accountant"), ("PE", "permanent establishment"), ("ODI", "overseas direct investment report (해외직접투자 신고)"), ("세무사", "Korean licensed tax accountant"), ("FTA", "Foreign Exchange Transactions Act (외국환거래법)"), ("USPTO", "United States Patent and Trademark Office"), ("IP", "intellectual property"), ("FCC", "Federal Communications Commission"), ("Prop 65", "California Proposition 65 warning label"), ("CL3", "in wall cable fire rating (Class 3)"), ("ISF", "Importer Security Filing (the ocean shipment filing to CBP)"), ("MOQ", "minimum order quantity"), ("PO", "purchase order"), ("QC", "quality control"), ("SKU", "stock keeping unit"), ("SOP", "standard operating procedure"), ("VA", "virtual assistant"), ("Gbps", "gigabits per second"), ("FX", "foreign exchange rate (KRW per USD)")]
for a, b in GLOSS:
    W(cv, r, 1, a); W(cv, r, 2, b); r += 1
cv.page_setup.orientation = "landscape"; cv.page_setup.fitToWidth = 1; cv.page_setup.fitToHeight = 0; cv.sheet_properties.pageSetUpPr.fitToPage = True

# ---- Risk Register ----
ws = wb.create_sheet("Risk Register")
HDR = ["ID", "Risk", "Category", "Probability (1 to 5)", "Impact (1 to 5)", "Score", "Early warning signal", "Mitigation", "Owner", "Trigger for action", "Kill criterion (Yes or No)", "Basis", "Rank"]
WIDTHS = [6, 54, 20, 12, 10, 8, 42, 60, 18, 42, 12, 20, 7]
for i, w in enumerate(WIDTHS, 1): ws.column_dimensions[get_column_letter(i)].width = w
HROW = 8; FIRST = HROW + 1; LAST = FIRST + len(R) - 1; RANGE_END = LAST + 20   # 20 spare rows for risks added by hand
PC = "$D$%d:$D$%d" % (FIRST, RANGE_END); IC = "$E$%d:$E$%d" % (FIRST, RANGE_END); SC = "$F$%d:$F$%d" % (FIRST, RANGE_END)
KC = "$K$%d:$K$%d" % (FIRST, RANGE_END); CC = "$C$%d:$C$%d" % (FIRST, RANGE_END); AC = "$A$%d:$A$%d" % (FIRST, RANGE_END); RK = "$M$%d:$M$%d" % (FIRST, RANGE_END)
W(ws, 1, 1, "Risk register: BRAND LLC US launch (Lucy, Hyejin Yoo, 100% owner; H2 Company, Seoul)", bold=True, size=12)
W(ws, 2, 1, "Edit the yellow cells only (Probability, Impact, Kill criterion). Score = Probability times Impact. Red 15 to 25, amber 10 to 14, green 1 to 9. Spare rows below the last risk take new entries; the counts and the Heat Map include them. Built " + BUILT + ".", italic=True, color="595959")
W(ws, 3, 2, "Summary (live counts over the register)", bold=True, fill=FILL_SEC)
labels = ["Risks", "Red (15 to 25)", "Amber (10 to 14)", "Green (1 to 9)", "Kill = Yes", "Average score", "Highest score"]
for j, lab in enumerate(labels):
    W(ws, 4, 2 + j, lab, bold=True, fill=FILL_HDR, wrap=True, border=True)
W(ws, 5, 2, "=COUNTA(%s)" % AC, fmt="0", border=True)
W(ws, 5, 3, '=COUNTIF(%s,">=15")' % SC, fmt="0", border=True, fill=FILL_RED)
W(ws, 5, 4, '=COUNTIFS(%s,">=10",%s,"<15")' % (SC, SC), fmt="0", border=True, fill=FILL_AMB)
W(ws, 5, 5, '=COUNTIFS(%s,">=1",%s,"<10")' % (SC, SC), fmt="0", border=True, fill=FILL_GRN)
W(ws, 5, 6, '=COUNTIF(%s,"Yes")' % KC, fmt="0", border=True)
W(ws, 5, 7, "=IF(B5=0,0,AVERAGE(%s))" % SC, fmt=FMT_NUM1, border=True)
W(ws, 5, 8, "=MAX(%s)" % SC, fmt="0", border=True)
ws.row_dimensions[4].height = 30
W(ws, 6, 2, "Check: band counts add up to the risk count (0 = ok)", italic=True, color="595959"); W(ws, 6, 3, "=B5-SUM(C5:E5)", fmt="0", border=True)
for i, h in enumerate(HDR, 1):
    W(ws, HROW, i, h, bold=True, fill=FILL_HDR, wrap=True, border=True)
ws.row_dimensions[HROW].height = 44
CH = {2: 54, 7: 42, 8: 60, 10: 42, 3: 20, 9: 18, 12: 20}
for n, row in enumerate(R):
    rr = FIRST + n
    rid, risk, cat, p, i, ews, mit, owner, trig, kill, basis = row
    W(ws, rr, 1, rid, border=True)
    W(ws, rr, 2, risk, wrap=True, border=True)
    W(ws, rr, 3, cat, wrap=True, border=True)
    W(ws, rr, 4, p, color=BLUE, fill=FILL_IN, border=True, align="center")
    W(ws, rr, 5, i, color=BLUE, fill=FILL_IN, border=True, align="center")
    W(ws, rr, 6, "=D%d*E%d" % (rr, rr), fmt="0", border=True, align="center", bold=True)
    W(ws, rr, 7, ews, wrap=True, border=True)
    W(ws, rr, 8, mit, wrap=True, border=True)
    W(ws, rr, 9, owner, wrap=True, border=True)
    W(ws, rr, 10, trig, wrap=True, border=True)
    W(ws, rr, 11, kill, color=BLUE, fill=FILL_IN, border=True, align="center")
    W(ws, rr, 12, basis, wrap=True, border=True)
    W(ws, rr, 13, "=IF(F%d=\"\",\"\",RANK(F%d,%s)+COUNTIF($F$%d:F%d,F%d)-1)" % (rr, rr, SC, FIRST, rr, rr), fmt="0", border=True, align="center")
    vals = {2: risk, 7: ews, 8: mit, 10: trig, 3: cat, 9: owner, 12: basis}
    lines = max(math.ceil(len(str(vals[c])) / CH[c]) for c in CH)
    ws.row_dimensions[rr].height = max(15, 13.2 * lines + 4)
for rr in range(LAST + 1, RANGE_END + 1):
    W(ws, rr, 4, None, color=BLUE, fill=FILL_IN); W(ws, rr, 5, None, color=BLUE, fill=FILL_IN); W(ws, rr, 11, None, color=BLUE, fill=FILL_IN)
    W(ws, rr, 6, '=IF(OR(D%d="",E%d=""),"",D%d*E%d)' % (rr, rr, rr, rr), fmt="0", align="center", bold=True)
    W(ws, rr, 13, "=IF(F%d=\"\",\"\",RANK(F%d,%s)+COUNTIF($F$%d:F%d,F%d)-1)" % (rr, rr, SC, FIRST, rr, rr), fmt="0", align="center")
W(ws, RANGE_END + 1, 1, "End of list. Add new risks in the spare rows above this line so the summary and the Heat Map count them.", italic=True, color="595959")
dvn = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True, showErrorMessage=True, errorTitle="Scale", error="Enter a whole number from 1 to 5")
ws.add_data_validation(dvn); dvn.add("D%d:E%d" % (FIRST, RANGE_END))
dvk = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True, showErrorMessage=True, errorTitle="Kill criterion", error="Yes or No")
ws.add_data_validation(dvk); dvk.add("K%d:K%d" % (FIRST, RANGE_END))
dvb = DataValidation(type="list", formula1='"' + EB + "," + ME + '"', allow_blank=True, showErrorMessage=True, errorTitle="Basis", error="Evidence based or Management estimate")
ws.add_data_validation(dvb); dvb.add("L%d:L%d" % (FIRST, RANGE_END))
dvc = DataValidation(type="list", formula1='"' + ",".join(CATS) + '"', allow_blank=True, showErrorMessage=False)
ws.add_data_validation(dvc); dvc.add("C%d:C%d" % (FIRST, RANGE_END))
ws.auto_filter.ref = "A%d:M%d" % (HROW, LAST)
ws.freeze_panes = "C%d" % FIRST
# conditional formatting on Score (red, amber, green) and on the kill flag
rng = "F%d:F%d" % (FIRST, RANGE_END)
ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["15"], fill=FILL_RED, font=Font(name=FONT, size=SZ, bold=True, color="9C0006"), stopIfTrue=True))
ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["10"], fill=FILL_AMB, font=Font(name=FONT, size=SZ, bold=True, color="9C5700"), stopIfTrue=True))
ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=FILL_GRN, font=Font(name=FONT, size=SZ, bold=True, color="006100")))
ws.conditional_formatting.add("K%d:K%d" % (FIRST, RANGE_END), CellIsRule(operator="equal", formula=['"Yes"'], font=Font(name=FONT, size=SZ, bold=True, color="9C0006")))
ws.print_title_rows = "%d:%d" % (HROW, HROW)
ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_area = "A1:M%d" % (LAST + 1)
ws.page_margins.left = ws.page_margins.right = 0.3; ws.page_margins.top = ws.page_margins.bottom = 0.4

# ---- Kill Criteria ----
kc = wb.create_sheet("Kill Criteria")
for col, w in zip("ABCDEFGHIJK", [50, 30, 30, 11, 15, 10, 15, 10, 15, 10, 44]): kc.column_dimensions[col].width = w
W(kc, 1, 1, "Kill criteria: Day 90 gate (end of Month 3) and Month 12 gate, with the model values beside each threshold", bold=True, size=12)
W(kc, 2, 1, "Thresholds and model values come from build/figures.json (gate_thresholds and scenarios.*.gates). " + fig["scenarios"]["Medium"]["gates"]["day90"]["result_rule"] + ". Kill means: " + fig["kill_criteria"]["kill_means"] + ". Observed metrics (conversion, suppressed listing, IP complaint) are read from Seller Central at the gate; the status cells here cover the modeled metrics only.", italic=True, color="595959", wrap=True)
kc.merge_cells("A2:K2"); kc.row_dimensions[2].height = 44
# numeric threshold block
W(kc, 4, 1, "Gate thresholds (numbers; the status formulas read these cells)", bold=True, fill=FILL_SEC)
for j, h in enumerate(["Threshold", "Value", "Unit"]):
    W(kc, 5, 1 + j, h, bold=True, fill=FILL_HDR, border=True)
TH_ORDER = ["d90_conv_pass", "d90_conv_kill", "d90_tacos_pass", "d90_tacos_kill", "d90_ret_pass", "d90_ret_kill", "d90_cpu_pass", "d90_cpu_kill", "d90_run_pass", "d90_run_kill", "d90_ip_days",
            "m12_rev_pass", "m12_rev_kill", "m12_cum_pass", "m12_cum_kill", "m12_tacos_pass", "m12_tacos_kill", "m12_ret_pass", "m12_ret_kill", "m12_gap_kill"]
TH = {}
r = 6
for key in TH_ORDER:
    unit = GT[key]["unit"]; v = GT[key]["value"]
    W(kc, r, 1, GT[key]["label"], border=True, wrap=True)
    W(kc, r, 2, v, fmt=(FMT_PCT if unit == "%" else FMT_MONEY if unit == "$" else "0"), border=True, color=GREEN)
    W(kc, r, 3, {"%": "percent", "$": "USD", "months": "months", "days": "days"}[unit], border=True)
    TH[key] = "$B$%d" % r; r += 1
TH_END = r - 1
r += 1

SCEN = ["Low", "Medium", "High"]
def gate_table(r, title, gate_key, rows):
    """rows: list of (kill_criteria row, metric key or None, status kind)."""
    W(kc, r, 1, title, bold=True, fill=FILL_SEC); r += 1
    hdr = ["Metric", "Pass threshold", "Kill trigger", "Source", "Low model value", "Low status", "Medium model value", "Medium status", "High model value", "High status", "Note"]
    for j, h in enumerate(hdr, 1): W(kc, r, j, h, bold=True, fill=FILL_HDR, wrap=True, border=True)
    kc.row_dimensions[r].height = 30; r += 1
    first = r
    status_cells = {s: [] for s in SCEN}
    for kcrow, mkey, kind, note in rows:
        W(kc, r, 1, kcrow["metric"], wrap=True, border=True)
        W(kc, r, 2, kcrow["pass"], wrap=True, border=True)
        W(kc, r, 3, kcrow["kill"], wrap=True, border=True)
        W(kc, r, 4, kcrow["source"], border=True)
        for si, s in enumerate(SCEN):
            vc, sc = 5 + 2 * si, 6 + 2 * si
            vcol = get_column_letter(vc)
            if mkey is None:
                W(kc, r, vc, PH("observed at the gate"), border=True, wrap=True, italic=True)
                W(kc, r, sc, "observed", border=True, align="center", italic=True)
                continue
            m = S[s]["gates"][gate_key]["metrics"][mkey]
            v = m["value"]
            fmtv = {"pct": FMT_PCT, "money": FMT_MONEY, "money2": FMT_MONEY2, "months": FMT_NUM1}[kind[0]]
            W(kc, r, vc, v, fmt=fmtv, border=True, color=GREEN, align="right")
            pk, kk = kind[1], kind[2]
            if kind[3] == "high_bad":      # pass if value <= pass threshold, kill if value > kill threshold
                f = '=IF(%s%d>%s,"KILL",IF(%s%d<=%s,"PASS","REVIEW"))' % (vcol, r, TH[kk], vcol, r, TH[pk])
            elif kind[3] == "low_bad":     # pass if value >= pass threshold, kill if value < kill threshold
                f = '=IF(%s%d<%s,"KILL",IF(%s%d>=%s,"PASS","REVIEW"))' % (vcol, r, TH[kk], vcol, r, TH[pk])
            else:                          # operating cash flow: pass if positive; kill if negative with a funding gap above the limit
                gap_cell = "%s%d" % (vcol, GAP_ROW)
                f = '=IF(%s%d>0,"PASS",IF(%s>%s,"KILL","REVIEW"))' % (vcol, r, gap_cell, TH[kk])
            W(kc, r, sc, f, border=True, align="center", bold=True)
            status_cells[s].append("%s%d" % (get_column_letter(sc), r))
        W(kc, r, 11, note, wrap=True, border=True, italic=True, color="595959")
        kc.row_dimensions[r].height = 15 * max(2, math.ceil(len(kcrow["metric"]) / 48), math.ceil(len(kcrow["pass"]) / 28), math.ceil(len(kcrow["kill"]) / 28), math.ceil(len(note) / 42), 3 if mkey is None else 1) + 4
        r += 1
    last = r - 1
    # result rows
    W(kc, r, 1, "Gate result from the status cells above (modeled metrics only)", bold=True, border=True, wrap=True)
    for si, s in enumerate(SCEN):
        sc = 6 + 2 * si; L_ = get_column_letter(sc)
        rngs = "%s%d:%s%d" % (L_, first, L_, last)
        W(kc, r, sc, '=IF(COUNTIF(%s,"KILL")>0,"KILL",IF(COUNTIF(%s,"REVIEW")>0,"REVIEW","PASS"))' % (rngs, rngs), bold=True, border=True, align="center")
    res_row = r; r += 1
    W(kc, r, 1, "Gate result in the model (figures.json)", border=True)
    for si, s in enumerate(SCEN):
        W(kc, r, 6 + 2 * si, fig["kill_criteria"]["results"][s][gate_key], border=True, align="center", color=GREEN)
    W(kc, r + 1, 1, "Check: formula result equals the model result (0 = ok)", italic=True, color="595959")
    for si, s in enumerate(SCEN):
        L_ = get_column_letter(6 + 2 * si)
        W(kc, r + 1, 6 + 2 * si, "=IF(%s%d=%s%d,0,1)" % (L_, res_row, L_, r), fmt="0", border=True, align="center")
    for rr in (res_row, r):
        for si in range(3):
            c = 6 + 2 * si; L_ = get_column_letter(c)
            kc.conditional_formatting.add("%s%d" % (L_, rr), CellIsRule(operator="equal", formula=['"KILL"'], fill=FILL_RED, font=Font(name=FONT, size=SZ, bold=True, color="9C0006")))
            kc.conditional_formatting.add("%s%d" % (L_, rr), CellIsRule(operator="equal", formula=['"REVIEW"'], fill=FILL_AMB, font=Font(name=FONT, size=SZ, bold=True, color="9C5700")))
            kc.conditional_formatting.add("%s%d" % (L_, rr), CellIsRule(operator="equal", formula=['"PASS"'], fill=FILL_GRN, font=Font(name=FONT, size=SZ, bold=True, color="006100")))
    for si in range(3):
        L_ = get_column_letter(6 + 2 * si)
        rngs = "%s%d:%s%d" % (L_, first, L_, last)
        kc.conditional_formatting.add(rngs, CellIsRule(operator="equal", formula=['"KILL"'], fill=FILL_RED, font=Font(name=FONT, size=SZ, bold=True, color="9C0006")))
        kc.conditional_formatting.add(rngs, CellIsRule(operator="equal", formula=['"REVIEW"'], fill=FILL_AMB, font=Font(name=FONT, size=SZ, bold=True, color="9C5700")))
        kc.conditional_formatting.add(rngs, CellIsRule(operator="equal", formula=['"PASS"'], fill=FILL_GRN, font=Font(name=FONT, size=SZ, bold=True, color="006100")))
    return r + 3, res_row

# scenario context block (cash figures the operating cash flow formula reads)
W(kc, r, 1, "Scenario context from the model", bold=True, fill=FILL_SEC); r += 1
for j, h in enumerate(["Item", "", "", "", "Low", "", "Medium", "", "High", "", "Note"], 1):
    if h: W(kc, r, j, h, bold=True, fill=FILL_HDR, border=True)
r += 1
ctx = [("Funding gap (1.2 times the negative minimum cash balance)", lambda s: S[s]["cash"]["funding_gap"], FMT_MONEY, "Read by the Month 12 operating cash flow status formula"),
       ("Funding needed by", lambda s: S[s]["cash"]["funding_needed_by_month"], '"Month "0', "The month before the balance first goes negative"),
       ("Peak cash need", lambda s: S[s]["cash"]["peak_need"], FMT_MONEY, ""),
       ("Peak cash need month", lambda s: S[s]["cash"]["peak_need_month"], '"Month "0', ""),
       ("Cash runway at the Day 90 gate after funding, months", lambda s: S[s]["cash"]["runway_months_day90"], FMT_NUM1, "Cash at the end of Month 3 after the funding line, divided by the average operating burn of Months 4 to 6"),
       ("18 month net (Amazon US)", lambda s: S[s]["m18"]["net"], FMT_MONEY, ""),
       ("Cumulative breakeven month", lambda s: S[s]["m18"]["breakeven_month"], '"Month "0', "none = no cumulative breakeven in 18 months"),
       ("Month 12 revenue", lambda s: S[s]["gates"]["month12"]["metrics"]["rev"]["value"], FMT_MONEY, ""),
       ("Year 5 revenue", lambda s: S[s]["annual"]["Y5"]["revenue"], FMT_MONEY, ""),
       ("Year 5 net", lambda s: S[s]["annual"]["Y5"]["net"], FMT_MONEY, "")]
GAP_ROW = r
for label, fn, fm, note in ctx:
    W(kc, r, 1, label, border=True, wrap=True)
    for si, s in enumerate(SCEN):
        v = fn(s)
        if v is None: W(kc, r, 5 + 2 * si, "none", border=True, color=GREEN, align="right")
        else: W(kc, r, 5 + 2 * si, v, fmt=fm, border=True, color=GREEN, align="right")
    W(kc, r, 11, note, italic=True, color="595959", wrap=True, border=True)
    r += 1
r += 1

d90 = fig["kill_criteria"]["day90"]; m12 = fig["kill_criteria"]["month12"]
D90_ROWS = [
    (d90[0], None, None, "Observed in Seller Central (unit session percentage per child ASIN)."),
    (d90[1], "tacos", ("pct", "d90_tacos_pass", "d90_tacos_kill", "high_bad"), "Model value is the Month 3 ad spend rate. The falling month over month test is observed."),
    (d90[2], "ret", ("pct", "d90_ret_pass", "d90_ret_kill", "high_bad"), "Model value is the scenario return rate in Month 3."),
    (d90[3], "cpu", ("money2", "d90_cpu_pass", "d90_cpu_kill", "low_bad"), "Model value is the hero contribution per unit before advertising at the scenario ASP; the 3 of 5 SKU test is observed."),
    (d90[4], "runway", ("months", "d90_run_pass", "d90_run_kill", "low_bad"), "Model value assumes the funding line is in place by the month shown in the context block."),
    (d90[5], None, None, "Observed: any suppression or IP complaint open more than " + str(gt("d90_ip_days")) + " days is a kill."),
]
M12_ROWS = [
    (m12[0], "rev", ("money", "m12_rev_pass", "m12_rev_kill", "low_bad"), "All channels in Month 12 (Amazon US only in Year 1)."),
    (m12[1], "cum", ("money", "m12_cum_pass", "m12_cum_kill", "low_bad"), "Cumulative net after fixed costs, seller fee and launch costs."),
    (m12[2], "tacos", ("pct", "m12_tacos_pass", "m12_tacos_kill", "high_bad"), "Month 12 ad spend rate."),
    (m12[3], "ret", ("pct", "m12_ret_pass", "m12_ret_kill", "high_bad"), "Scenario return rate in Month 12."),
    (m12[4], "opcf", ("money", None, "m12_gap_kill", "opcf"), "Value is the lower of Months 11 and 12. The status formula reads the funding gap in the context block for the kill test; the full three month test sits in 03_Financial_Model_5Y.xlsx, Funding and Breakeven."),
]
r, D90_RES = gate_table(r, "Day 90 gate (end of Month 3). Pass requires every row; any kill trigger stops further capital", "day90", D90_ROWS)
r, M12_RES = gate_table(r, "Month 12 gate. Pass requires every row; any kill trigger stops further capital", "month12", M12_ROWS)
W(kc, r, 1, "Register rows tied to a kill trigger (Kill criterion = Yes)", bold=True, border=True)
W(kc, r, 5, "='Risk Register'!F5", fmt="0", color=GREEN, border=True, align="center"); r += 1
W(kc, r, 1, "Plan B if a gate kills the Amazon plan: B2B installer distribution plus continued Korea wholesale (03_Financial_Model_5Y.xlsx, Plan B tab)", border=True, wrap=True)
pb = fig["plan_b"]
W(kc, r, 5, pb["Y1"]["revenue"], fmt=FMT_MONEY, color=GREEN, border=True); W(kc, r, 6, "Y1 revenue", border=True)
W(kc, r, 7, pb["Y1"]["net"], fmt=FMT_MONEY, color=GREEN, border=True); W(kc, r, 8, "Y1 net", border=True)
W(kc, r, 9, pb["Y3"]["net"], fmt=FMT_MONEY, color=GREEN, border=True); W(kc, r, 10, "Y3 net", border=True)
kc.row_dimensions[r].height = 48
kc.freeze_panes = "A4"
kc.page_setup.orientation = "landscape"; kc.page_setup.fitToWidth = 1; kc.page_setup.fitToHeight = 0; kc.sheet_properties.pageSetUpPr.fitToPage = True

# ---- Heat Map ----
hm = wb.create_sheet("Heat Map")
for col, w in zip("ABCDEFGHIJ", [26, 12, 12, 12, 12, 12, 12, 4, 30, 14]): hm.column_dimensions[col].width = w
W(hm, 1, 1, "Heat map: count of risks by probability and impact (COUNTIFS over the Risk Register tab, spare rows included)", bold=True, size=12)
W(hm, 2, 1, "Rows are probability (5 at the top), columns are impact. Cell color is the score band of that cell (red 15 to 25, amber 10 to 14, green 1 to 9). Bold = at least one risk in the cell.", italic=True, color="595959")
RR = "'Risk Register'!"
W(hm, 4, 1, "Probability \\ Impact", bold=True, fill=FILL_HDR, border=True, align="center")
for i in range(1, 6): W(hm, 4, 1 + i, "Impact " + str(i), bold=True, fill=FILL_HDR, border=True, align="center")
W(hm, 4, 7, "Total", bold=True, fill=FILL_HDR, border=True, align="center")
for pi, p in enumerate([5, 4, 3, 2, 1]):
    rr = 5 + pi
    W(hm, rr, 1, "Probability " + str(p), bold=True, fill=FILL_HDR, border=True)
    for i in range(1, 6):
        sc = p * i
        fill = FILL_RED if sc >= 15 else FILL_AMB if sc >= 10 else FILL_GRN
        W(hm, rr, 1 + i, "=COUNTIFS(%s%s,%d,%s%s,%d)" % (RR, PC, p, RR, IC, i), fmt="0", fill=fill, border=True, align="center", color=GREEN)
    W(hm, rr, 7, "=SUM(B%d:F%d)" % (rr, rr), fmt="0", border=True, align="center", fill=FILL_TOT)
    hm.row_dimensions[rr].height = 24
W(hm, 10, 1, "Total", bold=True, fill=FILL_TOT, border=True)
for c in range(2, 8):
    L_ = get_column_letter(c); W(hm, 10, c, "=SUM(%s5:%s9)" % (L_, L_), fmt="0", border=True, align="center", fill=FILL_TOT, bold=True)
W(hm, 11, 1, "Check (0 = ok)", italic=True, color="595959"); W(hm, 11, 2, "=G10-%sB5" % RR, fmt="0", border=True)
hm.conditional_formatting.add("B5:F9", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, size=12, bold=True, color="000000")))
# kill criterion grid
W(hm, 13, 1, "Risks tied to a kill trigger (Kill criterion = Yes), same grid", bold=True, fill=FILL_SEC)
W(hm, 14, 1, "Probability \\ Impact", bold=True, fill=FILL_HDR, border=True, align="center")
for i in range(1, 6): W(hm, 14, 1 + i, "Impact " + str(i), bold=True, fill=FILL_HDR, border=True, align="center")
W(hm, 14, 7, "Total", bold=True, fill=FILL_HDR, border=True, align="center")
for pi, p in enumerate([5, 4, 3, 2, 1]):
    rr = 15 + pi
    W(hm, rr, 1, "Probability " + str(p), bold=True, fill=FILL_HDR, border=True)
    for i in range(1, 6):
        sc = p * i
        fill = FILL_RED if sc >= 15 else FILL_AMB if sc >= 10 else FILL_GRN
        W(hm, rr, 1 + i, '=COUNTIFS(%s%s,%d,%s%s,%d,%s%s,"Yes")' % (RR, PC, p, RR, IC, i, RR, KC), fmt="0", fill=fill, border=True, align="center", color=GREEN)
    W(hm, rr, 7, "=SUM(B%d:F%d)" % (rr, rr), fmt="0", border=True, align="center", fill=FILL_TOT)
W(hm, 20, 1, "Total", bold=True, fill=FILL_TOT, border=True)
for c in range(2, 8):
    L_ = get_column_letter(c); W(hm, 20, c, "=SUM(%s15:%s19)" % (L_, L_), fmt="0", border=True, align="center", fill=FILL_TOT, bold=True)
hm.conditional_formatting.add("B15:F19", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, size=12, bold=True, color="000000")))
# by band and by category
W(hm, 22, 1, "By score band", bold=True, fill=FILL_SEC); W(hm, 22, 2, "Risks", bold=True, fill=FILL_HDR, border=True); W(hm, 22, 3, "Share", bold=True, fill=FILL_HDR, border=True); W(hm, 22, 4, "Kill = Yes", bold=True, fill=FILL_HDR, border=True)
bands = [("Red (15 to 25)", '">=15"', None, FILL_RED), ("Amber (10 to 14)", '">=10"', '"<15"', FILL_AMB), ("Green (1 to 9)", '">=1"', '"<10"', FILL_GRN)]
for bi, (lab, lo, hi, fill) in enumerate(bands):
    rr = 23 + bi
    W(hm, rr, 1, lab, border=True, fill=fill)
    if hi: W(hm, rr, 2, "=COUNTIFS(%s%s,%s,%s%s,%s)" % (RR, SC, lo, RR, SC, hi), fmt="0", border=True, color=GREEN)
    else: W(hm, rr, 2, "=COUNTIF(%s%s,%s)" % (RR, SC, lo), fmt="0", border=True, color=GREEN)
    W(hm, rr, 3, "=IF(%sB5=0,0,B%d/%sB5)" % (RR, rr, RR), fmt=FMT_PCT, border=True, color=GREEN)
    if hi: W(hm, rr, 4, '=COUNTIFS(%s%s,%s,%s%s,%s,%s%s,"Yes")' % (RR, SC, lo, RR, SC, hi, RR, KC), fmt="0", border=True, color=GREEN)
    else: W(hm, rr, 4, '=COUNTIFS(%s%s,%s,%s%s,"Yes")' % (RR, SC, lo, RR, KC), fmt="0", border=True, color=GREEN)
W(hm, 27, 1, "By category", bold=True, fill=FILL_SEC)
for j, h in enumerate(["Risks", "Average score", "Highest score", "Red", "Kill = Yes"], 2): W(hm, 27, j, h, bold=True, fill=FILL_HDR, border=True, wrap=True)
hm.row_dimensions[27].height = 30
for ci, cat in enumerate(CATS):
    rr = 28 + ci
    W(hm, rr, 1, cat, border=True)
    W(hm, rr, 2, "=COUNTIF(%s%s,A%d)" % (RR, CC, rr), fmt="0", border=True, color=GREEN)
    W(hm, rr, 3, "=IF(B%d=0,0,SUMIFS(%s%s,%s%s,A%d)/B%d)" % (rr, RR, SC, RR, CC, rr, rr), fmt=FMT_NUM1, border=True, color=GREEN)
    W(hm, rr, 4, "=IF(B%d=0,0,SUMPRODUCT(MAX((%s%s=A%d)*%s%s*%s%s)))" % (rr, RR, CC, rr, RR, PC, RR, IC), fmt="0", border=True, color=GREEN)
    W(hm, rr, 5, '=COUNTIFS(%s%s,A%d,%s%s,">=15")' % (RR, CC, rr, RR, SC), fmt="0", border=True, color=GREEN)
    W(hm, rr, 6, '=COUNTIFS(%s%s,A%d,%s%s,"Yes")' % (RR, CC, rr, RR, KC), fmt="0", border=True, color=GREEN)
TR = 28 + len(CATS)
W(hm, TR, 1, "Total", bold=True, fill=FILL_TOT, border=True); W(hm, TR, 2, "=SUM(B28:B%d)" % (TR - 1), fmt="0", bold=True, fill=FILL_TOT, border=True)
W(hm, TR, 5, "=SUM(E28:E%d)" % (TR - 1), fmt="0", bold=True, fill=FILL_TOT, border=True); W(hm, TR, 6, "=SUM(F28:F%d)" % (TR - 1), fmt="0", bold=True, fill=FILL_TOT, border=True)
# top 8 by rank (deck slide 33 reads these)
TOPR = TR + 3
W(hm, TOPR - 1, 1, "Top 8 risks by score (rank from the register; ties keep register order)", bold=True, fill=FILL_SEC)
W(hm, TOPR, 1, "Position", bold=True, fill=FILL_HDR, border=True)
for j, h in enumerate(["Rank", "ID", "Score", "Probability", "Impact", "Kill = Yes"], 2): W(hm, TOPR, j, h, bold=True, fill=FILL_HDR, border=True)
W(hm, TOPR, 9, "Risk", bold=True, fill=FILL_HDR, border=True); W(hm, TOPR, 10, "Owner", bold=True, fill=FILL_HDR, border=True)
hm.column_dimensions["I"].width = 90; hm.column_dimensions["J"].width = 20
for kk in range(1, 9):
    rr = TOPR + kk
    W(hm, rr, 1, "Rank " + str(kk), border=True)
    W(hm, rr, 2, kk, fmt="0", border=True, align="center")
    mt = "MATCH(B%d,%s%s,0)" % (rr, RR, RK)
    W(hm, rr, 3, "=IFERROR(INDEX(%s%s,%s),\"\")" % (RR, AC, mt), border=True, color=GREEN, align="center")
    W(hm, rr, 4, "=IFERROR(INDEX(%s%s,%s),\"\")" % (RR, SC, mt), fmt="0", border=True, color=GREEN, align="center")
    W(hm, rr, 5, "=IFERROR(INDEX(%s%s,%s),\"\")" % (RR, PC, mt), fmt="0", border=True, color=GREEN, align="center")
    W(hm, rr, 6, "=IFERROR(INDEX(%s%s,%s),\"\")" % (RR, IC, mt), fmt="0", border=True, color=GREEN, align="center")
    W(hm, rr, 7, "=IFERROR(INDEX(%s%s,%s),\"\")" % (RR, KC, mt), border=True, color=GREEN, align="center")
    W(hm, rr, 9, "=IFERROR(INDEX(%s$B$%d:$B$%d,%s),\"\")" % (RR, FIRST, RANGE_END, mt), border=True, color=GREEN, wrap=True)
    W(hm, rr, 10, "=IFERROR(INDEX(%s$I$%d:$I$%d,%s),\"\")" % (RR, FIRST, RANGE_END, mt), border=True, color=GREEN, wrap=True)
    hm.row_dimensions[rr].height = 44
hm.conditional_formatting.add("D%d:D%d" % (TOPR + 1, TOPR + 8), CellIsRule(operator="greaterThanOrEqual", formula=["15"], fill=FILL_RED, font=Font(name=FONT, size=SZ, bold=True, color="9C0006"), stopIfTrue=True))
hm.conditional_formatting.add("D%d:D%d" % (TOPR + 1, TOPR + 8), CellIsRule(operator="greaterThanOrEqual", formula=["10"], fill=FILL_AMB, font=Font(name=FONT, size=SZ, bold=True, color="9C5700"), stopIfTrue=True))
hm.conditional_formatting.add("D%d:D%d" % (TOPR + 1, TOPR + 8), CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=FILL_GRN, font=Font(name=FONT, size=SZ, bold=True, color="006100")))
hm.freeze_panes = "A4"
hm.page_setup.orientation = "landscape"; hm.page_setup.fitToWidth = 1; hm.page_setup.fitToHeight = 0; hm.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- check mode (after $RECALC)
if "--check" in sys.argv:
    wbv = load_workbook(PATH, data_only=True)
    rs = wbv["Risk Register"]; kv = wbv["Kill Criteria"]; hv = wbv["Heat Map"]
    out = {"risks": rs["B5"].value, "red": rs["C5"].value, "amber": rs["D5"].value, "green": rs["E5"].value, "kill_yes": rs["F5"].value, "avg": rs["G5"].value, "max": rs["H5"].value, "check": rs["C6"].value}
    out["sample_rows"] = [[rs.cell(row=rr, column=c).value for c in (1, 4, 5, 6, 13)] for rr in range(FIRST, FIRST + 5)]
    out["heat_total"] = hv["G10"].value; out["heat_check"] = hv["B11"].value
    out["heat_grid"] = [[hv.cell(row=rr, column=c).value for c in range(2, 7)] for rr in range(5, 10)]
    out["top8"] = [[hv.cell(row=rr, column=c).value for c in (2, 3, 4)] for rr in range(TOPR + 1, TOPR + 9)]
    out["day90_formula"] = [kv.cell(row=D90_RES, column=c).value for c in (6, 8, 10)]
    out["day90_model"] = [kv.cell(row=D90_RES + 1, column=c).value for c in (6, 8, 10)]
    out["month12_formula"] = [kv.cell(row=M12_RES, column=c).value for c in (6, 8, 10)]
    out["month12_model"] = [kv.cell(row=M12_RES + 1, column=c).value for c in (6, 8, 10)]
    out["gate_checks"] = [kv.cell(row=rr, column=c).value for rr in (D90_RES + 2, M12_RES + 2) for c in (6, 8, 10)]
    errs = []
    for ws_ in wbv.worksheets:
        for row in ws_.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"): errs.append((ws_.title, cell.coordinate, cell.value))
    out["error_cells"] = errs
    print(json.dumps(out, ensure_ascii=False, default=str))
    assert not errs, errs
    assert out["check"] == 0 and out["heat_check"] == 0, out
    assert out["day90_formula"] == out["day90_model"] and out["month12_formula"] == out["month12_model"], out
    assert all(v == 0 for v in out["gate_checks"]), out
    assert out["risks"] == len(R) and out["heat_total"] == len(R), out
    sys.exit(0)

os.makedirs(OUT, exist_ok=True)
wb.save(PATH)
print(json.dumps({"file": PATH, "risks": len(R), "categories": len(CATS), "header_row": HROW, "first_row": FIRST, "last_row": LAST,
                  "kill_yes": sum(1 for x in R if x[9] == "Yes"), "evidence_based": sum(1 for x in R if x[10] == EB), "management_estimate": sum(1 for x in R if x[10] == ME),
                  "scores": sorted([x[3] * x[4] for x in R], reverse=True)}))
