#!/usr/bin/env python3
"""Final tie out (BUILD_SPEC section 8).

1. Reads the delivered workbook (data_only, default selector state Scenario 2 Medium, Duty 1 Copper, FX 1430) and
   build/figures.json. Through build/figures_keys.json (JSON key path -> defined name -> sheet!cell -> selector state)
   every mapped figure is re derived from the workbook and compared within 0.5%. Scenario blocks (Low, Medium, High)
   are re derived the way build/pull_figures.py does: a temporary copy per scenario in build/qa/tie_out/, selector set,
   LibreOffice recalc, values read back. The full_recalc variants (fiber, stress, FX 1300, FX 1550) are re derived the same way.
2. Extracts the text of both decks with markitdown and asserts that the headline figures, formatted by build/fmt.py,
   appear in both (the Korean deck uses the Korean month and year suffix forms, N개월차 and N년차, for Month N and Year N).
3. Checks the SOP docx, the listing kit docx, the two one pagers (pdftotext), the three xlsx deliverables and, when present,
   the README and the decision log for the capital need, the funding gap, the gate thresholds and the other figures they quote.
4. Runs build/lint_hyphens.py over output/.
5. Prints a results table (deliverable, checks, passed, failed, details), saves it to build/tie_out_result.md and exits 1 on any failure.

Usage: source build/env.sh && python3 build/tie_out.py [--skip-recalc]
"""
import json, os, re, shutil, subprocess, sys, datetime
from pathlib import Path

ROOT = Path(os.environ.get("ROOT", "/home/user/lakehouse-tracker/BRAND_US_Launch"))
BUILD = ROOT / "build"; OUT = ROOT / "output"; QA = BUILD / "qa" / "tie_out"
QA.mkdir(parents=True, exist_ok=True)
sys.path.insert(0, str(BUILD))
import fmt  # noqa: E402
from verify_model import recalc, Reader  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

WB = OUT / "03_Financial_Model_5Y.xlsx"
F = json.load(open(BUILD / "figures.json"))
MAP = json.load(open(BUILD / "figures_map.json"))
KEYS = json.load(open(BUILD / "figures_keys.json"))
FACTS = json.load(open(BUILD / "facts.json"))
SKIP_RECALC = "--skip-recalc" in sys.argv
TOL = 0.005

RESULTS = []   # (deliverable, label, ok, detail)

def check(deliverable, label, ok, detail=""):
    RESULTS.append((deliverable, label, bool(ok), detail))
    return bool(ok)

# ---------------------------------------------------------------- helpers
def norm(v):
    if v is None or v == "" or v == "none": return None
    return v

def close(a, b, tol=TOL):
    a = norm(a); b = norm(b)
    if a is None and b is None: return True
    if isinstance(a, bool) or isinstance(b, bool): return a == b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        if abs(b) < 1e-9: return abs(a) <= 0.01
        return abs(a - b) <= max(tol * abs(b), 0.01)
    return a == b

def get_path(obj, path):
    if "[]." in path:
        pre, post = path.split("[].", 1)
        return [get_path(x, post) for x in get_path(obj, pre)]
    cur = obj
    for part in path.split("."):
        cur = cur[part]
    return cur

def flatten(v):
    if isinstance(v, dict):
        out = []
        for x in v.values(): out += flatten(x)
        return out
    if isinstance(v, list):
        out = []
        for x in v: out += flatten(x)
        return out
    return [v]

def sfmt(x):
    if isinstance(x, float): return f"{x:,.4f}"
    return str(x)

def text_of_pptx(p):
    return subprocess.run(["markitdown", str(p)], capture_output=True, text=True).stdout

def text_of_docx(p):
    return subprocess.run(["pandoc", "-t", "plain", "--columns=400", str(p)], capture_output=True, text=True).stdout

def text_of_pdf(p):
    """Layout text (for reading) followed by the raw reading order text with every whitespace run squashed to one space,
    so a figure that wraps inside a table cell (KRW\n55.8M) is still found as one string."""
    lay = subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True, text=True).stdout
    raw = subprocess.run(["pdftotext", str(p), "-"], capture_output=True, text=True).stdout
    return lay + "\n==== raw ====\n" + re.sub(r"\s+", " ", raw)

def squash(t):
    return re.sub(r"[ \t]+", " ", t)

def ko_forms(s):
    """Korean deck forms of an English label: Month N -> N개월차, Year N -> N년차; keep the English form as an alternative."""
    alts = [s]
    k = re.sub(r"Month (\d+)", r"\1개월차", s)
    k = re.sub(r"Year (\d+)", r"\1년차", k)
    if k != s: alts.append(k)
    if s == "none": alts.append("없음")
    return alts

# ---------------------------------------------------------------- 1. workbook vs figures.json
DEL = "03_Financial_Model_5Y.xlsx"
wb_vals = load_workbook(WB, data_only=True)
ws_a = wb_vals["Assumptions"]
sel_state = (ws_a["C3"].value, ws_a["C4"].value, ws_a["C5"].value)
check(DEL, "delivered workbook at the default selector state (Scenario 2, Duty 1, FX 1430)", sel_state == (2, 1, 1430), f"found {sel_state}")
check(DEL, "figures.json meta.selector_state_delivered agrees", tuple(F["meta"]["selector_state_delivered"].values()) == (2, 1, 1430), str(F["meta"]["selector_state_delivered"]))

vd = Reader(str(WB))
none_cells = []
for name, ref in MAP.items():
    v = vd(name)
    if isinstance(v, list):
        if all(x is None for x in v): none_cells.append(name)
    elif v is None: none_cells.append(name)
check(DEL, "workbook is recalculated: every mapped cell holds a value (data_only, no None)", not none_cells, "None in: " + ", ".join(none_cells[:12]))

readers = {"any": vd}
if not SKIP_RECALC:
    for sel, sname in ((1, "Low"), (2, "Medium"), (3, "High")):
        tmp = QA / f"model_scenario_{sel}.xlsx"
        shutil.copy(WB, tmp)
        w = load_workbook(tmp); w["Assumptions"]["C3"] = sel; w.save(tmp)
        res = recalc(str(tmp))
        ok = res.get("status") == "success" and res.get("total_errors") == 0
        check(DEL, f"per scenario recalc of a temporary copy at Scenario {sel} ({sname}) succeeds with 0 errors", ok, json.dumps(res)[:200])
        r = Reader(str(tmp))
        check(DEL, f"temporary copy reads Scenario {sel} ({sname}) after recalc", r("Scenario") == sel and r("ScenarioName") == sname, f"{r('Scenario')} {r('ScenarioName')}")
        readers[sel] = r
else:
    readers[2] = vd

# every mapped key path
mism = []; n_keys = 0; n_vals = 0; skipped = []
for path, d in KEYS.items():
    scen = d["scenario"]
    reader_key = "any" if (scen == "any" or str(scen).startswith("any")) else scen
    if reader_key not in readers:
        skipped.append(path); continue
    r = readers[reader_key]
    jv = get_path(F, path)
    cv = r(d["name"])
    n_keys += 1
    if "[]." in path:                       # monthly rows: 24 values; ranges with 25 cells start at Month 0
        cells = cv[1:] if isinstance(cv, list) and len(cv) == 25 else cv
        jvals = jv
    elif isinstance(jv, (list, dict)):      # table blocks: flatten in column order
        jvals = flatten(jv); cells = cv if isinstance(cv, list) else [cv]
    else:
        jvals = [jv]; cells = cv if isinstance(cv, list) else [cv]
    if len(jvals) != len(cells):
        mism.append((path, f"length {len(jvals)} vs {len(cells)}")); continue
    for i, (a, b) in enumerate(zip(jvals, cells)):
        n_vals += 1
        if not close(a, b):
            mism.append((path + (f"[{i}]" if len(jvals) > 1 else ""), f"json {sfmt(a)} vs workbook {sfmt(b)}"))
# Medium block also re derived from the delivered workbook at its default state
if 2 in readers and readers[2] is not vd:
    for path, d in KEYS.items():
        if d["scenario"] != 2: continue
        jv = get_path(F, path); cv = vd(d["name"])
        if "[]." in path:
            cells = cv[1:] if isinstance(cv, list) and len(cv) == 25 else cv; jvals = jv
        else:
            jvals = flatten(jv) if isinstance(jv, (list, dict)) else [jv]; cells = cv if isinstance(cv, list) else [cv]
        for i, (a, b) in enumerate(zip(jvals, cells)):
            n_vals += 1
            if not close(a, b): mism.append((path + " (delivered workbook)", f"json {sfmt(a)} vs workbook {sfmt(b)}"))
check(DEL, f"every figures.json key path re derived from the workbook ({n_keys} key paths, {n_vals} values, tolerance 0.5%)", not mism and not skipped,
      ("; ".join(f"{p}: {m}" for p, m in mism[:10]) + (f" (+{len(mism)-10} more)" if len(mism) > 10 else "")) + (f" skipped {len(skipped)}" if skipped else ""))

# scenario summary engines vs the selector pulls, and the headline block vs the scenario blocks
ss_bad = []
for s in ("Low", "Medium", "High"):
    S = F["scenarios"][s]; SS = F["scenario_summary"][s]; HL = F["headline"][s]
    pairs = [("m18_revenue", S["m18"]["revenue"]), ("m18_net", S["m18"]["net"]), ("m18_net_margin", S["m18"]["net_margin"]), ("m18_breakeven_month", S["m18"]["breakeven_month"]),
             ("m18_run_rate", S["m18"]["run_rate"]), ("m18_exit_margin", S["m18"]["exit_margin"]), ("cash_peak_need", S["cash"]["peak_need"]), ("cash_peak_need_month", S["cash"]["peak_need_month"]),
             ("cash_funding_gap", S["cash"]["funding_gap"]), ("cash_funding_needed_by_month", S["cash"]["funding_needed_by_month"]), ("cash_runway_months_day90", S["cash"]["runway_months_day90"]),
             ("gate_day90", S["gates"]["day90"]["result"]), ("gate_month12", S["gates"]["month12"]["result"])]
    pairs += [(f"annual_Y{y}_revenue", S["annual"][f"Y{y}"]["revenue"]) for y in range(1, 6)] + [(f"annual_Y{y}_net", S["annual"][f"Y{y}"]["net"]) for y in range(1, 6)]
    for k, v in pairs:
        if not close(SS[k], v): ss_bad.append(f"{s}.{k}: summary {sfmt(SS[k])} vs scenario {sfmt(v)}")
    hl_pairs = [("m18_revenue", S["m18"]["revenue"]), ("m18_net", S["m18"]["net"]), ("breakeven_month", S["m18"]["breakeven_month"]), ("peak_need", S["cash"]["peak_need"]), ("funding_gap", S["cash"]["funding_gap"]),
                ("y5_net_margin", S["annual"]["Y5"]["net_margin"]), ("gate_day90", S["gates"]["day90"]["result"]), ("gate_month12", S["gates"]["month12"]["result"])]
    for k, v in hl_pairs:
        if not close(HL[k], v): ss_bad.append(f"headline.{s}.{k}")
    for y in range(1, 6):
        if not close(HL["annual_revenue"][y - 1], S["annual"][f"Y{y}"]["revenue"]) or not close(HL["annual_net"][y - 1], S["annual"][f"Y{y}"]["net"]): ss_bad.append(f"headline.{s}.Y{y}")
    for k in ("day90", "month12"):
        if F["kill_criteria"]["results"][s][k] != S["gates"][k]["result"]: ss_bad.append(f"kill_criteria.results.{s}.{k}")
check(DEL, "Scenario Summary engines, headline block and kill_criteria.results agree with the scenario blocks", not ss_bad, "; ".join(ss_bad[:8]))

# full recalc variants (duty and FX effects quoted by the deck, SOPs and risk register)
if not SKIP_RECALC:
    variants = [("fiber", ("duty", "fiber"), 2, 2, 1430), ("stress", ("duty", "stress"), 2, 3, 1430), ("fx1300", ("fx", "1300"), 2, 1, 1300), ("fx1550", ("fx", "1550"), 2, 1, 1550), ("medium_base", ("duty", "copper"), 2, 1, 1430)]
    vbad = []
    for tag, (blk, key), sel, duty, fx in variants:
        tmp = QA / f"model_variant_{tag}.xlsx"
        shutil.copy(WB, tmp)
        w = load_workbook(tmp); w["Assumptions"]["C3"] = sel; w["Assumptions"]["C4"] = duty; w["Assumptions"]["C5"] = fx; w.save(tmp)
        res = recalc(str(tmp))
        if not (res.get("status") == "success" and res.get("total_errors") == 0): vbad.append(f"{tag}: recalc {json.dumps(res)[:120]}"); continue
        r = Reader(str(tmp))
        J = F["full_recalc"][blk][key]
        pairs = [("landed_pct", r("LandedPctActive")), ("duty_rate", r("DutyRateActive")), ("m18_revenue", r("m18_revenue")), ("m18_net", r("m18_net")), ("m18_net_margin", r("m18_net_margin")),
                 ("m18_breakeven_month", r("m18_breakeven_month")), ("y5_revenue", r("annual_Y5_revenue")), ("y5_net", r("annual_Y5_net")), ("y5_net_margin", r("annual_Y5_net_margin")),
                 ("cum_net_5y", r("annual_Y5_cum_net")), ("peak_need", r("cash_peak_need")), ("peak_need_month", r("cash_peak_need_month")), ("funding_gap", r("cash_funding_gap")),
                 ("funding_needed_by_month", r("cash_funding_needed_by_month")), ("hero_net_per_unit", r("duty_" + {1: "copper", 2: "fiber", 3: "stress"}[duty] + "_hero_net"))]
        for k, cv in pairs:
            if not close(J[k], cv): vbad.append(f"{tag}.{k}: json {sfmt(J[k])} vs workbook {sfmt(cv)}")
    check(DEL, "full_recalc variants (Medium at fiber, stress, FX 1300, FX 1550, base) re derived by recalc of temporary copies", not vbad, "; ".join(vbad[:8]))

# ---------------------------------------------------------------- 2. decks
S = F["scenarios"]; L, M, H = S["Low"], S["Medium"], S["High"]; UE = F["unit_economics"]["hero"]; D = F["duty"]; FR = F["full_recalc"]; GT = F["gate_thresholds"]; PB = F["plan_b"]
money, money_k, money_m, pct, num, units, month_label = fmt.money, fmt.money_k, fmt.money_m, fmt.pct, fmt.num, fmt.units, fmt.month_label
def pct_auto(x): return pct(x, 0) if abs(x * 100 - round(x * 100)) < 1e-9 else pct(x, 1)
def gt(k):
    t = GT[k]; return pct_auto(t["value"]) if t["unit"] == "%" else money(t["value"]) if t["unit"] == "$" else num(t["value"])
def be(s): return "none" if s["m18"]["breakeven_month"] is None else month_label(s["m18"]["breakeven_month"])

HEADLINE = [
    ("Medium 18 month revenue", money_m(M["m18"]["revenue"])), ("Medium 18 month net", money_k(M["m18"]["net"])), ("Medium 18 month margin", pct(M["m18"]["net_margin"])),
    ("Low 18 month revenue", money_k(L["m18"]["revenue"])), ("Low 18 month net", money_k(L["m18"]["net"])),
    ("High 18 month revenue", money_m(H["m18"]["revenue"])), ("High 18 month net", money_k(H["m18"]["net"])), ("High 18 month margin", pct(H["m18"]["net_margin"])),
    ("Medium breakeven month", be(M)), ("High breakeven month", be(H)), ("Low breakeven (none)", be(L)),
    ("Medium Month 18 run rate", money_m(M["m18"]["run_rate"])), ("Medium exit margin", pct(M["m18"]["exit_margin"])),
    ("Medium Year 5 revenue", money_m(M["annual"]["Y5"]["revenue"])), ("Medium Year 5 net", money_k(M["annual"]["Y5"]["net"])), ("Medium Year 5 margin", pct(M["annual"]["Y5"]["net_margin"])),
    ("High Year 5 revenue", money_m(H["annual"]["Y5"]["revenue"])), ("High Year 5 net", money_m(H["annual"]["Y5"]["net"])),
    ("Low Year 5 revenue", money_k(L["annual"]["Y5"]["revenue"])), ("Low Year 5 net", money_k(L["annual"]["Y5"]["net"])),
    ("Medium peak cash need", money_k(M["cash"]["peak_need"])), ("Medium peak need month", month_label(M["cash"]["peak_need_month"])),
    ("Low peak cash need", money_k(L["cash"]["peak_need"])), ("Low peak need month", month_label(L["cash"]["peak_need_month"])),
    ("High peak cash need", money_k(H["cash"]["peak_need"])), ("High peak need month", month_label(H["cash"]["peak_need_month"])),
    ("Medium funding gap", money_k(M["cash"]["funding_gap"])), ("Medium funding needed by", month_label(M["cash"]["funding_needed_by_month"])),
    ("Low funding gap", money_k(L["cash"]["funding_gap"])), ("High funding gap", money_k(H["cash"]["funding_gap"])), ("High funding needed by", month_label(H["cash"]["funding_needed_by_month"])),
    ("Opening capital", money_k(M["cash"]["opening_capital"])),
    ("Medium Day 90 runway", num(M["cash"]["runway_months_day90"], 1)),
    ("Hero price", money(UE["price"], 2)), ("Hero list price", money(UE["list"], 2)), ("Hero landed cost", money(UE["landed_cost"], 2)), ("Hero landed %", pct_auto(UE["landed_pct"])),
    ("Hero referral", money(UE["referral"], 2)), ("Hero FBA", money(UE["fba"], 2)), ("Hero advertising", money(UE["ads"], 2)), ("Hero returns and storage", money(UE["returns_storage"], 2)),
    ("Hero net", money(UE["net"], 2)), ("Hero net %", pct(UE["net_pct"])), ("Hero Korea wholesale gross", money(UE["korea_wholesale_gross"], 2)), ("Hero contribution before ads at ASP 42", money(UE["contribution_before_ads_at_asp42"], 2)),
    ("Cost base H2 landed", money(F["unit_economics"]["cost_bases"]["h2"]["landed"], 2)), ("Cost base H2 net", money(F["unit_economics"]["cost_bases"]["h2"]["net"], 2)),
    ("Cost base Alibaba low net", money(F["unit_economics"]["cost_bases"]["alibaba_low"]["net"], 2)), ("Cost base Alibaba high net", money(F["unit_economics"]["cost_bases"]["alibaba_high"]["net"], 2)),
    ("Duty copper", pct(D["copper"]["rate"])), ("Duty fiber", pct(D["fiber"]["rate"])), ("Duty stress", pct(D["stress"]["rate"])), ("Landed % Medium fiber", pct(D["fiber"]["landed_pct_medium"])),
    ("Hero net fiber", money(D["fiber"]["hero_net"], 2)), ("Hero net stress", money(D["stress"]["hero_net"], 2)),
    ("Fiber 18 month net", money_k(FR["duty"]["fiber"]["m18_net"])), ("Stress 18 month net", money_k(FR["duty"]["stress"]["m18_net"])),
    ("Fiber funding gap", money_k(FR["duty"]["fiber"]["funding_gap"])), ("Stress funding gap", money_k(FR["duty"]["stress"]["funding_gap"])),
    ("Gate Day 90 TACOS pass", gt("d90_tacos_pass")), ("Gate Day 90 TACOS kill", gt("d90_tacos_kill")), ("Gate Day 90 return pass", gt("d90_ret_pass")), ("Gate Day 90 return kill", gt("d90_ret_kill")),
    ("Gate Day 90 contribution pass", gt("d90_cpu_pass")), ("Gate Day 90 contribution kill", gt("d90_cpu_kill")), ("Gate Month 12 revenue pass", gt("m12_rev_pass")), ("Gate Month 12 revenue kill", gt("m12_rev_kill")),
    ("Gate Month 12 cumulative kill", gt("m12_cum_kill")), ("Gate Month 12 TACOS pass", gt("m12_tacos_pass")), ("Gate Month 12 TACOS kill", gt("m12_tacos_kill")), ("Gate Month 12 funding gap kill", gt("m12_gap_kill")),
    ("Medium Month 12 revenue", money(M["gates"]["month12"]["metrics"]["rev"]["value"])), ("Medium Month 12 cumulative net", money(M["gates"]["month12"]["metrics"]["cum"]["value"])),
    ("Gate results Low", f"{L['gates']['day90']['result']}"), ("Gate results Low Month 12", f"{L['gates']['month12']['result']}"),
    ("Plan B Year 1 revenue", money_k(PB["Y1"]["revenue"])), ("Plan B Year 1 net", money_k(PB["Y1"]["net"])), ("Plan B Year 3 revenue", money_k(PB["Y3"]["revenue"])), ("Plan B Year 3 net", money_k(PB["Y3"]["net"])),
]
for lang, fname in (("en", "01_Business_Proposal_EN.pptx"), ("ko", "02_Business_Proposal_KO.pptx")):
    txt = text_of_pptx(OUT / fname)
    (QA / f"deck_{lang}.md").write_text(txt)
    check(fname, "markitdown extracted the deck text (37 slides)", txt.count("<!-- Slide number:") == 37, f"{txt.count('<!-- Slide number:')} slides")
    for label, s in HEADLINE:
        alts = ko_forms(s) if lang == "ko" else [s]
        check(fname, f"{label}: {s}", any(a in txt for a in alts), "not found: " + " | ".join(alts))
    # no stale figures from the brief or the earlier spec sentence
    for stale in ("$457,804", "$457K", "24.7%", "$455K"):
        check(fname, f"stale High figure absent: {stale}", stale not in txt, "found")

# ---------------------------------------------------------------- 3. documents
def doc_checks(fname, txt, items):
    for label, alts in items:
        if isinstance(alts, str): alts = [alts]
        check(fname, f"{label}: {alts[0]}", any(a in txt for a in alts), "not found")

opening = M["cash"]["opening_capital"]; FX = F["meta"]["fx"]
def inv_days_label(): return None

# 05 SOPs
sops = squash(text_of_docx(OUT / "05_Operating_Processes.docx")); (QA / "sops.txt").write_text(sops)
mm = M["monthly"]
sop_items = [
    ("opening capital", [money(opening), money_k(opening)]), ("Medium peak cash need", money_k(M["cash"]["peak_need"])), ("Medium peak need month", month_label(M["cash"]["peak_need_month"])),
    ("Medium funding gap", money_k(M["cash"]["funding_gap"])), ("Medium funding needed by", month_label(M["cash"]["funding_needed_by_month"])),
    ("Low peak cash need", money_k(L["cash"]["peak_need"])), ("Low funding gap", money_k(L["cash"]["funding_gap"])), ("High peak cash need", money_k(H["cash"]["peak_need"])), ("High funding gap", money_k(H["cash"]["funding_gap"])),
    ("Medium 18 month revenue", money_m(M["m18"]["revenue"])), ("Medium 18 month net", money_k(M["m18"]["net"])), ("Medium 18 month margin", pct(M["m18"]["net_margin"])), ("Medium breakeven", be(M)),
    ("hero price", money(UE["price"], 2)), ("hero list", money(UE["list"], 2)), ("hero landed", money(UE["landed_cost"], 2)), ("hero net", money(UE["net"], 2)), ("hero net %", pct(UE["net_pct"])), ("hero contribution at ASP 42", money(UE["contribution_before_ads_at_asp42"], 2)),
    ("hero duty copper", money(next(p for p in F["unit_economics"]["per_sku"] if p["sku"] == "8K_10M")["duty_copper"], 2)), ("hero duty fiber", money(next(p for p in F["unit_economics"]["per_sku"] if p["sku"] == "8K_10M")["duty_fiber"], 2)),
    ("hero air freight", money(next(p for p in F["unit_economics"]["per_sku"] if p["sku"] == "8K_10M")["air_freight"], 2)), ("hero ocean freight", money(next(p for p in F["unit_economics"]["per_sku"] if p["sku"] == "8K_10M")["ocean_freight"], 2)),
    ("landed % copper", pct(D["copper"]["landed_pct_medium"])), ("landed % fiber", pct(D["fiber"]["landed_pct_medium"])), ("landed % stress", pct(D["stress"]["landed_pct_medium"])),
    ("hero net fiber", money(D["fiber"]["hero_net"], 2)), ("fiber 18 month net", money_k(FR["duty"]["fiber"]["m18_net"])), ("stress 18 month net", money_k(FR["duty"]["stress"]["m18_net"])),
    ("duty copper", pct(D["copper"]["rate"])), ("duty fiber", pct(D["fiber"]["rate"])), ("duty stress", pct(D["stress"]["rate"])),
    ("gate TACOS pass", gt("d90_tacos_pass")), ("gate TACOS kill", gt("d90_tacos_kill")), ("gate return pass", gt("d90_ret_pass")), ("gate return kill", gt("d90_ret_kill")),
    ("gate contribution pass", gt("d90_cpu_pass")), ("gate contribution kill", gt("d90_cpu_kill")), ("gate Month 12 revenue pass", gt("m12_rev_pass")), ("gate Month 12 revenue kill", gt("m12_rev_kill")),
    ("gate Month 12 cumulative kill", gt("m12_cum_kill")), ("gate Month 12 TACOS pass", gt("m12_tacos_pass")), ("gate Month 12 TACOS kill", gt("m12_tacos_kill")), ("gate Month 12 return pass", gt("m12_ret_pass")), ("gate Month 12 return kill", gt("m12_ret_kill")), ("gate funding gap kill", gt("m12_gap_kill")),
    ("Medium Month 12 revenue", money(M["gates"]["month12"]["metrics"]["rev"]["value"])), ("Low Month 12 revenue", money(L["gates"]["month12"]["metrics"]["rev"]["value"])), ("Medium Month 12 cumulative net", money(M["gates"]["month12"]["metrics"]["cum"]["value"])),
    ("Plan B Year 1 revenue", money_k(PB["Y1"]["revenue"])), ("Plan B Year 1 net", money_k(PB["Y1"]["net"])), ("Plan B Year 3 revenue", money_k(PB["Y3"]["revenue"])), ("Plan B Year 3 net", money_k(PB["Y3"]["net"])),
    ("Medium Month 12 inventory on hand", money_k(mm[11]["inventory_on_hand"])), ("Medium Month 12 inventory in transit", money_k(mm[11]["inventory_in_transit"])),
    ("insurance trigger month Medium", "Medium " + month_label(M["insurance_trigger_month"])), ("insurance trigger margin", "threshold by " + money(M["insurance_trigger_margin_usd"])),
    ("lever 3 returns from, to", "from " + pct(F["levers"][2]["from"]) + " to " + pct(F["levers"][2]["to"])), ("returns Month 1 and Month 19 (Medium)", [pct_auto(mm[0]["returns_pct"]) + " Year 1 in the Medium model, " + pct_auto(mm[18]["returns_pct"])]),
    ("Form 5472 penalty", money(FACTS["entity"]["form_5472_penalty_usd"])), ("continuous bond", money(FACTS["logistics"]["broker_bond_annual_usd"])),
]
doc_checks("05_Operating_Processes.docx", sops, sop_items)
bad_rows = []
for m in range(1, 13):
    if not re.search(rf"Month {m}\s+{re.escape(units(mm[m-1]['units']))}\b", sops): bad_rows.append(f"Month {m} {units(mm[m-1]['units'])}")
check("05_Operating_Processes.docx", "reorder table: Medium units Months 1 to 12 match figures.json", not bad_rows, "; ".join(bad_rows))
sop_pdf = squash(text_of_pdf(OUT / "05_Operating_Processes.pdf"))
doc_checks("05_Operating_Processes.pdf", sop_pdf, [("Medium funding gap", money_k(M["cash"]["funding_gap"])), ("Medium peak cash need", money_k(M["cash"]["peak_need"])), ("opening capital", [money(opening), money_k(opening)]), ("gate Month 12 revenue pass", gt("m12_rev_pass"))])

# 09 listing kit
lk = squash(text_of_docx(OUT / "09_Amazon_Listing_Kit.docx")); (QA / "listing.txt").write_text(lk)
lk_items = [("hero price", money(UE["price"], 2)), ("hero net", money(UE["net"], 2)), ("hero net %", pct(UE["net_pct"])), ("hero landed %", pct_auto(UE["landed_pct"])),
            ("Alibaba low landed", money(F["unit_economics"]["cost_bases"]["alibaba_low"]["landed"], 2)), ("Alibaba high landed", money(F["unit_economics"]["cost_bases"]["alibaba_high"]["landed"], 2)),
            ("Alibaba low net", money(F["unit_economics"]["cost_bases"]["alibaba_low"]["net"], 2)), ("Alibaba high net", money(F["unit_economics"]["cost_bases"]["alibaba_high"]["net"], 2)),
            ("duty copper", pct(D["copper"]["rate"])), ("gate contribution pass", gt("d90_cpu_pass")), ("returns Year 1 and Year 2 (Medium)", "Medium " + pct(mm[0]["returns_pct"]) + " of revenue in Year 1, " + pct(mm[18]["returns_pct"]) + " from Year 2"),
            ("insurance trigger month", month_label(M["insurance_trigger_month"]))]
doc_checks("09_Amazon_Listing_Kit.docx", lk, lk_items)
sku_bad = []
for p in F["unit_economics"]["per_sku"]:
    label = p["sku"].replace("_", " ")
    lines = [ln for ln in lk.splitlines() if ln.strip().startswith(label + " ")]
    exp = [money(p["us_list"], 2), money(p["us_sale_price"], 2), money(p["landed_air"], 2), money(p["fba_fee"], 2), money(p["net_air"], 2), pct(p["net_air_pct"]),
           money(0.8 * p["us_sale_price"] - p["landed_air"] - p["fba_fee"], 2), money(0.8 * p["us_sale_price"] - p["landed_ocean"] - p["fba_fee"], 2)]
    if not any(all(e in ln for e in exp) for ln in lines):
        missing = [e for e in exp if not any(e in ln for ln in lines)]
        sku_bad.append(f"{label}: {missing}")
check("09_Amazon_Listing_Kit.docx", "unit economics table: list, sale price, landed air, FBA fee, net air, net %, contribution air and ocean per SKU", not sku_bad, "; ".join(sku_bad))
price_bad = []
for p in F["unit_economics"]["per_sku"]:
    label = p["sku"].replace("_", " ")
    if not re.search(re.escape(label) + r"\s+\d+ ft \(\d+ m\)\s+" + re.escape(money(p["us_list"], 2)) + r"\s+" + re.escape(money(p["us_sale_price"], 2)), lk): price_bad.append(label)
check("09_Amazon_Listing_Kit.docx", "child ASIN table: list and sale price per SKU", not price_bad, "; ".join(price_bad))

# 10 one pagers
op_common = [("opening capital", money(opening)), ("opening capital KRW", fmt.krw_m(opening * FX)),
             ("Medium peak cash need", money_k(M["cash"]["peak_need"])), ("Medium peak month", month_label(M["cash"]["peak_need_month"])), ("Low peak cash need", money_k(L["cash"]["peak_need"])), ("Low peak month", month_label(L["cash"]["peak_need_month"])),
             ("Medium funding gap", money_k(M["cash"]["funding_gap"])), ("Medium funding gap full", money(M["cash"]["funding_gap"])), ("Medium funding gap KRW", fmt.krw_m(M["cash"]["funding_gap"] * FX)),
             ("Low funding gap", money_k(L["cash"]["funding_gap"])), ("Low funding gap full", money(L["cash"]["funding_gap"])), ("Low funding gap KRW", fmt.krw_m(L["cash"]["funding_gap"] * FX)),
             ("Medium funding needed by", month_label(M["cash"]["funding_needed_by_month"])), ("High funding gap", money_k(H["cash"]["funding_gap"])), ("High funding needed by", month_label(H["cash"]["funding_needed_by_month"])),
             ("Medium 18 month revenue", money_m(M["m18"]["revenue"])), ("Medium 18 month net", money_k(M["m18"]["net"])), ("Medium 18 month margin", pct(M["m18"]["net_margin"])), ("Medium breakeven", be(M)),
             ("Low 18 month revenue", money_k(L["m18"]["revenue"])), ("Low 18 month net", money_k(L["m18"]["net"])),
             ("Medium Year 5 revenue", money_m(M["annual"]["Y5"]["revenue"])), ("Medium Year 5 net", money_k(M["annual"]["Y5"]["net"])), ("Low Year 5 revenue", money_m(L["annual"]["Y5"]["revenue"])), ("Low Year 5 net", money_k(L["annual"]["Y5"]["net"])),
             ("Low Month 12 gate", L["gates"]["month12"]["result"]), ("Medium Day 90 gate", M["gates"]["day90"]["result"]), ("duty copper", pct(D["copper"]["rate"])), ("duty fiber", pct(D["fiber"]["rate"])),
             ("Wyoming formation", money(FACTS["entity"]["formation_cost_usd"])), ("Wyoming annual", money(FACTS["entity"]["annual_cost_usd"])), ("Form 5472 penalty", money(FACTS["entity"]["form_5472_penalty_usd"])),
             ("trademark fee", money(FACTS["trademark"]["uspto_fee_per_class_usd"])), ("continuous bond", money(FACTS["logistics"]["broker_bond_annual_usd"])), ("Korean account report threshold", fmt.krw_m_short(FACTS["korea"]["overseas_financial_account_report_threshold_krw"]))]
for fname in ("10_One_Page_Summary_EN.pdf", "10_One_Page_Summary_KO.pdf"):
    t = squash(text_of_pdf(OUT / fname)); (QA / (fname[:-4] + ".txt")).write_text(t)
    check(fname, "one page", t.split("==== raw ====")[0].count("\f") <= 1, f"{t.count(chr(12))} form feeds")
    doc_checks(fname, t, op_common)

# 06 compliance checklist
def xlsx_strings(p):
    w = load_workbook(p, data_only=True); out = []
    for ws in w.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str): out.append(v)
    return w, "\n".join(out)
wb06, t06 = xlsx_strings(OUT / "06_Compliance_Checklist.xlsx")
doc_checks("06_Compliance_Checklist.xlsx", t06, [("opening capital", money(opening)), ("Medium funding gap by month", money_k(M["cash"]["funding_gap"]) + " funding gap by " + month_label(M["cash"]["funding_needed_by_month"])),
    ("insurance trigger month", month_label(M["insurance_trigger_month"]) + " (first month sales pass " + money(FACTS["amazon"]["insurance_trigger_monthly_sales_usd"])), ("duty copper and fiber", "Copper rate " + pct(D["copper"]["rate"]) + ", fiber rate " + pct(D["fiber"]["rate"])),
    ("duty stress", "Stress rate " + pct(D["stress"]["rate"])), ("ocean from month", "from " + month_label(7)), ("Form 5472 penalty", money(FACTS["entity"]["form_5472_penalty_usd"])), ("Wyoming annual", "about " + money(FACTS["entity"]["annual_cost_usd"])),
    ("trademark fee", money(FACTS["trademark"]["uspto_fee_per_class_usd"]) + " per class"), ("continuous bond", "about " + money(FACTS["logistics"]["broker_bond_annual_usd"]) + " per year"), ("seller fee", money(FACTS["amazon"]["professional_fee_monthly_usd"], 2))])
ws06 = wb06["Checklist"]
hdr06 = [c.value for c in ws06[ws06.min_row + 0]] if ws06 else []
check("06_Compliance_Checklist.xlsx", "Summary tab item count check cell reads 0", any(isinstance(r[0], str) and r[0].startswith("Check: area totals") and r[1] == 0 for r in wb06["Summary"].iter_rows(values_only=True)), "check row not 0")

# 07 90 day plan
wb07 = load_workbook(OUT / "07_90_Day_Launch_Plan.xlsx", data_only=True)
def label_pos(row, label, startswith=True):
    for i, v in enumerate(row):
        if isinstance(v, str) and (v.startswith(label) if startswith else label in v): return i
    return None
def row_by_label(ws, label, startswith=True, after_label=None):
    """The cells after the label (the label may sit in any column; merged cells leave None gaps, so only non None cells are returned).
    With after_label the search starts below the first row carrying that label."""
    started = after_label is None
    for row in ws.iter_rows(values_only=True):
        if not started:
            if label_pos(row, after_label) is not None: started = True
            continue
        i = label_pos(row, label, startswith)
        if i is not None: return [v for v in row[i + 1:] if v is not None]
    return None
def row_by_num(ws, n, label_sub, after_label=None):
    """Numbered table row: the cells from the row number on, after an optional section label."""
    started = after_label is None
    for row in ws.iter_rows(values_only=True):
        if not started:
            if label_pos(row, after_label) is not None: started = True
            continue
        for i, v in enumerate(row):
            if v == n and isinstance(row[i + 1] if i + 1 < len(row) else None, str) and label_sub in row[i + 1]: return list(row[i:])
    return None
g = wb07["Gate"]; bad07 = []
def cmp07(label, row, idx, exp):
    if row is None: bad07.append(f"{label}: row not found"); return
    v = row[idx] if idx < len(row) else None
    if not close(v, exp): bad07.append(f"{label}: sheet {sfmt(v)} vs model {sfmt(exp)}")
for n, sub, key, tp, tk in ((3, "TACOS in Month 3", "tacos", "d90_tacos_pass", "d90_tacos_kill"), (5, "Return rate in Month 3", "ret", "d90_ret_pass", "d90_ret_kill"), (6, "Contribution per unit", "cpu", "d90_cpu_pass", "d90_cpu_kill"), (7, "Cash runway", "runway", "d90_run_pass", "d90_run_kill")):
    row = row_by_num(g, n, sub)
    cmp07(f"Day 90 row {n} pass value", row, 4, GT[tp]["value"]); cmp07(f"Day 90 row {n} kill value", row, 5, GT[tk]["value"])
    for idx, sc in ((7, M), (8, L), (9, H)): cmp07(f"Day 90 row {n} {['Medium','Low','High'][idx-7]} model", row, idx, sc["gates"]["day90"]["metrics"][key]["value"])
    cmp07(f"Day 90 row {n} Medium result", row, 10, "PASS" if M["gates"]["day90"]["metrics"][key]["passes"] else ("KILL" if M["gates"]["day90"]["metrics"][key]["kills"] else "REVIEW"))
for n, sub, key, tp, tk in ((1, "Monthly revenue in Month 12", "rev", "m12_rev_pass", "m12_rev_kill"), (2, "Cumulative net at Month 12", "cum", "m12_cum_pass", "m12_cum_kill"), (3, "TACOS in Month 12", "tacos", "m12_tacos_pass", "m12_tacos_kill"), (4, "Return rate in Month 12", "ret", "m12_ret_pass", "m12_ret_kill")):
    row = row_by_num(g, n, sub, after_label="Reference: Month 12 gate")
    cmp07(f"Month 12 row {n} pass value", row, 4, GT[tp]["value"]); cmp07(f"Month 12 row {n} kill value", row, 5, GT[tk]["value"])
    for idx, sc in ((7, M), (8, L), (9, H)): cmp07(f"Month 12 row {n} {['Medium','Low','High'][idx-7]} model", row, idx, sc["gates"]["month12"]["metrics"][key]["value"])
row = row_by_num(g, 5, "Operating cash flow", after_label="Reference: Month 12 gate")
for idx, sc in ((7, M), (8, L), (9, H)): cmp07(f"Month 12 row 5 {['Medium','Low','High'][idx-7]} model", row, idx, sc["gates"]["month12"]["metrics"]["opcf"]["value"])
for label, getter in (("Units sold in Month 3", lambda s: s["monthly"][2]["units"]), ("Revenue in Month 3", lambda s: s["monthly"][2]["revenue"]), ("Net in Month 3", lambda s: s["monthly"][2]["net"]),
                      ("Cumulative net, Months 1 to 3", lambda s: s["monthly"][2]["cum_net"]), ("Contribution per unit after advertising, Month 3", lambda s: s["monthly"][2]["contribution_after_ads"]),
                      ("Cash balance at end of Month 3 before funding", lambda s: s["monthly"][2]["cash_balance"]), ("Cash balance at end of Month 3 after funding", lambda s: s["monthly"][2]["cash_balance_after_funding"]),
                      ("Funding gap in the model", lambda s: s["cash"]["funding_gap"]), ("Funding needed by month", lambda s: month_label(s["cash"]["funding_needed_by_month"])),
                      ("Peak cash need and month", lambda s: money_k(s["cash"]["peak_need"]) + " in " + month_label(s["cash"]["peak_need_month"])),
                      ("18 month revenue and net", lambda s: money_m(s["m18"]["revenue"]) + " and " + money_k(s["m18"]["net"])), ("Cumulative breakeven month", lambda s: be(s))):
    row = row_by_label(g, label)
    for idx, sc in ((0, M), (1, L), (2, H)): cmp07(f"{label} {['Medium','Low','High'][idx]}", row, idx, getter(sc))
cmp07("Overall Day 90 result (Medium model)", row_by_label(g, "Overall Day 90 result"), 0, M["gates"]["day90"]["result"])
cmp07("Overall Month 12 result (Medium model)", row_by_label(g, "Overall Month 12 result"), 0, M["gates"]["month12"]["result"])
check("07_90_Day_Launch_Plan.xlsx", "Gate tab: thresholds, model values and Month 3 context equal figures.json", not bad07, "; ".join(bad07[:8]))
cov = wb07["Cover"]; cov_t = "\n".join(str(c.value) for row in cov.iter_rows() for c in row if isinstance(c.value, str))
doc_checks("07_90_Day_Launch_Plan.xlsx", cov_t, [("opening capital", money(opening)), ("funding gap sentence", money_k(M["cash"]["funding_gap"]) + " needed by " + month_label(M["cash"]["funding_needed_by_month"]) + " (1.2 times the peak need of " + money_k(M["cash"]["peak_need"]) + " in " + month_label(M["cash"]["peak_need_month"]) + ")"),
    ("cash at end of Month 3", money(mm[2]["cash_balance"]) + " before funding; " + money(mm[2]["cash_balance_after_funding"]) + " after funding")])
costs = wb07["Costs"]; bad_c = []
r = row_by_label(costs, "Launch budget", after_label="By cost type")
if r is None or not close(r[2], FACTS["august"]["m18_assumptions"]["launch_one_time"]): bad_c.append("launch budget reference")
inv03 = sum(x for x in vd("monthly_inventory_cash_out")[0:4])
r = row_by_label(costs, "Inventory")
if r is None or not close(r[2], inv03): bad_c.append(f"inventory reference {r[2] if r else None} vs workbook Months 0 to 3 {inv03:,.2f}")
r = row_by_label(costs, "Fixed cost")
if r is None or not close(r[2], 3 * (FACTS["august"]["m18_assumptions"]["fixed_monthly"] + FACTS["amazon"]["professional_fee_monthly_usd"])): bad_c.append("fixed cost reference")
for lab in ("Check: owner total", "Check: week total", "Check: owner and week total"):
    r = row_by_label(costs, lab)
    if r is None or not close(r[0], 0): bad_c.append(lab)
r = row_by_label(costs, "Opening capital (Gantt tab)")
if r is None or not close(r[0], opening): bad_c.append("opening capital")
check("07_90_Day_Launch_Plan.xlsx", "Costs tab: model references (launch $9,500, inventory Months 0 to 3, fixed, opening capital) and check rows read 0", not bad_c, "; ".join(bad_c))
gantt_t = "\n".join(str(c.value) for row in wb07["Gantt"].iter_rows() for c in row if isinstance(c.value, str))
doc_checks("07_90_Day_Launch_Plan.xlsx", gantt_t, [("Gantt: Month 1 units and revenue", "Medium plan " + units(mm[0]["units"]) + " units, revenue " + money(mm[0]["revenue"])), ("Gantt: Month 2 units and revenue", "Medium plan " + units(mm[1]["units"]) + " units, revenue " + money(mm[1]["revenue"])),
    ("Gantt: Month 3 units and revenue", "Medium plan " + units(mm[2]["units"]) + " units, revenue " + money(mm[2]["revenue"])), ("Gantt: funding line", "draws " + money_k(M["cash"]["funding_gap"]) + " (1.2 times the peak need of " + money_k(M["cash"]["peak_need"]) + " in " + month_label(M["cash"]["peak_need_month"])),
    ("Gantt: landed % fiber lever", "from " + pct(D["copper"]["landed_pct_medium"]) + " to " + pct(D["fiber"]["landed_pct_medium"])), ("Gantt: opening capital", money(opening)), ("Gantt: insurance before month", "before " + month_label(M["insurance_trigger_month"]))])

# 08 risk register
wb08 = load_workbook(OUT / "08_Risk_Register.xlsx", data_only=True); kc = wb08["Kill Criteria"]; bad08 = []
def cmp08(label, row, idx, exp):
    if row is None: bad08.append(f"{label}: row not found"); return
    v = row[idx] if idx < len(row) else None
    if not close(v, exp): bad08.append(f"{label}: sheet {sfmt(v)} vs model {sfmt(exp)}")
for k, t in GT.items(): cmp08(f"threshold {k}", row_by_label(kc, t["label"]), 0, t["value"])
for label, getter in (("Funding gap (1.2 times", lambda s: s["cash"]["funding_gap"]), ("Funding needed by", lambda s: s["cash"]["funding_needed_by_month"]), ("Peak cash need month", lambda s: s["cash"]["peak_need_month"]),
                      ("Peak cash need", lambda s: s["cash"]["peak_need"]), ("Cash runway at the Day 90 gate", lambda s: s["cash"]["runway_months_day90"]), ("18 month net", lambda s: s["m18"]["net"]),
                      ("Cumulative breakeven month", lambda s: s["m18"]["breakeven_month"]), ("Month 12 revenue", lambda s: s["gates"]["month12"]["metrics"]["rev"]["value"]), ("Year 5 revenue", lambda s: s["annual"]["Y5"]["revenue"]), ("Year 5 net", lambda s: s["annual"]["Y5"]["net"])):
    row = row_by_label(kc, label)
    for idx, sc in ((0, L), (1, M), (2, H)): cmp08(f"{label} {['Low','Medium','High'][idx]}", row, idx, getter(sc))
for gate, metrics in (("day90", ("tacos", "ret", "cpu", "runway")), ("month12", ("rev", "cum", "tacos", "ret", "opcf"))):
    rows = {kcrit["metric"]: kcrit for kcrit in F["kill_criteria"][gate]}
    model_rows = [kcrit for kcrit in F["kill_criteria"][gate] if kcrit["source"] == "model"]
    for kcrit, mk in zip(model_rows, metrics):
        row = row_by_label(kc, kcrit["metric"], after_label={"day90": "Day 90 gate (end of Month 3). Pass requires", "month12": "Month 12 gate. Pass requires"}[gate])
        for idx, sc in ((3, L), (5, M), (7, H)): cmp08(f"{gate} {mk} {['Low','Medium','High'][(idx-3)//2]}", row, idx, sc["gates"][gate]["metrics"][mk]["value"])
        if row is not None:
            if row[0] != kcrit["pass"] or row[1] != kcrit["kill"]: bad08.append(f"{gate} {mk} threshold wording")
res_rows = [[v for v in row[label_pos(row, "Gate result in the model") + 1:] if v is not None] for row in kc.iter_rows(values_only=True) if label_pos(row, "Gate result in the model") is not None]
check("08_Risk_Register.xlsx", "Kill Criteria tab carries a Day 90 and a Month 12 gate result row from the model", len(res_rows) == 2, f"{len(res_rows)} rows")
for row, gate in zip(res_rows, ("day90", "month12")):
    for idx, sc in ((0, L), (1, M), (2, H)): cmp08(f"{gate} result {['Low','Medium','High'][idx]}", row, idx, sc["gates"][gate]["result"])
chk_rows = [[v for v in row[label_pos(row, "Check: formula result") + 1:] if v is not None] for row in kc.iter_rows(values_only=True) if label_pos(row, "Check: formula result") is not None]
for row in chk_rows:
    if not all(close(row[i], 0) for i in (0, 1, 2)): bad08.append("formula vs model check row not 0")
row = row_by_label(kc, "Plan B if a gate kills")
cmp08("Plan B Y1 revenue", row, 0, PB["Y1"]["revenue"]); cmp08("Plan B Y1 net", row, 2, PB["Y1"]["net"]); cmp08("Plan B Y3 net", row, 4, PB["Y3"]["net"])
check("08_Risk_Register.xlsx", "Kill Criteria tab: thresholds, model values, gate results and Plan B equal figures.json", not bad08, "; ".join(bad08[:8]))
wb08s = xlsx_strings(OUT / "08_Risk_Register.xlsx")[1]
sens = {(r["driver"], r["setting"]): r for r in F["sensitivity"]["table"]}
doc_checks("08_Risk_Register.xlsx", wb08s, [("Cover context: 18 month", "18 month revenue " + money_m(M["m18"]["revenue"]) + ", net " + money_k(M["m18"]["net"]) + " (" + pct(M["m18"]["net_margin"]) + "), cumulative breakeven " + be(M)),
    ("Cover context: Year 5", "Year 5 revenue " + money_m(M["annual"]["Y5"]["revenue"]) + ", net " + money_k(M["annual"]["Y5"]["net"])), ("Cover context: cash", "Peak cash need " + money_k(M["cash"]["peak_need"]) + " in " + month_label(M["cash"]["peak_need_month"]) + "; funding gap " + money_k(M["cash"]["funding_gap"]) + " needed by " + month_label(M["cash"]["funding_needed_by_month"])),
    ("stress case 18 month net", "Medium 18 month net falls from " + money_k(M["m18"]["net"]) + " to " + money_k(FR["duty"]["stress"]["m18_net"])), ("stress case Year 5 net (full recalc)", "Year 5 net from " + money_k(M["annual"]["Y5"]["net"]) + " to " + money_k(FR["duty"]["stress"]["y5_net"])),
    ("duty stress vs copper", "stress case " + pct(D["stress"]["rate"]) + " against the copper base " + pct(D["copper"]["rate"])), ("fiber lever upside", "lost upside of " + money_k(FR["duty"]["fiber"]["m18_net_vs_base"]) + " over 18 months"),
    ("working capital risk", "Medium peak cash need " + money_k(M["cash"]["peak_need"]) + " in " + month_label(M["cash"]["peak_need_month"]) + " against " + money_k(opening) + " opening capital, funding gap " + money_k(M["cash"]["funding_gap"]) + " needed by " + month_label(M["cash"]["funding_needed_by_month"])),
    ("FX 1300 18 month net (full recalc)", "at 1,300 KRW per USD Medium 18 month net is " + money_k(FR["fx"]["1300"]["m18_net"]) + " against " + money_k(M["m18"]["net"]) + " at 1,430 (1,550 gives " + money_k(FR["fx"]["1550"]["m18_net"])),
    ("return rate 5% sensitivity", "at 5% Medium 18 month net falls to " + money_k(sens[("Return rate", "5%")]["m18_net"])), ("ASP minus 10% sensitivity", "ASP minus 10% takes Medium 18 month net to " + money_k(sens[("ASP", "minus 10%")]["m18_net"])),
    ("ad spend plus 3 points sensitivity", "3 points of extra ad spend cost " + money_k(M["m18"]["net"] - sens[("Ad spend", "plus 3 points")]["m18_net"]) + " of Medium 18 month net"),
    ("demand shortfall Low figures", "Low 18 month revenue " + money_k(L["m18"]["revenue"]) + " and net " + money_k(L["m18"]["net"]) + ", Month 12 revenue " + money(L["gates"]["month12"]["metrics"]["rev"]["value"])),
    ("TACOS Day 90 figures", "Low case " + pct_auto(L["gates"]["day90"]["metrics"]["tacos"]["value"]) + ", Medium " + pct_auto(M["gates"]["day90"]["metrics"]["tacos"]["value"])), ("hero FOB", money(UE["fob_usd"], 2) + " FOB")])

# 00 README and 11 Decision Log (present only when their build agents have run)
for fname, items in (("00_README.md", [("Medium 18 month revenue", money_m(M["m18"]["revenue"])), ("Medium 18 month net", money_k(M["m18"]["net"])), ("Medium Year 5 revenue", money_m(M["annual"]["Y5"]["revenue"])), ("Medium funding gap", money_k(M["cash"]["funding_gap"]))]),
                     ("11_Decision_Log.md", [("FX 1300 18 month net", money_k(FR["fx"]["1300"]["m18_net"])), ("FX 1550 18 month net", money_k(FR["fx"]["1550"]["m18_net"])), ("fiber 18 month net", money_k(FR["duty"]["fiber"]["m18_net"])), ("Medium funding gap", money_k(M["cash"]["funding_gap"])), ("opening capital", [money_k(opening), money(opening)])])):
    p = OUT / fname
    if p.exists():
        doc_checks(fname, p.read_text(), items)
    else:
        check(fname, "file not present yet (built by another agent); figure checks skipped, not a failure", True, "absent")

# ---------------------------------------------------------------- 4. hyphen lint
lint = subprocess.run(["python3", str(BUILD / "lint_hyphens.py"), str(OUT)], capture_output=True, text=True)
last = [ln for ln in lint.stdout.splitlines() if ln.startswith("{")]
lint_res = json.loads(last[-1]) if last else {"files_checked": 0, "violating_lines": -1}
check("output/ (all files)", f"lint_hyphens.py: {lint_res.get('files_checked')} files checked, violating_lines {lint_res.get('violating_lines')}", lint_res.get("violating_lines") == 0, lint.stdout[:400].replace("\n", " | "))

# ---------------------------------------------------------------- 5. table and exit
order = []
for d, _, _, _ in RESULTS:
    if d not in order: order.append(d)
lines = ["| Deliverable | Checks | Passed | Failed | Details |", "|---|---|---|---|---|"]
total_fail = 0
for d in order:
    rows = [r for r in RESULTS if r[0] == d]
    passed = sum(1 for r in rows if r[2]); failed = len(rows) - passed; total_fail += failed
    fails = [f"{r[1]}" + (f" ({r[3]})" if r[3] else "") for r in rows if not r[2]]
    absent = [r for r in rows if r[3] == "absent"]
    detail = ("file absent, figure checks skipped (not a failure)" if absent else "all pass") if not fails else "; ".join(fails)[:600]
    lines.append(f"| {d} | {len(rows)} | {passed} | {failed} | {detail.replace('|', '/')} |")
total = len(RESULTS)
verdict = "PASS" if total_fail == 0 else "FAIL"
head = [f"# Tie out result: {verdict}", "", f"Run {datetime.datetime.now().strftime('%B %d, %Y %H:%M')} by build/tie_out.py. {total} checks, {total - total_fail} passed, {total_fail} failed. Tolerance 0.5%. Workbook: output/03_Financial_Model_5Y.xlsx at Scenario 2 (Medium), Duty 1 (Copper), FX 1430.",
        f"Hyphen lint: {lint_res.get('files_checked')} files, {lint_res.get('violating_lines')} violating lines.", ""]
table = "\n".join(head + lines) + "\n"
(BUILD / "tie_out_result.md").write_text(table)
print(table)
if total_fail:
    print("FAILED CHECKS:")
    for r in RESULTS:
        if not r[2]: print(f"  [{r[0]}] {r[1]} :: {r[3]}")
sys.exit(1 if total_fail else 0)
