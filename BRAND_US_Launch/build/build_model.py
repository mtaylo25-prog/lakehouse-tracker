#!/usr/bin/env python3
"""Builds output/03_Financial_Model_5Y.xlsx with live formulas (openpyxl), recalculates it with LibreOffice,
reads it back and checks the calibration targets in BUILD_SPEC sections 2.3 and 2.5.
Rerun: python3 build/build_model.py
Drivers live in build/model_core.py (single source, also used by the pure Python cross check).
"""
import json, os, sys, subprocess, shutil, datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, BarChart, Reference, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.marker import Marker

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, "output", "03_Financial_Model_5Y.xlsx")
sys.path.insert(0, HERE)
import model_core as mc

BUILT = "October 8, 2026"
FONT = "Arial"; SZ = 10
BLUE, BLACK, GREEN, GREY, RED = "0000FF", "000000", "008000", "666666", "C00000"
FILL_IN = PatternFill("solid", fgColor="FFFF00")
FILL_HDR = PatternFill("solid", fgColor="D9E1F2")
FILL_SEC = PatternFill("solid", fgColor="EDEDED")
FILL_HERO = PatternFill("solid", fgColor="FFF2CC")
FILL_OUT = PatternFill("solid", fgColor="E2EFDA")
FMT_MONEY = '$#,##0;($#,##0);0'
FMT_MONEY2 = '$#,##0.00;($#,##0.00);0'
FMT_PCT = '0.0%;(0.0%);0.0%'
FMT_PCT2 = '0.00%;(0.00%);0.00%'
FMT_INT = '#,##0;(#,##0);0'
FMT_NUM2 = '#,##0.00;(#,##0.00);0'
FMT_NUM1 = '#,##0.0;(#,##0.0);0'
FMT_KRW = '"KRW "#,##0;("KRW "#,##0);0'
THIN = Side(style="thin", color="BFBFBF")

wb = Workbook()
wb.remove(wb.active)
NAMES = {}      # defined name -> "'Sheet'!$C$5" (figures_map.json)
SHEETS = {}

def qs(name):
    return "'" + name.replace("'", "''") + "'"

def addsheet(name):
    ws = wb.create_sheet(name)
    SHEETS[name] = ws
    ws.sheet_view.showGridLines = True
    return ws

def defname(name, ws, ref):
    """Register a workbook level defined name. ref like $C$5 or $D$5:$AA$5."""
    full = f"{qs(ws.title)}!{ref}"
    wb.defined_names[name] = DefinedName(name, attr_text=full)
    NAMES[name] = full
    return name

def absref(col, row):
    return f"${L(col)}${row}"

def W(ws, r, c, v, fmt=None, kind="f", bold=False, italic=False, wrap=False, align=None, fill=None, color=None, size=None, border=False):
    """Write a cell with the house style. kind: in (input), f (formula or value), t (text), h (header), s (section), note."""
    cell = ws.cell(row=r, column=c, value=v)
    col = BLACK
    if kind == "in":
        col = BLUE; cell.fill = FILL_IN
    elif kind == "f":
        if isinstance(v, str) and v.startswith("=") and "!" in v:
            col = GREEN
    elif kind == "h":
        bold = True; cell.fill = FILL_HDR
    elif kind == "s":
        bold = True; cell.fill = FILL_SEC
    elif kind == "note":
        italic = True; col = GREY
    elif kind == "out":
        cell.fill = FILL_OUT; bold = True
    if kind != "in" and isinstance(v, str) and v.startswith("=") and "!" in v:
        col = GREEN     # legend: a formula that reads another sheet is green, whatever the cell kind
    if color: col = color
    cell.font = Font(name=FONT, size=size or SZ, bold=bold, italic=italic, color=col)
    if fmt: cell.number_format = fmt
    if wrap or align:
        cell.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical="top" if wrap else None)
    if fill: cell.fill = fill
    if border: cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    return cell

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

def mcol(m):
    """Column index for model month m (Month 0 in column C)."""
    return 3 + m

def ML(m):
    return L(mcol(m))

# ----------------------------------------------------------------------------- Cover
cover = addsheet("Cover")
W(cover, 1, 1, "BRAND LLC: 5 Year Financial Model", bold=True, size=16, kind="t")
W(cover, 2, 1, "US launch of hybrid active optical HDMI cables on Amazon, then eBay and Walmart, B2B installers, a DTC site, Amazon Canada and a pro line", kind="t")
W(cover, 3, 1, f"Build date: {BUILT}. Workbook: output/03_Financial_Model_5Y.xlsx. Generator: build/build_model.py.", kind="note")
r = 5
W(cover, r, 1, "How to use", kind="s"); r += 1
for line in [
    "1. Change the three selectors on the Assumptions tab: Scenario (1 Low, 2 Medium, 3 High), DutyScenario (1 Copper, 2 Fiber, 3 Stress), FXRate (KRW per USD).",
    "2. Every other tab reads the Active column on Assumptions, so one selector drives every sheet. The Scenario Summary tab shows all three scenarios side by side without changing the selector.",
    "3. Edit yellow cells only. Black cells are formulas. Green cells are links to another sheet.",
    "4. Months are model months: Month 1 is the first month with inventory live on Amazon (assumed early 2027). Year 1 is Months 1 to 12.",
    "5. Months 1 to 18 reproduce the August 2026 Amazon model. Year 2 drivers (ocean freight, fixed cost step up, compliance, FBA fee inflation) start in Month 19.",
    "6. Recalculate after edits (Excel does this automatically). The build script recalculates with LibreOffice and checks that there are no formula errors.",
]:
    W(cover, r, 1, line, kind="t", wrap=False); r += 1
r += 1
W(cover, r, 1, "Legend", kind="s"); r += 1
W(cover, r, 1, "Input: yellow fill with blue text", kind="in"); r += 1
W(cover, r, 1, "Formula: black text", kind="t"); r += 1
c = W(cover, r, 1, "Link to another sheet: green text", kind="t"); c.font = Font(name=FONT, size=SZ, color=GREEN); r += 1
W(cover, r, 1, "Key output: green fill", kind="out"); r += 1
r += 1
W(cover, r, 1, "Tabs in order", kind="s"); r += 1
TABS = ["Cover", "Assumptions", "FX", "SKU Unit Economics", "Landed Cost & Tariff Scenarios", "Monthly PL Y1 Y2", "Annual PL Y1 to Y5", "Cash Flow and Working Capital", "Funding and Breakeven", "Sensitivity", "KPI Dashboard", "Levers", "Scenario Summary", "Plan B"]
TAB_DESC = {
    "Cover": "this page", "Assumptions": "Low, Medium, High side by side, the Active column and the three selectors",
    "FX": "KRW per USD input, the three FX scenarios, conversions of capital, the August deal and profit repatriation",
    "SKU Unit Economics": "per SKU cost build (FOB, freight, duty by scenario, fees, FBA, referral, ads, returns, net) and the three cost bases",
    "Landed Cost & Tariff Scenarios": "duty table, landed COGS % build for the active scenario, waterfall data",
    "Monthly PL Y1 Y2": "Months 1 to 24, Amazon US plus the Year 2 channels, 18 month block",
    "Annual PL Y1 to Y5": "Years 1 to 5 by channel with the check rows to the monthly tab",
    "Cash Flow and Working Capital": "inventory purchases, in transit stock, payout lag, cash balance, peak need, runway, funding gap",
    "Funding and Breakeven": "funding plan, Day 90 and Month 12 gates with PASS, REVIEW or KILL, Plan B summary",
    "Sensitivity": "one way sensitivity of 18 month net and Year 5 net, tornado data",
    "KPI Dashboard": "KPI table and the five charts", "Levers": "the twelve levers mapped to the assumption each one moves",
    "Scenario Summary": "all three scenarios side by side, computed live by three parallel engines", "Plan B": "B2B installer distribution plus continued Korea wholesale if Amazon fails",
}
for t in TABS:
    W(cover, r, 1, t, bold=True, kind="t"); W(cover, r, 2, TAB_DESC[t], kind="t"); r += 1
r += 1
W(cover, r, 1, "Basis tags: Evidence based = brief section 1, the August deal, the August model or a verified regulatory fact. Management estimate = everything else.", kind="note"); r += 1
W(cover, r, 1, "Largest unverified assumption: the factory cost. At the Alibaba benchmark FOB ($17 to $22) the hero SKU loses money at $44.99. See SKU Unit Economics.", kind="note", color=RED)
widths(cover, {"A": 34, "B": 120})

# ----------------------------------------------------------------------------- Assumptions
asm = addsheet("Assumptions")
W(asm, 1, 1, "Assumptions", bold=True, size=14, kind="t")
W(asm, 2, 1, "Low, Medium and High side by side. The Active column (F) is what every other sheet reads. Edit yellow cells only.", kind="note")
W(asm, 3, 1, "Scenario selector (1 Low, 2 Medium, 3 High)", bold=True, kind="t"); W(asm, 3, 3, 2, kind="in", fmt="0")
W(asm, 3, 4, '=CHOOSE(Scenario,"Low","Medium","High")', kind="f", bold=True); W(asm, 3, 5, "active scenario name", kind="note")
W(asm, 4, 1, "Duty scenario selector (1 Copper, 2 Fiber, 3 Stress)", bold=True, kind="t"); W(asm, 4, 3, 1, kind="in", fmt="0")
W(asm, 4, 4, '=CHOOSE(DutyScenario,"Copper","Fiber","Stress")', kind="f", bold=True); W(asm, 4, 5, "active duty scenario name", kind="note")
W(asm, 5, 1, "FX rate, KRW per USD (scenarios 1,300 / 1,430 / 1,550)", bold=True, kind="t"); W(asm, 5, 3, mc.FX_DEFAULT, kind="in", fmt="#,##0")
W(asm, 5, 4, '="FX "&TEXT(FXRate,"#,##0")&" KRW per USD"', kind="f")
defname("Scenario", asm, "$C$3"); defname("ScenarioName", asm, "$D$3"); defname("DutyScenario", asm, "$C$4"); defname("DutyScenarioName", asm, "$D$4"); defname("FXRate", asm, "$C$5")
dv1 = DataValidation(type="list", formula1='"1,2,3"', allow_blank=False, showErrorMessage=True, errorTitle="Scenario", error="Enter 1 (Low), 2 (Medium) or 3 (High)")
dv2 = DataValidation(type="list", formula1='"1,2,3"', allow_blank=False, showErrorMessage=True, errorTitle="Duty scenario", error="Enter 1 (Copper), 2 (Fiber) or 3 (Stress)")
dv3 = DataValidation(type="whole", operator="between", formula1="900", formula2="2500", allow_blank=False, showErrorMessage=True, errorTitle="FX", error="Enter KRW per USD between 900 and 2,500 (scenarios 1300, 1430, 1550)")
for dv, ref in ((dv1, "C3"), (dv2, "C4"), (dv3, "C5")):
    asm.add_data_validation(dv); dv.add(ref)
hdr = ["Driver", "Unit", "Low", "Medium", "High", "Active", "Basis", "Note"]
for i, h in enumerate(hdr, 1): W(asm, 7, i, h, kind="h")
ROW = {}
r = 8
UNIT_FMT = {"%": FMT_PCT, "$": FMT_MONEY2, "x": FMT_NUM2, "units": FMT_INT, "month": "0", "pts": FMT_PCT, "months": FMT_NUM1, "days": "0"}

def drow(key, label, unit, vals, basis, note=""):
    global r
    ROW[key] = r
    W(asm, r, 1, label, kind="t"); W(asm, r, 2, unit, kind="t")
    fmt = UNIT_FMT.get(unit, FMT_NUM2)
    for j, v in enumerate(vals):
        W(asm, r, 3 + j, v, kind="in", fmt=fmt)
    W(asm, r, 6, f"=INDEX($C{r}:$E{r},1,Scenario)", kind="f", fmt=fmt, bold=True)
    W(asm, r, 7, basis, kind="t"); W(asm, r, 8, note, kind="t")
    r += 1

def section(title):
    global r
    r += 1
    W(asm, r, 1, title, kind="s")
    for c in range(2, 9): asm.cell(row=r, column=c).fill = FILL_SEC
    r += 1

section("Eighteen month block (August 2026 Amazon model; Months 1 to 18 reproduce it)")
for k in ["asp_y12", "landed_base", "returns_y1", "ads_p1", "ads_p2", "ads_p3", "ads_p4", "fixed_y1", "unit_mult"]:
    d = mc.D[k]; drow(k, d["label"], d["unit"], d["vals"], d["basis"], d["note"])
SHARED_LBL = {"referral": ("Amazon referral fee %", "%", mc.EB), "fba": ("FBA fulfillment fee per unit, Year 1 (hero basis)", "$", mc.EB), "storage": ("Storage and inbound % of revenue", "%", mc.EB),
              "seller_fee": ("Professional seller fee per month", "$", mc.EB), "launch": ("One time launch costs in Month 1", "$", mc.EB)}
for k, (lbl, unit, basis) in SHARED_LBL.items():
    drow(k, lbl, unit, [mc.SHARED[k]] * 3, basis, "August model; same in every scenario")

section("Year 2 to Year 5 drivers (Year 2 changes start in Month 19, after the August horizon)")
for k in ["y2_start", "ads_p5", "m19_add", "lost_q_factor", "lost_q_start", "returns_y2", "fba_infl", "landed_chg_y2", "landed_chg_y3", "landed_chg_y4", "landed_chg_y5", "g_y3", "g_y4", "g_y5", "asp_chg", "roadmap_uplift",
          "ads_y3", "ads_y4", "ads_y5", "fixed_y2", "fixed_y3", "fixed_y4", "fixed_y5", "ebay_pct_y2", "ebay_pct_y3", "ebay_pct_y4", "ebay_pct_y5",
          "b2b_y2", "b2b_y3", "b2b_y4", "b2b_y5", "dtc_y3", "dtc_y4", "dtc_y5", "ca_y4", "ca_y5", "pro_y4", "pro_y5"]:
    d = mc.D[k]; drow(k, d["label"], d["unit"], d["vals"], d["basis"], d["note"])

section("Channel economics (same in every scenario; management estimates)")
CH_LBL = {"ebay_referral": ("eBay and Walmart referral %", "%"), "ebay_mcf": ("eBay and Walmart multi channel fulfillment per unit", "$"), "ebay_ads": ("eBay and Walmart ads % of revenue", "%"), "ebay_returns": ("eBay and Walmart returns %", "%"),
          "b2b_price_factor": ("B2B price as a share of retail ASP", "%"), "b2b_3pl": ("B2B 3PL pick pack ship per unit", "$"), "b2b_commission": ("B2B sales commission %", "%"), "b2b_marketing": ("B2B marketing %", "%"), "b2b_returns": ("B2B returns %", "%"),
          "dtc_payment": ("DTC payment processing %", "%"), "dtc_ship": ("DTC fulfillment and shipping per unit", "$"), "dtc_ads": ("DTC ads %", "%"), "dtc_returns": ("DTC returns %", "%"),
          "ca_asp_factor": ("Amazon Canada ASP as a share of US ASP (8% below)", "%"), "ca_fba": ("Amazon Canada FBA fee per unit", "$"), "ca_referral": ("Amazon Canada referral %", "%"), "ca_ads": ("Amazon Canada ads %", "%"), "ca_returns": ("Amazon Canada returns %", "%"), "ca_duty": ("Duty on US to Canada transfer [PLACEHOLDER: broker check]", "%"),
          "pro_asp": ("Pro line ASP", "$"), "pro_landed": ("Pro line landed COGS %", "%"), "pro_b2b_share": ("Pro line share sold B2B (rest on Amazon)", "%"), "pro_fba": ("Pro line FBA fee per unit (large tier)", "$"), "pro_amz_ads": ("Pro line Amazon ads %", "%"), "pro_amz_returns": ("Pro line Amazon returns %", "%"), "pro_storage": ("Pro line Amazon storage and inbound %", "%")}
for k, (lbl, unit) in CH_LBL.items():
    drow(k, lbl, unit, [mc.CH[k]] * 3, mc.ME, "")

section("Compliance costs per year from Year 2 (Year 1 is inside the August fixed 850 per month and the 9,500 launch costs)")
COMP_LBL = {"form5472": ("Form 5472 preparation (with pro forma Form 1120)", "$", mc.ME), "wyoming": ("Wyoming annual report", "$", mc.EB), "bond": ("Customs continuous bond renewal", "$", mc.EB), "insurance": ("Product liability insurance ($1M cover) per year once monthly sales pass $10K", "$", mc.ME), "trademark_y5": ("Trademark maintenance in Year 5 [PLACEHOLDER: attorney quote]", "$", mc.ME)}
for k, (lbl, unit, basis) in COMP_LBL.items():
    drow(k, lbl, unit, [mc.COMPLIANCE[k]] * 3, basis, "")

section("Working capital")
for k in ["cover", "ocean_from", "lead_air", "lead_ocean", "inv_days"]:
    d = mc.D[k]; drow(k, d["label"], d["unit"], d["vals"], d["basis"], d["note"])
drow("opening_capital", "Opening capital (self funded, August ask)", "$", [mc.OPENING_CAPITAL] * 3, mc.EB, "DECISIONS.md")
drow("deposit_pct", "PO deposit at order (balance at shipment)", "%", [mc.DEPOSIT_PCT] * 3, mc.ME, "30% at order, 70% at shipment")
drow("payout_lag", "Amazon payout lag, days (receivable = lag/30 of the month's net proceeds)", "days", [mc.PAYOUT_LAG_DAYS] * 3, mc.EB, "brief")
drow("recv_days", "Receivable days of revenue, Years 3 to 5", "days", [mc.RECV_DAYS_ANNUAL] * 3, mc.ME, "")
drow("cushion", "Funding gap cushion on the negative minimum balance", "x", [mc.FUNDING_CUSHION] * 3, mc.ME, "1.2 times")
drow("fx_base", "FX embedded in the KRW cost base, KRW per USD", "x", [mc.FX_DEFAULT] * 3, mc.EB, "the only verified cost data is in KRW; a USD factory quote removes the FX effect")
for k in ["fx_base"]: asm.cell(row=ROW[k], column=6).number_format = "#,##0"
for k in ["fx_base"]:
    for c in (3, 4, 5): asm.cell(row=ROW[k], column=c).number_format = "#,##0"

section("Gate thresholds (spec 2.7; same in every scenario). The Funding and Breakeven tab and the three Scenario Summary engines read these cells, so one edit moves every PASS and KILL test and its wording")
for k, lbl, unit, val in mc.GATES:
    drow(k, lbl, unit, [val] * 3, mc.ME, "spec 2.7 (the brief asks for the gates in numbers; the numbers are management estimates)")
    if unit == "$" and abs(val) >= 1000:
        for c in (3, 4, 5, 6): asm.cell(row=ROW[k], column=c).number_format = FMT_MONEY
    defname("gate_threshold_" + k[2:], asm, f"$F${ROW[k]}")

section("Unit ramp by month (Amazon US). Column I holds the August Medium ramp; Low and High scale it; Months 19 to 24 extend it; the Low lost quarter applies by formula")
W(asm, r, 1, "Month", kind="h"); W(asm, r, 2, "Index", kind="h"); W(asm, r, 3, "Low", kind="h"); W(asm, r, 4, "Medium", kind="h"); W(asm, r, 5, "High", kind="h"); W(asm, r, 6, "Active", kind="h"); W(asm, r, 7, "Basis", kind="h"); W(asm, r, 8, "Note", kind="h"); W(asm, r, 9, "August Medium ramp (input)", kind="h")
r += 1
RAMP = {}
for m in range(1, 25):
    RAMP[m] = r
    W(asm, r, 1, f"Month {m}", kind="t"); W(asm, r, 2, m, kind="t", fmt="0")
    if m <= 18:
        W(asm, r, 9, mc.MEDIUM_UNITS_18[m - 1], kind="in", fmt=FMT_INT)
        for j, col in enumerate("CDE"):
            W(asm, r, 3 + j, f"=ROUND($I{r}*{col}${ROW['unit_mult']}*IF(AND($B{r}>={col}${ROW['lost_q_start']},$B{r}<={col}${ROW['lost_q_start']}+2),{col}${ROW['lost_q_factor']},1),0)", kind="f", fmt=FMT_INT)
        W(asm, r, 7, mc.EB if True else "", kind="t"); W(asm, r, 8, "August deck ramp times the scenario multiplier", kind="t")
    else:
        k = m - 18
        for j, col in enumerate("CDE"):
            W(asm, r, 3 + j, f"=ROUND(({col}${RAMP[18]}+{col}${ROW['m19_add']}*{k})*IF(AND($B{r}>={col}${ROW['lost_q_start']},$B{r}<={col}${ROW['lost_q_start']}+2),{col}${ROW['lost_q_factor']},1),0)", kind="f", fmt=FMT_INT)
        W(asm, r, 7, mc.ME, kind="t"); W(asm, r, 8, "Month 18 plus the monthly addition; Low flat", kind="t")
    W(asm, r, 6, f"=INDEX($C{r}:$E{r},1,Scenario)", kind="f", fmt=FMT_INT, bold=True)
    r += 1
asm.freeze_panes = "C8"
widths(asm, {"A": 62, "B": 8, "C": 12, "D": 12, "E": 12, "F": 12, "G": 20, "H": 90, "I": 14})

def A(key, col="F"):
    """Reference to an Assumptions cell. col F = Active; C, D, E = Low, Medium, High (parallel engines only)."""
    if isinstance(key, tuple) and key[0] == "ramp":
        return f"Assumptions!${col}${RAMP[key[1]]}"
    return f"Assumptions!${col}${ROW[key]}"

# ----------------------------------------------------------------------------- SKU Unit Economics
sku = addsheet("SKU Unit Economics")
W(sku, 1, 1, "SKU Unit Economics", bold=True, size=14, kind="t")
W(sku, 2, 1, "FOB China = H2 landed Korea cost in KRW divided by FX times the FOB factor. Duty columns use the duty table. Air and ocean landed costs use the active duty scenario. Hero row highlighted.", kind="note")
cr = 17
W(sku, cr, 1, "Cost build constants", kind="s"); cr += 1
CONST = {}
def crow(key, label, val, fmt, basis, note=""):
    global cr
    CONST[key] = cr
    W(sku, cr, 1, label, kind="t"); W(sku, cr, 2, val, kind="in", fmt=fmt); W(sku, cr, 3, basis, kind="t"); W(sku, cr, 4, note, kind="t"); cr += 1
crow("fob_factor", "FOB factor on the H2 Korea landed cost (removes Korea inbound freight and handling)", mc.FOB_FACTOR, FMT_NUM2, mc.ME)
crow("air", "Air freight DDP per kg packaged", mc.AIR_PER_KG, FMT_MONEY2, mc.ME, "consolidated express, first lots")
crow("ocean", "Ocean LCL plus drayage per kg packaged", mc.OCEAN_PER_KG, FMT_MONEY2, mc.ME)
crow("lb_kg", "Kilograms per pound", mc.LB_TO_KG, "0.0000", mc.EB)
crow("mpf", "Merchandise processing and harbor maintenance fees, % of FOB, ocean only", mc.MPF_HMF_OCEAN, FMT_PCT2, mc.ME)
crow("broker", "Broker entry fee and continuous bond, amortized per unit", mc.BROKER_PER_UNIT, FMT_MONEY2, mc.ME, "$150 per entry plus $500 per year bond at launch volumes")
crow("inbound", "FBA inbound placement and prep per unit", mc.INBOUND_PER_UNIT, FMT_MONEY2, mc.ME)
crow("qc", "Third party QC and FNSKU labeling at the factory per unit", mc.QC_PER_UNIT, FMT_MONEY2, mc.ME)
crow("ue_ads", "Advertising % of price for unit economics (August hero: $5.40 on $44.99)", mc.HERO_AUG["ads_pct"], FMT_PCT, mc.EB)
crow("ue_ret", "Returns and storage reserve % of price (3% returns plus 2% storage and inbound)", mc.HERO_AUG["returns_storage_pct"], FMT_PCT, mc.EB)
crow("ue_fba_aug", "FBA fee used in the August hero unit economics", mc.HERO_AUG["fba"], FMT_MONEY2, mc.EB)
crow("alib_lo", "Alibaba benchmark FOB, low, comparable 30 ft (9 m) AOC", mc.ALIBABA_LOW, FMT_MONEY2, mc.EB, "quote, flagged")
crow("alib_hi", "Alibaba benchmark FOB, high", mc.ALIBABA_HIGH, FMT_MONEY2, mc.EB, "quote, flagged")
C = lambda k: f"$B${CONST[k]}"
DUTY_SHEET = "Landed Cost & Tariff Scenarios"   # 30 characters: Excel limits sheet names to 31, so the spec name (32) is shortened with an ampersand
DQ = qs(DUTY_SHEET)
# duty rate cells on the Landed tab are rows 5,6,7 column G (written below); fixed addresses
DUTY_RATE = {0: f"{DQ}!$G$5", 1: f"{DQ}!$G$6", 2: f"{DQ}!$G$7"}
SKU_HDR = ["SKU", "Spec (internal only)", "Length m", "H2 cost KRW", "FOB USD", "US list", "US sale price (ASP)", "Packaged lb", "kg", "Air freight", "Ocean freight", "Duty copper", "Duty fiber", "Duty stress", "Duty active", "MPF and HMF (ocean)", "Broker and bond", "FBA inbound and prep", "QC and labels",
           "Landed air", "Landed ocean", "Landed air %", "Landed ocean %", "FBA fee", "Referral 15%", "Ads", "Returns and storage", "Net air", "Net air %", "Net ocean", "Net ocean %", "Launch SKU"]
for i, h in enumerate(SKU_HDR, 1): W(sku, 4, i, h, kind="h", wrap=True)
sku.row_dimensions[4].height = 42
SKU_ROW = {}
for i, s in enumerate(mc.SKUS):
    rr = 5 + i; SKU_ROW[s[0]] = rr
    fill = FILL_HERO if s[0] == mc.HERO else None
    W(sku, rr, 1, s[0], kind="t", bold=(s[0] == mc.HERO), fill=fill); W(sku, rr, 2, s[1], kind="t", fill=fill); W(sku, rr, 3, s[2], kind="in", fmt="0")
    W(sku, rr, 4, s[3], kind="in", fmt=FMT_KRW)
    W(sku, rr, 5, f"=D{rr}/FXRate*{C('fob_factor')}", fmt=FMT_MONEY2)
    W(sku, rr, 6, s[4], kind="in", fmt=FMT_MONEY2); W(sku, rr, 7, s[5], kind="in", fmt=FMT_MONEY2); W(sku, rr, 8, s[6], kind="in", fmt=FMT_NUM2)
    W(sku, rr, 9, f"=H{rr}*{C('lb_kg')}", fmt="0.000"); W(sku, rr, 10, f"=I{rr}*{C('air')}", fmt=FMT_MONEY2); W(sku, rr, 11, f"=I{rr}*{C('ocean')}", fmt=FMT_MONEY2)
    W(sku, rr, 12, f"=E{rr}*{DUTY_RATE[0]}", fmt=FMT_MONEY2); W(sku, rr, 13, f"=E{rr}*{DUTY_RATE[1]}", fmt=FMT_MONEY2); W(sku, rr, 14, f"=E{rr}*{DUTY_RATE[2]}", fmt=FMT_MONEY2)
    W(sku, rr, 15, f"=INDEX(L{rr}:N{rr},1,DutyScenario)", fmt=FMT_MONEY2)
    W(sku, rr, 16, f"=E{rr}*{C('mpf')}", fmt=FMT_MONEY2); W(sku, rr, 17, f"={C('broker')}", fmt=FMT_MONEY2); W(sku, rr, 18, f"={C('inbound')}", fmt=FMT_MONEY2); W(sku, rr, 19, f"={C('qc')}", fmt=FMT_MONEY2)
    W(sku, rr, 20, f"=E{rr}+J{rr}+O{rr}+Q{rr}+R{rr}+S{rr}", fmt=FMT_MONEY2); W(sku, rr, 21, f"=E{rr}+K{rr}+O{rr}+P{rr}+Q{rr}+R{rr}+S{rr}", fmt=FMT_MONEY2)
    W(sku, rr, 22, f"=T{rr}/G{rr}", fmt=FMT_PCT); W(sku, rr, 23, f"=U{rr}/G{rr}", fmt=FMT_PCT)
    W(sku, rr, 24, s[7], kind="in", fmt=FMT_MONEY2); W(sku, rr, 25, f"=G{rr}*{A('referral')}", fmt=FMT_MONEY2); W(sku, rr, 26, f"=G{rr}*{C('ue_ads')}", fmt=FMT_MONEY2); W(sku, rr, 27, f"=G{rr}*{C('ue_ret')}", fmt=FMT_MONEY2)
    W(sku, rr, 28, f"=G{rr}-T{rr}-X{rr}-Y{rr}-Z{rr}-AA{rr}", fmt=FMT_MONEY2, bold=True); W(sku, rr, 29, f"=AB{rr}/G{rr}", fmt=FMT_PCT)
    W(sku, rr, 30, f"=G{rr}-U{rr}-X{rr}-Y{rr}-Z{rr}-AA{rr}", fmt=FMT_MONEY2, bold=True); W(sku, rr, 31, f"=AD{rr}/G{rr}", fmt=FMT_PCT)
    W(sku, rr, 32, s[8], kind="in")
    if fill:
        for c in range(1, 33):
            cell = sku.cell(row=rr, column=c)
            if cell.fill is None or cell.fill.fgColor is None or not str(cell.fill.fgColor.rgb).endswith("FFFF00"):
                cell.fill = FILL_HERO     # inputs on the hero row keep the yellow input fill per the Cover legend
            f = cell.font; cell.font = Font(name=f.name, size=f.size, bold=True, italic=f.italic, color=f.color)
            cell.border = Border(top=THIN, bottom=THIN)
W(sku, 15, 1, "Basis: H2 cost KRW evidence based (August deal). FOB factor, freight, fees, weights, US prices and FBA fees are management estimates pending the Amazon rate card check. The spec column is internal: packaging and listings use the category names (Ultra High Speed HDMI Cable; Premium High Speed HDMI Cable pending certification).", kind="note")
HR = SKU_ROW[mc.HERO]
# hero August block
hr0 = cr + 1
W(sku, hr0, 1, "Hero SKU unit economics, August model (evidence based): 8K AOC 10M on Amazon US", kind="s"); 
HB = {}
rows = [("price", "Price", mc.HERO_AUG["price"], "in", FMT_MONEY2), ("list", "List price", 59.99, "in", FMT_MONEY2),
        ("landed_pct", "Landed cost % (assumed 38% pending the customs ruling)", mc.HERO_AUG["landed_pct"], "in", FMT_PCT),
        ("landed", "Landed cost", None, "f", FMT_MONEY2), ("referral_pct", "Referral %", None, "f", FMT_PCT), ("referral", "Referral", None, "f", FMT_MONEY2),
        ("fba", "FBA fulfillment", None, "f", FMT_MONEY2), ("ads_pct", "Advertising % (blended)", None, "f", FMT_PCT), ("ads", "Advertising", None, "f", FMT_MONEY2),
        ("ret_pct", "Returns and storage reserve %", None, "f", FMT_PCT), ("ret", "Returns and storage reserve", None, "f", FMT_MONEY2),
        ("net", "Net per unit", None, "f", FMT_MONEY2), ("net_pct", "Net %", None, "f", FMT_PCT), ("korea", "Korea wholesale gross on the same unit", mc.HERO_AUG["korea_gross"], "in", FMT_MONEY2),
        ("cpu_b", "Contribution per unit before advertising", None, "f", FMT_MONEY2)]
rr = hr0 + 1
for key, lbl, val, kind, fmt in rows:
    HB[key] = rr
    W(sku, rr, 1, lbl, kind="t")
    if kind == "in": W(sku, rr, 2, val, kind="in", fmt=fmt)
    rr += 1
W(sku, HB["landed"], 2, f"=B{HB['price']}*B{HB['landed_pct']}", fmt=FMT_MONEY2)
W(sku, HB["referral_pct"], 2, f"={A('referral')}", fmt=FMT_PCT)
W(sku, HB["referral"], 2, f"=B{HB['price']}*B{HB['referral_pct']}", fmt=FMT_MONEY2)
W(sku, HB["fba"], 2, f"={C('ue_fba_aug')}", fmt=FMT_MONEY2)
W(sku, HB["ads_pct"], 2, f"={C('ue_ads')}", fmt=FMT_PCT)
W(sku, HB["ads"], 2, f"=B{HB['price']}*B{HB['ads_pct']}", fmt=FMT_MONEY2)
W(sku, HB["ret_pct"], 2, f"={C('ue_ret')}", fmt=FMT_PCT)
W(sku, HB["ret"], 2, f"=B{HB['price']}*B{HB['ret_pct']}", fmt=FMT_MONEY2)
W(sku, HB["net"], 2, f"=B{HB['price']}-B{HB['landed']}-B{HB['referral']}-B{HB['fba']}-B{HB['ads']}-B{HB['ret']}", fmt=FMT_MONEY2, kind="out")
W(sku, HB["net_pct"], 2, f"=B{HB['net']}/B{HB['price']}", fmt=FMT_PCT, kind="out")
W(sku, HB["cpu_b"], 2, f"=B{HB['net']}+B{HB['ads']}", fmt=FMT_MONEY2)
for key, nm in [("price", "ue_hero_price"), ("list", "ue_hero_list"), ("landed", "ue_hero_landed_cost"), ("landed_pct", "ue_hero_landed_pct"), ("referral", "ue_hero_referral"), ("fba", "ue_hero_fba"), ("ads", "ue_hero_ads"), ("ret", "ue_hero_returns_storage"), ("net", "ue_hero_net"), ("net_pct", "ue_hero_net_pct"), ("korea", "ue_hero_korea_wholesale_gross"), ("cpu_b", "ue_hero_contribution_before_ads")]:
    defname(nm, sku, f"$B${HB[key]}")
W(sku, rr, 1, "Hero FOB share of the August price (hero FOB USD at the FX embedded in the KRW cost base, 1,430, divided by $44.99): drives the duty and FX effect on landed COGS %. Held at the base FX so the FX delta formula (spec 2.4) is the only FX effect", kind="t"); W(sku, rr, 2, f"=$D${HR}/{A('fx_base')}*{C('fob_factor')}/$B${HB['price']}", fmt=FMT_PCT2, bold=True)
defname("HeroFobShare", sku, f"$B${rr}"); HFS_ROW = rr; rr += 1
W(sku, rr, 1, "Hero FOB USD at the active FX rate", kind="t"); W(sku, rr, 2, f"=E{HR}", fmt=FMT_MONEY2); defname("HeroFobUsd", sku, f"$B${rr}"); HFU_ROW = rr; rr += 1
# cost bases
rr += 1
W(sku, rr, 1, "Three cost bases for the hero SKU at $44.99 (air freight, active duty scenario). The gap between (a) and (c) is the single largest unverified assumption in the plan.", kind="s"); rr += 1
for i, h in enumerate(["Cost basis", "FOB USD", "Freight, duty, fees", "Landed cost", "Landed %", "Net per unit at $44.99 (August fees)", "Net %", "Basis"], 1): W(sku, rr, i, h, kind="h")
rr += 1
CB = {}
def cbrow(key, label, fob_f, basis):
    global rr
    CB[key] = rr
    W(sku, rr, 1, label, kind="t"); W(sku, rr, 2, fob_f, fmt=FMT_MONEY2)
    W(sku, rr, 3, f"=J{HR}+B{rr}*{DQ}!$G$9+Q{HR}+R{HR}+S{HR}", fmt=FMT_MONEY2)
    W(sku, rr, 4, f"=B{rr}+C{rr}", fmt=FMT_MONEY2); W(sku, rr, 5, f"=D{rr}/B{HB['price']}", fmt=FMT_PCT)
    W(sku, rr, 6, f"=B{HB['price']}-D{rr}-B{HB['referral']}-B{HB['fba']}-B{HB['ads']}-B{HB['ret']}", fmt=FMT_MONEY2, bold=True); W(sku, rr, 7, f"=F{rr}/B{HB['price']}", fmt=FMT_PCT)
    W(sku, rr, 8, basis, kind="t"); rr += 1
cbrow("h2", "(a) H2 cost base (August deal KRW, FOB factor, bottom up)", f"=E{HR}", "Evidence based cost, management estimate build")
CB["august"] = rr
W(sku, rr, 1, "(b) August model assumption (38% of price)", kind="t"); W(sku, rr, 2, "", kind="t"); W(sku, rr, 3, "", kind="t"); W(sku, rr, 4, f"=B{HB['landed']}", fmt=FMT_MONEY2); W(sku, rr, 5, f"=D{rr}/B{HB['price']}", fmt=FMT_PCT)
W(sku, rr, 6, f"=B{HB['price']}-D{rr}-B{HB['referral']}-B{HB['fba']}-B{HB['ads']}-B{HB['ret']}", fmt=FMT_MONEY2, bold=True); W(sku, rr, 7, f"=F{rr}/B{HB['price']}", fmt=FMT_PCT); W(sku, rr, 8, "Evidence based (August deck assumption)", kind="t"); rr += 1
cbrow("alibaba_low", "(c) Alibaba benchmark FOB $17 (comparable 30 ft AOC)", f"={C('alib_lo')}", "Evidence based quote, flagged")
cbrow("alibaba_high", "(c) Alibaba benchmark FOB $22", f"={C('alib_hi')}", "Evidence based quote, flagged")
for k in ("h2", "august", "alibaba_low", "alibaba_high"):
    defname(f"cost_base_{k}_landed", sku, f"$D${CB[k]}"); defname(f"cost_base_{k}_net", sku, f"$F${CB[k]}")
W(sku, rr, 1, "At the Alibaba cost the hero SKU loses money at $44.99. Verify the factory FOB quote in USD before the first purchase order.", kind="note", color=RED); rr += 1
defname("sku_table", sku, f"$A$5:$AF${4 + len(mc.SKUS)}")
sku.freeze_panes = "C5"
widths(sku, {"A": 44, "B": 26, "C": 12, "D": 14, "E": 10, "F": 10, "G": 12, "H": 10})
for c in range(9, 33): sku.column_dimensions[L(c)].width = 11

# ----------------------------------------------------------------------------- Landed Cost and Tariff Scenarios
lan = addsheet(DUTY_SHEET)
W(lan, 1, 1, "Landed Cost and Tariff Scenarios", bold=True, size=14, kind="t")
W(lan, 2, 1, "Duty is a selector (DutyScenario on Assumptions). The 12.5% Section 301 forced labor rate applies whichever classification the CBP binding ruling gives.", kind="note")
for i, h in enumerate(["Scenario", "HTS", "MFN", "Section 301 list rate", "Section 301 forced labor (late July 2026)", "Pending second Section 301 action", "Total duty rate", "Basis"], 1): W(lan, 4, i, h, kind="h", wrap=True)
lan.row_dimensions[4].height = 42
for i, d in enumerate(mc.DUTY):
    rr = 5 + i
    W(lan, rr, 1, d[0], kind="t", bold=True); W(lan, rr, 2, d[1], kind="t")
    for j in range(4): W(lan, rr, 3 + j, d[2 + j], kind="in", fmt=FMT_PCT)
    W(lan, rr, 7, f"=SUM(C{rr}:F{rr})", fmt=FMT_PCT, bold=True); W(lan, rr, 8, d[6], kind="t")
d = mc.DUTY_FIBER_LIST3_CHECK
W(lan, 8, 1, d[0] + " (check row only, not selectable)", kind="t"); W(lan, 8, 2, d[1], kind="t")
for j in range(4): W(lan, 8, 3 + j, d[2 + j], kind="in", fmt=FMT_PCT)
W(lan, 8, 7, "=SUM(C8:F8)", fmt=FMT_PCT, bold=True); W(lan, 8, 8, "Shown as a check row only", kind="t")
W(lan, 9, 1, "Active duty rate (selected scenario)", kind="t", bold=True); W(lan, 9, 7, "=INDEX($G$5:$G$7,DutyScenario,1)", fmt=FMT_PCT, kind="out")
W(lan, 10, 1, "Copper rate (the rate embedded in the base landed COGS %)", kind="t"); W(lan, 10, 7, "=$G$5", fmt=FMT_PCT)
defname("DutyRateActive", lan, "$G$9"); defname("DutyRateCopper", lan, "$G$10")
defname("duty_copper_rate", lan, "$G$5"); defname("duty_fiber_rate", lan, "$G$6"); defname("duty_stress_rate", lan, "$G$7"); defname("duty_fiber_list3_check_rate", lan, "$G$8")
# landed % build
W(lan, 12, 1, "Landed COGS % build for the active scenario (P&L)", kind="s")
LB = {}
items = [("base", "Base landed COGS % (Assumptions, Active: Low 43%, Medium 38%, High 32%)", f"={A('landed_base')}", FMT_PCT),
         ("duty_delta", "Duty delta = (active duty rate minus copper rate) times hero FOB share of price", "=(DutyRateActive-DutyRateCopper)*HeroFobShare", FMT_PCT2),
         ("freight_delta", "Freight mode delta (points; 0 because the base already assumes air then ocean; the Year 2 landed change carries the ocean benefit)", 0.0, FMT_PCT2),
         ("fx_delta", "FX delta = hero FOB share times (FX embedded in the cost base divided by FXRate, minus 1)", f"=HeroFobShare*({A('fx_base')}/FXRate-1)", FMT_PCT2),
         ("active", "Active landed COGS % used by the P&L", None, FMT_PCT)]
rr = 13
for key, lbl, f, fmt in items:
    LB[key] = rr
    W(lan, rr, 1, lbl, kind="t")
    if key == "freight_delta": W(lan, rr, 3, f, kind="in", fmt=fmt)
    elif key == "active": W(lan, rr, 3, f"=C{LB['base']}+C{LB['duty_delta']}+C{LB['freight_delta']}+C{LB['fx_delta']}", fmt=fmt, kind="out")
    else: W(lan, rr, 3, f, fmt=fmt)
    rr += 1
defname("LandedPctActive", lan, f"$C${LB['active']}"); defname("LandedFreightDelta", lan, f"$C${LB['freight_delta']}"); defname("LandedFxDelta", lan, f"$C${LB['fx_delta']}"); defname("LandedDutyDelta", lan, f"$C${LB['duty_delta']}")
W(lan, rr, 1, "Note: the High base of 32% already assumes the fiber ruling, ocean freight and volume pricing. Selecting Fiber duty with High counts the ruling twice. Use Medium plus Fiber to see the ruling effect (about 33.5%, the August lever 38% to 33%).", kind="note"); rr += 2
# scenario x duty table
W(lan, rr, 1, "Landed COGS % by scenario and duty scenario at the active FX rate (used by the Scenario Summary engines)", kind="s"); rr += 1
for i, h in enumerate(["Scenario", "Copper", "Fiber", "Stress", "At active duty", "Hero net per unit at active duty (August fees, $44.99)"], 1): W(lan, rr, i, h, kind="h", wrap=True)
rr += 1
LS = {}
for i, sname in enumerate(mc.SCEN_NAMES):
    col = "CDE"[i]
    LS[sname] = rr
    W(lan, rr, 1, sname, kind="t", bold=True)
    for j in range(3):
        W(lan, rr, 2 + j, f"={A('landed_base', col)}+($G${5 + j}-DutyRateCopper)*HeroFobShare+LandedFreightDelta+LandedFxDelta", fmt=FMT_PCT)
    W(lan, rr, 5, f"=INDEX(B{rr}:D{rr},1,DutyScenario)", fmt=FMT_PCT, bold=True)
    W(lan, rr, 6, f"=ue_hero_price*(1-E{rr})-ue_hero_referral-ue_hero_fba-ue_hero_ads-ue_hero_returns_storage", fmt=FMT_MONEY2)
    defname(f"LandedPct{sname}", lan, f"$E${rr}")
    rr += 1
for j, dn in enumerate(["copper", "fiber", "stress"]):
    defname(f"duty_{dn}_landed_pct_medium", lan, f"${'BCD'[j]}${LS['Medium']}")
rr += 1
# waterfall and hero net per duty scenario (Medium base)
W(lan, rr, 1, "Hero net per unit by duty scenario (Medium base 38%, August fees, price $44.99) and the waterfall data for the chart", kind="s"); rr += 1
for i, h in enumerate(["Duty scenario", "Duty rate", "Landed %", "Hero net per unit", "Waterfall label", "Invisible base", "Bar value", "Bar color"], 1): W(lan, rr, i, h, kind="h")
rr += 1
HN = {}
for j, dn in enumerate(["Copper", "Fiber", "Stress"]):
    HN[dn] = rr
    W(lan, rr, 1, dn, kind="t", bold=True); W(lan, rr, 2, f"=$G${5 + j}", fmt=FMT_PCT); W(lan, rr, 3, f"={'BCD'[j]}{LS['Medium']}", fmt=FMT_PCT)
    W(lan, rr, 4, f"=ue_hero_price*(1-C{rr})-ue_hero_referral-ue_hero_fba-ue_hero_ads-ue_hero_returns_storage", fmt=FMT_MONEY2, bold=True)
    defname(f"duty_{dn.lower()}_hero_net", lan, f"$D${rr}")
    rr += 1
wf0 = rr
WF = [("Copper net", f"=0", f"=D{HN['Copper']}", "value"), ("Fiber ruling uplift", f"=D{HN['Copper']}", f"=D{HN['Fiber']}-D{HN['Copper']}", "up"), ("Fiber net", "=0", f"=D{HN['Fiber']}", "value"),
      ("Stress penalty", f"=D{HN['Stress']}", f"=D{HN['Copper']}-D{HN['Stress']}", "down"), ("Stress net", "=0", f"=D{HN['Stress']}", "value")]
for lbl, base, val, color in WF:
    W(lan, rr, 5, lbl, kind="t"); W(lan, rr, 6, base, fmt=FMT_MONEY2); W(lan, rr, 7, val, fmt=FMT_MONEY2); W(lan, rr, 8, color, kind="t"); rr += 1
defname("waterfall_block", lan, f"$E${wf0}:$G${rr - 1}")
WF_ROWS = (wf0, rr - 1)
rr += 1
W(lan, rr, 1, "Tariff stack by scenario (for the stacked chart)", kind="s"); rr += 1
for i, h in enumerate(["Scenario", "MFN", "Section 301 list", "Section 301 forced labor", "Pending second action", "Total"], 1): W(lan, rr, i, h, kind="h")
rr += 1
ts0 = rr
for j, dn in enumerate(["Copper", "Fiber", "Stress"]):
    W(lan, rr, 1, dn, kind="t"); 
    for k in range(4): W(lan, rr, 2 + k, f"={'CDEF'[k]}{5 + j}", fmt=FMT_PCT)
    W(lan, rr, 6, f"=G{5 + j}", fmt=FMT_PCT, bold=True); rr += 1
defname("tariff_stack_block", lan, f"$A${ts0}:$F${rr - 1}")
TS_ROWS = (ts0, rr - 1)
lan.freeze_panes = "A5"
widths(lan, {"A": 70, "B": 24, "C": 12, "D": 14, "E": 20, "F": 16, "G": 12, "H": 60})

# ----------------------------------------------------------------------------- FX
fx = addsheet("FX")
W(fx, 1, 1, "FX", bold=True, size=14, kind="t")
W(fx, 2, 1, "The only verified cost data is in KRW (August deal). FX moves landed COGS % through the hero FOB share. A USD factory quote removes this effect.", kind="note")
W(fx, 4, 1, "Active FX rate, KRW per USD (input on Assumptions)", kind="t", bold=True); W(fx, 4, 2, "=FXRate", fmt="#,##0")
for i, h in enumerate(["Item", "Active", "Scenario 1,300", "Scenario 1,430", "Scenario 1,550", "Basis"], 1): W(fx, 6, i, h, kind="h")
W(fx, 7, 1, "FX rate, KRW per USD", kind="t"); W(fx, 7, 2, "=FXRate", fmt="#,##0")
for j, v in enumerate(mc.FX_SCEN): W(fx, 7, 3 + j, v, kind="in", fmt="#,##0")
W(fx, 7, 6, "Evidence based default 1,430; scenarios per the brief", kind="t")
FXR = {}
def fxrow(key, label, f_tmpl, fmt, basis=""):
    """f_tmpl uses {fx} for the FX cell of that column."""
    global rr
    FXR[key] = rr
    W(fx, rr, 1, label, kind="t")
    for j, col in enumerate("BCDE"):
        W(fx, rr, 2 + j, f_tmpl.format(fx=f"{col}$7"), fmt=fmt)
    W(fx, rr, 6, basis, kind="t"); rr += 1
rr = 8
W(fx, rr, 1, "August 2026 deal totals (inputs from inputs/august_deal_totals.csv; H2 Company to Dioptek, 1,400 units)", kind="s")
for c in range(2, 7): fx.cell(row=rr, column=c).fill = FILL_SEC
rr += 1
FXI = {}
for key, lbl, val, fmtv in [("units", "August deal units", mc.DEAL["units"], FMT_INT), ("cost", "August deal cost, KRW", mc.DEAL["cost_krw"], FMT_KRW), ("rev", "August deal revenue ex VAT, KRW", mc.DEAL["revenue_krw"], FMT_KRW), ("gp", "August deal gross profit, KRW", mc.DEAL["gross_profit_krw"], FMT_KRW)]:
    FXI[key] = rr
    W(fx, rr, 1, lbl, kind="t"); W(fx, rr, 2, val, kind="in", fmt=fmtv); W(fx, rr, 6, "Evidence based: " + mc.DEAL["source"] + " (inputs/august_deal_totals.csv)", kind="t"); rr += 1
W(fx, rr, 1, "August deal gross margin (gross profit divided by revenue)", kind="t"); W(fx, rr, 2, f"=B{FXI['gp']}/B{FXI['rev']}", fmt=FMT_PCT); W(fx, rr, 6, "20.3% in the brief", kind="t"); FXI["margin"] = rr; rr += 1
defname("august_deal_units", fx, f"$B${FXI['units']}"); defname("august_deal_cost_krw", fx, f"$B${FXI['cost']}"); defname("august_deal_revenue_krw", fx, f"$B${FXI['rev']}"); defname("august_deal_gross_profit_krw", fx, f"$B${FXI['gp']}"); defname("august_deal_gross_margin", fx, f"$B${FXI['margin']}")
rr += 1
W(fx, rr, 1, "Conversions at the active FX rate and the three FX scenarios", kind="s")
for c in range(2, 7): fx.cell(row=rr, column=c).fill = FILL_SEC
rr += 1
fxrow("hero_fob", "Hero FOB USD (8K 10M, KRW 11,470 times the FOB factor)", f"='SKU Unit Economics'!$D${HR}/{{fx}}*'SKU Unit Economics'!{C('fob_factor')}", FMT_MONEY2, "Evidence based KRW cost")
fxrow("fx_delta", "Landed COGS % delta vs the 1,430 cost base (hero FOB share times (1,430 / FX minus 1))", f"=HeroFobShare*({A('fx_base')}/{{fx}}-1)", FMT_PCT2, "Model formula")
fxrow("capital_krw", "Opening capital $39,000 in KRW", f"={A('opening_capital')}*{{fx}}", FMT_KRW, "DECISIONS.md")
fxrow("deal_cost_usd", "August deal cost in USD", f"=$B${FXI['cost']}/{{fx}}", FMT_MONEY, "August deal (input above)")
fxrow("deal_rev_usd", "August deal revenue ex VAT in USD", f"=$B${FXI['rev']}/{{fx}}", FMT_MONEY, "August deal (input above)")
fxrow("deal_gp_usd", "August deal gross profit in USD", f"=$B${FXI['gp']}/{{fx}}", FMT_MONEY, "August deal (input above; about $3,500 at 1,430)")
fxrow("hero_cost_krw", "Hero H2 cost KRW 11,470 in USD (Korea landed)", f"='SKU Unit Economics'!$D${HR}/{{fx}}", FMT_MONEY2, "August deal")
W(fx, rr, 1, "Profit repatriation examples (active scenario)", kind="s"); rr += 1
FX_REPAT_ROW = rr  # filled after the annual tab exists (Year 1 net, Year 5 net)
rr += 3
W(fx, rr, 1, "Korean overseas financial account report threshold: KRW 500M month end aggregate (flag for the tax accountant). Overseas direct investment report before capital moves (any stake 10% or more, any amount).", kind="note")
fx.freeze_panes = "B7"
widths(fx, {"A": 78, "B": 16, "C": 16, "D": 16, "E": 16, "F": 40})

# ============================================================================= ENGINES
# One generator writes the monthly P&L, the annual P&L, the cash flow and the gates. The main tabs use the
# Active column (F); the Scenario Summary tab uses the same generator three times with columns C, D, E.

def refto(ws_from, ws_to, col, row, absolute=False):
    """Reference from ws_from to a cell on ws_to; sheet qualified only when they differ (so link colors are right)."""
    a = f"${L(col)}${row}" if absolute else f"{L(col)}{row}"
    return a if ws_from is ws_to else f"{qs(ws_to.title)}!{a}"

Y_COLS = {1: "C", 2: "D", 3: "E", 4: "F", 5: "G"}

def write_pl(ws, r0, col, title, landed_ref):
    """Monthly P&L engine. Returns dict with 'ws' and row numbers."""
    R = {"ws": ws}
    W(ws, r0, 1, title, kind="s")
    for c in range(2, 30): ws.cell(row=r0, column=c).fill = FILL_SEC
    r = r0 + 1
    W(ws, r, 1, "Line", kind="h"); W(ws, r, 2, "Unit", kind="h"); W(ws, r, 3, "Month 0 (pre launch)", kind="h")
    for m in range(1, 25): W(ws, r, mcol(m), f"Month {m}", kind="h")
    W(ws, r, 28, "Year 1", kind="h"); W(ws, r, 29, "Year 2", kind="h")
    r += 1; R["idx"] = r
    W(ws, r, 1, "Month index", kind="t"); W(ws, r, 3, 0, kind="t", fmt="0")
    for m in range(1, 25): W(ws, r, mcol(m), m, kind="t", fmt="0")
    W(ws, r, 28, "Months 1 to 12", kind="note"); W(ws, r, 29, "Months 13 to 24", kind="note")
    r += 1
    AA = lambda k: A(k, col)
    def line(key, label, unit, f_of_m, fmt, bold=False, ysum=True, yfunc=None, kind="f", m0=None):
        nonlocal r
        R[key] = r
        W(ws, r, 1, label, kind="t", bold=bold); W(ws, r, 2, unit, kind="t")
        if m0 is not None: W(ws, r, 3, m0, fmt=fmt)
        for m in range(1, 25):
            c = ML(m); p = ML(m - 1)
            W(ws, r, mcol(m), f_of_m(m, c, p), fmt=fmt, kind=kind, bold=bold)
        if ysum:
            W(ws, r, 28, f"=SUM(D{r}:O{r})", fmt=fmt, bold=bold); W(ws, r, 29, f"=SUM(P{r}:AA{r})", fmt=fmt, bold=bold)
        elif yfunc:
            W(ws, r, 28, yfunc("AB", "D", "O"), fmt=fmt, bold=bold); W(ws, r, 29, yfunc("AC", "P", "AA"), fmt=fmt, bold=bold)
        r += 1
    I = lambda c: f"{c}${R['idx']}"
    W(ws, r, 1, "Amazon US (Months 1 to 18 reproduce the August model; Year 2 driver changes start in Month 19)", kind="t", bold=True); r += 1
    line("units", "Units sold", "units", lambda m, c, p: f"={AA(('ramp', m))}", FMT_INT, bold=True)
    line("asp", "ASP", "$ per unit", lambda m, c, p: f"={AA('asp_y12')}", FMT_MONEY2, ysum=False, yfunc=lambda yc, a, b: f"=IF({yc}{R['units']}=0,0,{yc}{r + 1}/{yc}{R['units']})")
    line("rev", "Revenue", "$", lambda m, c, p: f"={c}{R['units']}*{c}{R['asp']}", FMT_MONEY, bold=True)
    line("cogs", "Landed COGS", "$", lambda m, c, p: f"={c}{R['rev']}*{c}{r + 21}", FMT_MONEY)   # landed % row written below (r+21 resolved after block; fixed by placeholder)
    cogs_row = R["cogs"]
    line("ref", "Referral fee", "$", lambda m, c, p: f"={c}{R['rev']}*{AA('referral')}", FMT_MONEY)
    line("fba", "FBA fulfillment", "$", lambda m, c, p: f"={c}{R['units']}*{c}{r + 21}", FMT_MONEY)
    fba_row = R["fba"]
    line("ret", "Returns", "$", lambda m, c, p: f"={c}{R['rev']}*{c}{r + 20}", FMT_MONEY)
    ret_row = R["ret"]
    line("sto", "Storage and inbound", "$", lambda m, c, p: f"={c}{R['rev']}*{AA('storage')}", FMT_MONEY)
    line("ads", "Advertising", "$", lambda m, c, p: f"={c}{R['rev']}*{c}{r + 17}", FMT_MONEY)
    ads_row = R["ads"]
    line("contrib", "Contribution, Amazon US (after advertising)", "$", lambda m, c, p: f"={c}{R['rev']}-{c}{R['cogs']}-{c}{R['ref']}-{c}{R['fba']}-{c}{R['ret']}-{c}{R['sto']}-{c}{R['ads']}", FMT_MONEY, bold=True)
    Y2S = AA('y2_start')
    line("fixed", "Fixed (software, insurance, misc; people from Year 3)", "$", lambda m, c, p: f"=IF({I(c)}>={Y2S},{AA('fixed_y2')},{AA('fixed_y1')})", FMT_MONEY)
    line("fee", "Professional seller fee", "$", lambda m, c, p: f"={AA('seller_fee')}", FMT_MONEY)
    line("launch", "One time launch costs", "$", lambda m, c, p: f"=IF({I(c)}=1,{AA('launch')},0)", FMT_MONEY)
    line("comp", "Compliance (Form 5472, Wyoming report, bond, insurance; spread over the months from the Year 2 start month to Month 24)", "$", lambda m, c, p: f"=IF({I(c)}>={Y2S},({AA('form5472')}+{AA('wyoming')}+{AA('bond')}+{AA('insurance')})/(24-{Y2S}+1),0)", FMT_MONEY)
    line("net", "Net, Amazon US after fixed costs (August basis)", "$", lambda m, c, p: f"={c}{R['contrib']}-{c}{R['fixed']}-{c}{R['fee']}-{c}{R['launch']}-{c}{R['comp']}", FMT_MONEY, bold=True)
    line("cum", "Cumulative net", "$", lambda m, c, p: (f"={c}{R['net']}" if m == 1 else f"={p}{r}+{c}{R['net']}"), FMT_MONEY, ysum=False, yfunc=lambda yc, a, b: f"={b}{r}")
    line("margin", "Net margin", "%", lambda m, c, p: f"=IF({c}{R['rev']}=0,0,{c}{R['net']}/{c}{R['rev']})", FMT_PCT, ysum=False, yfunc=lambda yc, a, b: f"=IF({yc}{R['rev']}=0,0,{yc}{R['net']}/{yc}{R['rev']})")
    line("tacos", "TACOS (advertising divided by revenue)", "%", lambda m, c, p: f"=IF({c}{R['rev']}=0,0,{c}{R['ads']}/{c}{R['rev']})", FMT_PCT, ysum=False, yfunc=lambda yc, a, b: f"=IF({yc}{R['rev']}=0,0,{yc}{R['ads']}/{yc}{R['rev']})")
    line("cpu_b", "Contribution per unit before advertising", "$ per unit", lambda m, c, p: f"=IF({c}{R['units']}=0,0,({c}{R['contrib']}+{c}{R['ads']})/{c}{R['units']})", FMT_MONEY2, ysum=False, yfunc=lambda yc, a, b: f"=IF({yc}{R['units']}=0,0,({yc}{R['contrib']}+{yc}{R['ads']})/{yc}{R['units']})")
    line("cpu_a", "Contribution per unit after advertising", "$ per unit", lambda m, c, p: f"=IF({c}{R['units']}=0,0,{c}{R['contrib']}/{c}{R['units']})", FMT_MONEY2, ysum=False, yfunc=lambda yc, a, b: f"=IF({yc}{R['units']}=0,0,{yc}{R['contrib']}/{yc}{R['units']})")
    line("beflag", "Cumulative breakeven flag (1 when cumulative net is at or above 0)", "flag", lambda m, c, p: f"=IF({c}{R['cum']}>=0,1,0)", "0", ysum=False)
    r += 1
    W(ws, r, 1, "Monthly drivers (from Assumptions)", kind="t", bold=True); r += 1
    line("landed_pct", "Landed COGS % of revenue (base plus duty, FX and the Year 2 change from the Year 2 start month)", "%", lambda m, c, p: f"=IF({I(c)}>={Y2S},{landed_ref}+{AA('landed_chg_y2')},{landed_ref})", FMT_PCT, ysum=False, yfunc=lambda yc, a, b: f"=IF({yc}{R['rev']}=0,0,{yc}{R['cogs']}/{yc}{R['rev']})")
    line("fba_fee", "FBA fee per unit (inflation from the Year 2 start month)", "$ per unit", lambda m, c, p: f"={AA('fba')}*IF({I(c)}>={Y2S},1+{AA('fba_infl')},1)", FMT_MONEY2, ysum=False, yfunc=lambda yc, a, b: f"=IF({yc}{R['units']}=0,0,{yc}{R['fba']}/{yc}{R['units']})")
    line("ret_pct", "Returns % (Year 2 rate from the Year 2 start month)", "%", lambda m, c, p: f"=IF({I(c)}>={Y2S},{AA('returns_y2')},{AA('returns_y1')})", FMT_PCT, ysum=False, yfunc=lambda yc, a, b: f"=IF({yc}{R['rev']}=0,0,{yc}{R['ret']}/{yc}{R['rev']})")
    line("ads_pct", "Ad spend % (phases: Months 1 to 3, 4 to 6, 7 to 12, 13 to 18, then the Year 2 start month to 24)", "%", lambda m, c, p: f"=IF({I(c)}<=3,{AA('ads_p1')},IF({I(c)}<=6,{AA('ads_p2')},IF({I(c)}<=12,{AA('ads_p3')},IF({I(c)}<{Y2S},{AA('ads_p4')},{AA('ads_p5')}))))", FMT_PCT, ysum=False, yfunc=lambda yc, a, b: f"=IF({yc}{R['rev']}=0,0,{yc}{R['ads']}/{yc}{R['rev']})")
    # fix the forward references written with r offsets
    for m in range(1, 25):
        c = ML(m)
        ws.cell(row=cogs_row, column=mcol(m), value=f"={c}{R['rev']}*{c}{R['landed_pct']}")
        ws.cell(row=fba_row, column=mcol(m), value=f"={c}{R['units']}*{c}{R['fba_fee']}")
        ws.cell(row=ret_row, column=mcol(m), value=f"={c}{R['rev']}*{c}{R['ret_pct']}")
        ws.cell(row=ads_row, column=mcol(m), value=f"={c}{R['rev']}*{c}{R['ads_pct']}")
    r += 1
    W(ws, r, 1, "Year 2 channels (Months 13 to 24) so Year 2 ties to the annual tab", kind="t", bold=True); r += 1
    line("eb_rev", "eBay and Walmart overflow revenue (% of Amazon US revenue)", "$", lambda m, c, p: f"=IF({I(c)}>=13,{c}{R['rev']}*{AA('ebay_pct_y2')},0)", FMT_MONEY, bold=True)
    line("eb_units", "eBay and Walmart units", "units", lambda m, c, p: f"=IF({c}{R['asp']}=0,0,{c}{R['eb_rev']}/{c}{R['asp']})", FMT_INT)
    line("eb_cogs", "eBay and Walmart landed COGS", "$", lambda m, c, p: f"={c}{R['eb_rev']}*{c}{R['landed_pct']}", FMT_MONEY)
    line("eb_ref", "eBay and Walmart referral", "$", lambda m, c, p: f"={c}{R['eb_rev']}*{AA('ebay_referral')}", FMT_MONEY)
    line("eb_mcf", "eBay and Walmart multi channel fulfillment", "$", lambda m, c, p: f"={c}{R['eb_units']}*{AA('ebay_mcf')}", FMT_MONEY)
    line("eb_ads", "eBay and Walmart advertising", "$", lambda m, c, p: f"={c}{R['eb_rev']}*{AA('ebay_ads')}", FMT_MONEY)
    line("eb_ret", "eBay and Walmart returns", "$", lambda m, c, p: f"={c}{R['eb_rev']}*{AA('ebay_returns')}", FMT_MONEY)
    line("eb_con", "eBay and Walmart contribution", "$", lambda m, c, p: f"={c}{R['eb_rev']}-{c}{R['eb_cogs']}-{c}{R['eb_ref']}-{c}{R['eb_mcf']}-{c}{R['eb_ads']}-{c}{R['eb_ret']}", FMT_MONEY, bold=True)
    line("b2b_rev", "B2B installer and AV integrator revenue (Year 2 total spread evenly)", "$", lambda m, c, p: f"=IF({I(c)}>=13,{AA('b2b_y2')}/12,0)", FMT_MONEY, bold=True)
    line("b2b_units", "B2B units (price 60% of retail ASP)", "units", lambda m, c, p: f"=IF({c}{R['asp']}=0,0,{c}{R['b2b_rev']}/({c}{R['asp']}*{AA('b2b_price_factor')}))", FMT_INT)
    line("b2b_cogs", "B2B landed COGS", "$", lambda m, c, p: f"={c}{R['b2b_units']}*{c}{R['asp']}*{c}{R['landed_pct']}", FMT_MONEY)
    line("b2b_3pl", "B2B 3PL pick pack ship", "$", lambda m, c, p: f"={c}{R['b2b_units']}*{AA('b2b_3pl')}", FMT_MONEY)
    line("b2b_com", "B2B sales commission", "$", lambda m, c, p: f"={c}{R['b2b_rev']}*{AA('b2b_commission')}", FMT_MONEY)
    line("b2b_mkt", "B2B marketing", "$", lambda m, c, p: f"={c}{R['b2b_rev']}*{AA('b2b_marketing')}", FMT_MONEY)
    line("b2b_ret", "B2B returns", "$", lambda m, c, p: f"={c}{R['b2b_rev']}*{AA('b2b_returns')}", FMT_MONEY)
    line("b2b_con", "B2B contribution", "$", lambda m, c, p: f"={c}{R['b2b_rev']}-{c}{R['b2b_cogs']}-{c}{R['b2b_3pl']}-{c}{R['b2b_com']}-{c}{R['b2b_mkt']}-{c}{R['b2b_ret']}", FMT_MONEY, bold=True)
    r += 1
    W(ws, r, 1, "All channels", kind="t", bold=True); r += 1
    line("tot_rev", "Revenue, all channels", "$", lambda m, c, p: f"={c}{R['rev']}+{c}{R['eb_rev']}+{c}{R['b2b_rev']}", FMT_MONEY, bold=True)
    line("tot_net", "Net, all channels", "$", lambda m, c, p: f"={c}{R['net']}+{c}{R['eb_con']}+{c}{R['b2b_con']}", FMT_MONEY, bold=True)
    line("tot_cum", "Cumulative net, all channels", "$", lambda m, c, p: (f"={c}{R['tot_net']}" if m == 1 else f"={p}{r}+{c}{R['tot_net']}"), FMT_MONEY, ysum=False, yfunc=lambda yc, a, b: f"={b}{r}")
    line("tot_margin", "Net margin, all channels", "%", lambda m, c, p: f"=IF({c}{R['tot_rev']}=0,0,{c}{R['tot_net']}/{c}{R['tot_rev']})", FMT_PCT, ysum=False, yfunc=lambda yc, a, b: f"=IF({yc}{R['tot_rev']}=0,0,{yc}{R['tot_net']}/{yc}{R['tot_rev']})")
    r += 1
    W(ws, r, 1, "Eighteen month block (Amazon US only, August basis, Months 1 to 18)", kind="s")
    for c in range(2, 30): ws.cell(row=r, column=c).fill = FILL_SEC
    r += 1
    def kv(key, label, f, fmt):
        nonlocal r
        R[key] = r
        W(ws, r, 1, label, kind="t"); W(ws, r, 3, f, fmt=fmt, kind="out"); r += 1
    kv("m18_rev", "18 month revenue", f"=SUM(D{R['rev']}:U{R['rev']})", FMT_MONEY)
    kv("m18_net", "18 month net", f"=SUM(D{R['net']}:U{R['net']})", FMT_MONEY)
    kv("m18_margin", "18 month net margin", f"=IF(C{R['m18_rev']}=0,0,C{R['m18_net']}/C{R['m18_rev']})", FMT_PCT)
    kv("m18_be", "Cumulative breakeven month within 18 months (none if never)", f'=IFERROR(MATCH(1,D{R["beflag"]}:U{R["beflag"]},0),"none")', "0")
    kv("m18_runrate", "Month 18 run rate (Month 18 revenue times 12)", f"=U{R['rev']}*12", FMT_MONEY)
    kv("m18_exit", "Exit margin (Month 18 net margin)", f"=U{R['margin']}", FMT_PCT)
    kv("m18_units", "18 month units", f"=SUM(D{R['units']}:U{R['units']})", FMT_INT)
    kv("be24", "Cumulative breakeven month within 24 months (all channels)", f'=IFERROR(MATCH(1,D{R["beflag"]}:AA{R["beflag"]},0),"none")', "0")
    kv("m24_rev", "24 month revenue, all channels", f"=SUM(D{R['tot_rev']}:AA{R['tot_rev']})", FMT_MONEY)
    kv("m24_net", "24 month net, all channels", f"=SUM(D{R['tot_net']}:AA{R['tot_net']})", FMT_MONEY)
    R["end"] = r
    return R

def write_annual(ws, r0, col, title, PL, landed_ref):
    R = {"ws": ws}
    pws = PL["ws"]
    W(ws, r0, 1, title, kind="s")
    for c in range(2, 9): ws.cell(row=r0, column=c).fill = FILL_SEC
    r = r0 + 1
    W(ws, r, 1, "Line", kind="h"); W(ws, r, 2, "Unit", kind="h")
    for y in range(1, 6): W(ws, r, 2 + y, f"Year {y}", kind="h")
    W(ws, r, 8, "Basis", kind="h")
    r += 1
    AA = lambda k: A(k, col)
    def PLy(key, y):   # year total of a monthly row
        return refto(ws, pws, 28 if y == 1 else 29, PL[key])
    def line(key, label, unit, f12, f345, fmt, bold=False, basis=""):
        nonlocal r
        R[key] = r
        W(ws, r, 1, label, kind="t", bold=bold); W(ws, r, 2, unit, kind="t")
        for y in range(1, 6):
            yc = Y_COLS[y]; pc = Y_COLS[y - 1] if y > 1 else None
            f = f12(y, yc) if y <= 2 else f345(y, yc, pc)
            W(ws, r, 2 + y, f, fmt=fmt, bold=bold)
        W(ws, r, 8, basis, kind="t")
        r += 1
    W(ws, r, 1, "Amazon US drivers", kind="t", bold=True); r += 1
    line("drv_units", "Amazon US units", "units", lambda y, yc: f"={PLy('units', y)}", lambda y, yc, pc: f"={pc}{r}*(1+{AA(f'g_y{y}')})", FMT_INT, basis="Years 1 and 2 from the monthly tab; then unit growth driver")
    line("asp", "ASP (Years 3 to 5: yearly change plus the Year 3 roadmap mix uplift)", "$ per unit", lambda y, yc: f"={PLy('asp', y)}", lambda y, yc, pc: f"={pc}{r}*(1+{AA('asp_chg')})*IF({y}=3,1+{AA('roadmap_uplift')},1)", FMT_MONEY2)
    line("landed", "Landed COGS % (active base plus duty, FX and the yearly change)", "%", lambda y, yc: f"={PLy('landed_pct', y)}", lambda y, yc, pc: f"={landed_ref}+{AA(f'landed_chg_y{y}')}", FMT_PCT)
    line("ads_pct", "Ad spend % of revenue", "%", lambda y, yc: f"={PLy('ads_pct', y)}", lambda y, yc, pc: f"={AA(f'ads_y{y}')}", FMT_PCT)
    line("ret_pct", "Returns %", "%", lambda y, yc: f"={PLy('ret_pct', y)}", lambda y, yc, pc: f"={AA('returns_y2')}", FMT_PCT)
    line("fba_fee", "FBA fee per unit (3% inflation per year)", "$ per unit", lambda y, yc: f"={PLy('fba_fee', y)}", lambda y, yc, pc: f"={AA('fba')}*(1+{AA('fba_infl')})^({y}-1)", FMT_MONEY2)
    r += 1
    CHS = [("amz", "Amazon US"), ("ebay", "eBay and Walmart overflow"), ("b2b", "B2B installer and AV integrator"), ("dtc", "DTC site"), ("ca", "Amazon Canada"), ("pro", "Pro line (CL3, plenum, conduit pull, long run fiber KVM)")]
    ROWS = ["units", "rev", "cogs", "fees", "ful", "ads", "ret", "con"]
    LBL = {"units": "Units", "rev": "Revenue", "cogs": "Landed COGS", "fees": "Channel fees (referral, payment, commission)", "ful": "Fulfillment (FBA, MCF, 3PL, shipping)", "ads": "Advertising and marketing", "ret": "Returns and storage", "con": "Contribution"}
    FMT = {"units": FMT_INT}
    def chan(ch, label):
        nonlocal r
        W(ws, r, 1, label, kind="t", bold=True); r += 1
        for i, k in enumerate(ROWS): R[f"{ch}_{k}"] = r + i     # pre assign so units can reference revenue
        pl_map = {"amz": dict(units="units", rev="rev", cogs="cogs", fees="ref", ful="fba", ads="ads", ret=("ret", "sto"), con="contrib"),
                  "ebay": dict(units="eb_units", rev="eb_rev", cogs="eb_cogs", fees="eb_ref", ful="eb_mcf", ads="eb_ads", ret="eb_ret", con="eb_con"),
                  "b2b": dict(units="b2b_units", rev="b2b_rev", cogs="b2b_cogs", fees="b2b_com", ful="b2b_3pl", ads="b2b_mkt", ret="b2b_ret", con="b2b_con")}
        def f12(y, yc, k, ch):
            if ch in pl_map:
                src = pl_map[ch][k]
                if isinstance(src, tuple): return "=" + "+".join(PLy(s, y) for s in src)
                return f"={PLy(src, y)}"
            return "=0"
        def f345(y, yc, pc, k, ch):
            u = f"{yc}{R[f'{ch}_units']}"; rv = f"{yc}{R[f'{ch}_rev']}"
            asp = f"{yc}{R['asp']}"; ld = f"{yc}{R['landed']}"
            if ch == "amz":
                d = {"units": f"={yc}{R['drv_units']}", "rev": f"={u}*{asp}", "cogs": f"={rv}*{ld}", "fees": f"={rv}*{AA('referral')}", "ful": f"={u}*{yc}{R['fba_fee']}", "ads": f"={rv}*{yc}{R['ads_pct']}", "ret": f"={rv}*({yc}{R['ret_pct']}+{AA('storage')})"}
            elif ch == "ebay":
                d = {"rev": f"={yc}{R['amz_rev']}*{AA(f'ebay_pct_y{y}')}", "units": f"=IF({asp}=0,0,{rv}/{asp})", "cogs": f"={rv}*{ld}", "fees": f"={rv}*{AA('ebay_referral')}", "ful": f"={u}*{AA('ebay_mcf')}", "ads": f"={rv}*{AA('ebay_ads')}", "ret": f"={rv}*{AA('ebay_returns')}"}
            elif ch == "b2b":
                d = {"rev": f"={AA(f'b2b_y{y}')}", "units": f"=IF({asp}=0,0,{rv}/({asp}*{AA('b2b_price_factor')}))", "cogs": f"={u}*{asp}*{ld}", "fees": f"={rv}*{AA('b2b_commission')}", "ful": f"={u}*{AA('b2b_3pl')}", "ads": f"={rv}*{AA('b2b_marketing')}", "ret": f"={rv}*{AA('b2b_returns')}"}
            elif ch == "dtc":
                d = {"rev": f"={AA(f'dtc_y{y}')}", "units": f"=IF({asp}=0,0,{rv}/{asp})", "cogs": f"={rv}*{ld}", "fees": f"={rv}*{AA('dtc_payment')}", "ful": f"={u}*{AA('dtc_ship')}", "ads": f"={rv}*{AA('dtc_ads')}", "ret": f"={rv}*{AA('dtc_returns')}"}
            elif ch == "ca":
                d = {"rev": (f"={AA(f'ca_y{y}')}" if y >= 4 else "=0"), "units": f"=IF({asp}=0,0,{rv}/({asp}*{AA('ca_asp_factor')}))", "cogs": f"={u}*{asp}*{ld}*(1+{AA('ca_duty')})", "fees": f"={rv}*{AA('ca_referral')}", "ful": f"={u}*{AA('ca_fba')}", "ads": f"={rv}*{AA('ca_ads')}", "ret": f"={rv}*{AA('ca_returns')}"}
            else:
                sh = AA("pro_b2b_share")
                d = {"rev": (f"={AA(f'pro_y{y}')}" if y >= 4 else "=0"), "units": f"=IF({AA('pro_asp')}=0,0,{rv}/{AA('pro_asp')})", "cogs": f"={rv}*{AA('pro_landed')}",
                     "fees": f"={rv}*({sh}*{AA('b2b_commission')}+(1-{sh})*{AA('referral')})", "ful": f"={u}*({sh}*{AA('b2b_3pl')}+(1-{sh})*{AA('pro_fba')})",
                     "ads": f"={rv}*({sh}*{AA('b2b_marketing')}+(1-{sh})*{AA('pro_amz_ads')})", "ret": f"={rv}*({sh}*{AA('b2b_returns')}+(1-{sh})*({AA('pro_amz_returns')}+{AA('pro_storage')}))"}
            if k == "con":
                return f"={rv}-{yc}{R[f'{ch}_cogs']}-{yc}{R[f'{ch}_fees']}-{yc}{R[f'{ch}_ful']}-{yc}{R[f'{ch}_ads']}-{yc}{R[f'{ch}_ret']}"
            return d[k]
        for k in ROWS:
            key = f"{ch}_{k}"
            if k == "con":
                line(key, LBL[k], "$", lambda y, yc, ch=ch: f"={yc}{R[f'{ch}_rev']}-{yc}{R[f'{ch}_cogs']}-{yc}{R[f'{ch}_fees']}-{yc}{R[f'{ch}_ful']}-{yc}{R[f'{ch}_ads']}-{yc}{R[f'{ch}_ret']}",
                     lambda y, yc, pc, k=k, ch=ch: f345(y, yc, pc, k, ch), FMT_MONEY, bold=True)
            else:
                line(key, LBL[k], "units" if k == "units" else "$", lambda y, yc, k=k, ch=ch: f12(y, yc, k, ch), lambda y, yc, pc, k=k, ch=ch: f345(y, yc, pc, k, ch), FMT.get(k, FMT_MONEY), bold=(k == "rev"))
    for ch, label in CHS:
        chan(ch, label)
    r += 1
    W(ws, r, 1, "Totals", kind="t", bold=True); r += 1
    def tot(key, label, k, bold=False):
        line(key, label, "$", lambda y, yc: "=" + "+".join(f"{yc}{R[f'{ch}_{k}']}" for ch, _ in CHS), lambda y, yc, pc: "=" + "+".join(f"{yc}{R[f'{ch}_{k}']}" for ch, _ in CHS), FMT_MONEY, bold=bold)
    tot("revenue", "Revenue, all channels", "rev", bold=True)
    tot("cogs_t", "Landed COGS", "cogs"); tot("fees_t", "Channel fees", "fees"); tot("ful_t", "Fulfillment", "ful"); tot("ads_t", "Advertising and marketing", "ads"); tot("ret_t", "Returns and storage", "ret")
    tot("con_t", "Contribution, all channels", "con", bold=True)
    line("fixed", "Fixed costs (software, insurance, people, tools)", "$", lambda y, yc: f"={PLy('fixed', y)}", lambda y, yc, pc: f"={AA(f'fixed_y{y}')}*12", FMT_MONEY, basis="Years 1 and 2 from the monthly tab; Year 3 adds Mark consulting retainer, VA, tools and insurance")
    line("fee", "Professional seller fee", "$", lambda y, yc: f"={PLy('fee', y)}", lambda y, yc, pc: f"={AA('seller_fee')}*12", FMT_MONEY)
    line("launch", "One time launch costs", "$", lambda y, yc: f"={PLy('launch', y)}", lambda y, yc, pc: "=0", FMT_MONEY)
    line("comp", "Compliance (Form 5472, Wyoming report, bond, insurance, trademark maintenance)", "$", lambda y, yc: f"={PLy('comp', y)}", lambda y, yc, pc: f"={AA('form5472')}+{AA('wyoming')}+{AA('bond')}+{AA('insurance')}+IF({y}=5,{AA('trademark_y5')},0)", FMT_MONEY, basis=mc.COMPLIANCE_Y1_NOTE)
    line("net", "Net", "$", lambda y, yc: f"={yc}{R['con_t']}-{yc}{R['fixed']}-{yc}{R['fee']}-{yc}{R['launch']}-{yc}{R['comp']}", lambda y, yc, pc: f"={yc}{R['con_t']}-{yc}{R['fixed']}-{yc}{R['fee']}-{yc}{R['launch']}-{yc}{R['comp']}", FMT_MONEY, bold=True)
    line("margin", "Net margin", "%", lambda y, yc: f"=IF({yc}{R['revenue']}=0,0,{yc}{R['net']}/{yc}{R['revenue']})", lambda y, yc, pc: f"=IF({yc}{R['revenue']}=0,0,{yc}{R['net']}/{yc}{R['revenue']})", FMT_PCT, bold=True)
    line("cum_net", "Cumulative net, 5 years", "$", lambda y, yc: (f"={yc}{R['net']}" if y == 1 else f"={Y_COLS[y-1]}{r}+{yc}{R['net']}"), lambda y, yc, pc: f"={pc}{r}+{yc}{R['net']}", FMT_MONEY)
    line("amz_con_pct", "Amazon US contribution margin after advertising (compression check)", "%", lambda y, yc: f"=IF({yc}{R['amz_rev']}=0,0,{yc}{R['amz_con']}/{yc}{R['amz_rev']})", lambda y, yc, pc: f"=IF({yc}{R['amz_rev']}=0,0,{yc}{R['amz_con']}/{yc}{R['amz_rev']})", FMT_PCT)
    r += 1
    W(ws, r, 1, "Check rows (must be 0): annual minus monthly tab", kind="t", bold=True); r += 1
    R["chk_rev"] = r
    W(ws, r, 1, "Revenue check, Year 1 and Year 2 (annual minus monthly)", kind="t")
    W(ws, r, 3, f"=C{R['revenue']}-{refto(ws, pws, 28, PL['tot_rev'])}", fmt=FMT_MONEY); W(ws, r, 4, f"=D{R['revenue']}-{refto(ws, pws, 29, PL['tot_rev'])}", fmt=FMT_MONEY); r += 1
    R["chk_net"] = r
    W(ws, r, 1, "Net check, Year 1 and Year 2 (annual minus monthly)", kind="t")
    W(ws, r, 3, f"=C{R['net']}-{refto(ws, pws, 28, PL['tot_net'])}", fmt=FMT_MONEY); W(ws, r, 4, f"=D{R['net']}-{refto(ws, pws, 29, PL['tot_net'])}", fmt=FMT_MONEY); r += 1
    R["end"] = r
    return R

def write_cash(ws, r0, col, title, PL, AN, landed_ref):
    """Cash flow and working capital engine: Months 0 to 24 (columns C to AA), forecast helper Months 25 to 30 (AB to AG), annual Years 3 to 5."""
    R = {"ws": ws}
    pws = PL["ws"]; aws = AN["ws"]
    AA = lambda k: A(k, col)
    P = lambda key, m: refto(ws, pws, mcol(m), PL[key])
    W(ws, r0, 1, title, kind="s")
    for c in range(2, 34): ws.cell(row=r0, column=c).fill = FILL_SEC
    r = r0 + 1
    W(ws, r, 1, "Line", kind="h"); W(ws, r, 2, "Unit", kind="h"); W(ws, r, 3, "Month 0 (pre launch)", kind="h")
    for m in range(1, 25): W(ws, r, mcol(m), f"Month {m}", kind="h")
    for m in range(25, 31): W(ws, r, mcol(m), f"Month {m} forecast (planning only)", kind="h", wrap=True)
    ws.row_dimensions[r].height = 40
    r += 1; R["idx"] = r
    W(ws, r, 1, "Month index", kind="t")
    for m in range(0, 31): W(ws, r, mcol(m), m, kind="t", fmt="0")
    r += 1
    I = lambda c: f"{c}${R['idx']}"
    def line(key, label, unit, f_of_m, fmt, months=range(1, 25), m0=None, bold=False, kind="f"):
        nonlocal r
        R[key] = r
        W(ws, r, 1, label, kind="t", bold=bold); W(ws, r, 2, unit, kind="t")
        if m0 is not None: W(ws, r, 3, m0, fmt=fmt, kind=kind, bold=bold)
        for m in months:
            c = ML(m); p = ML(m - 1) if m > 0 else None
            W(ws, r, mcol(m), f_of_m(m, c, p), fmt=fmt, kind=kind, bold=bold)
        r += 1
    W(ws, r, 1, "Inventory plan (units; one pool at the P&L landed cost)", kind="t", bold=True); r += 1
    line("units", "Units sold, all channels (Months 25 to 30: Month 24 times 1 plus Year 3 growth)", "units",
         lambda m, c, p: (f"={P('units', m)}+{P('eb_units', m)}+{P('b2b_units', m)}" if m <= 24 else f"=$AA{r}*(1+{AA('g_y3')})"), FMT_INT, months=range(1, 31), m0="=0")
    line("landed_u", "Landed cost per unit (landed % times ASP)", "$ per unit", lambda m, c, p: (f"={P('landed_pct', m)}*{P('asp', m)}" if m <= 24 else f"=$AA{r}"), FMT_MONEY2, months=range(1, 31))
    line("fob_u", "FOB per unit paid to the factory (hero FOB USD at the active FX)", "$ per unit", lambda m, c, p: "=HeroFobUsd", FMT_MONEY2, months=range(1, 31))
    line("nonfob_u", "Freight, duty, broker, inbound and QC per unit (landed minus FOB), paid on arrival", "$ per unit", lambda m, c, p: f"={c}{R['landed_u']}-{c}{R['fob_u']}", FMT_MONEY2, months=range(1, 31))
    line("target", "Target closing stock (cover months times the average of the next three months of sales)", "units", lambda m, c, p: f"={AA('cover')}*AVERAGE({ML(m + 1)}{R['units']}:{ML(m + 3)}{R['units']})", FMT_INT, months=range(1, 27))
    R["opening"] = r; R["arrivals"] = r + 1; R["closing"] = r + 2
    line("opening", "Opening stock", "units", lambda m, c, p: ("=0" if m == 1 else f"={p}{R['closing']}"), FMT_INT, months=range(1, 27))
    line("arrivals", "Arrivals (purchase order quantity landing this month)", "units", lambda m, c, p: f"=MAX(0,{c}{R['target']}-({c}{R['opening']}-{c}{R['units']}))", FMT_INT, months=range(1, 27), bold=True)
    line("closing", "Closing stock", "units", lambda m, c, p: f"={c}{R['opening']}+{c}{R['arrivals']}-{c}{R['units']}", FMT_INT, months=range(1, 27))
    line("lead", "Lead time, months (air lead time until the ocean switch month, then the ocean lead time; both on Assumptions)", "months", lambda m, c, p: f"=IF({I(c)}>={AA('ocean_from')},{AA('lead_ocean')},{AA('lead_air')})", "0", months=range(1, 27))
    line("ordm", "Order month for this arrival (deposit paid)", "month", lambda m, c, p: f"={I(c)}-{c}{R['lead']}", "0", months=range(1, 27))
    line("shipm", "Shipment month for this arrival (balance paid; goods in transit until arrival)", "month", lambda m, c, p: f"={I(c)}-1", "0", months=range(1, 27))
    line("fobv", "FOB value of arrivals", "$", lambda m, c, p: f"={c}{R['arrivals']}*{c}{R['fob_u']}", FMT_MONEY, months=range(1, 27))
    line("nonfobv", "Freight, duty and fees on arrivals", "$", lambda m, c, p: f"={c}{R['arrivals']}*{c}{R['nonfob_u']}", FMT_MONEY, months=range(1, 27))
    r += 1
    W(ws, r, 1, "Cash flow (direct method)", kind="t", bold=True); r += 1
    rng = lambda key: f"$D${R[key]}:$AC${R[key]}"
    line("dep", "PO deposits paid (30% at order)", "$", lambda m, c, p: f"={AA('deposit_pct')}*SUMPRODUCT(({rng('ordm')}={I(c)})*{rng('fobv')})", FMT_MONEY, months=range(0, 25))
    line("balp", "PO balances paid (70% at shipment)", "$", lambda m, c, p: f"=(1-{AA('deposit_pct')})*SUMPRODUCT(({rng('shipm')}={I(c)})*{rng('fobv')})", FMT_MONEY, months=range(0, 25))
    line("nfp", "Freight, duty, broker, inbound and QC paid on arrival", "$", lambda m, c, p: (f"={c}{R['nonfobv']}" if m >= 1 else "=0"), FMT_MONEY, months=range(0, 25))
    line("inv_cash", "Inventory cash out", "$", lambda m, c, p: f"={c}{R['dep']}+{c}{R['balp']}+{c}{R['nfp']}", FMT_MONEY, months=range(0, 25), bold=True)
    line("proc_amz", "Amazon net proceeds (revenue less referral, FBA, returns, storage, advertising)", "$", lambda m, c, p: (f"={P('rev', m)}-{P('ref', m)}-{P('fba', m)}-{P('ret', m)}-{P('sto', m)}-{P('ads', m)}" if m >= 1 else "=0"), FMT_MONEY, months=range(0, 25))
    line("proc_oth", "Other channel net proceeds (eBay and Walmart, B2B)", "$", lambda m, c, p: (f"={P('eb_rev', m)}-{P('eb_ref', m)}-{P('eb_mcf', m)}-{P('eb_ads', m)}-{P('eb_ret', m)}+{P('b2b_rev', m)}-{P('b2b_3pl', m)}-{P('b2b_com', m)}-{P('b2b_mkt', m)}-{P('b2b_ret', m)}" if m >= 1 else "=0"), FMT_MONEY, months=range(0, 25))
    line("recv", "Receivable at month end (payout lag days / 30 of the month's net proceeds)", "$", lambda m, c, p: f"=({c}{R['proc_amz']}+{c}{R['proc_oth']})*{AA('payout_lag')}/30", FMT_MONEY, months=range(0, 25))
    line("cash_in", "Cash in (proceeds received)", "$", lambda m, c, p: (f"={c}{R['proc_amz']}+{c}{R['proc_oth']}-{c}{R['recv']}+{p}{R['recv']}" if m >= 1 else "=0"), FMT_MONEY, months=range(0, 25), bold=True)
    line("opex", "Fixed costs, seller fee, launch costs, compliance", "$", lambda m, c, p: (f"={P('fixed', m)}+{P('fee', m)}+{P('launch', m)}+{P('comp', m)}" if m >= 1 else "=0"), FMT_MONEY, months=range(0, 25))
    line("cash_out", "Cash out", "$", lambda m, c, p: f"={c}{R['inv_cash']}+{c}{R['opex']}", FMT_MONEY, months=range(0, 25), bold=True)
    line("ncf", "Net cash flow", "$", lambda m, c, p: f"={c}{R['cash_in']}-{c}{R['cash_out']}", FMT_MONEY, months=range(0, 25), bold=True)
    line("bal", "Cash balance before new funding (opening capital plus cumulative net cash flow)", "$", lambda m, c, p: (f"={AA('opening_capital')}+{c}{R['ncf']}" if m == 0 else f"={p}{r}+{c}{R['ncf']}"), FMT_MONEY, months=range(0, 25), bold=True)
    line("negflag", "Negative balance flag", "flag", lambda m, c, p: f"=IF({c}{R['bal']}<0,1,0)", "0", months=range(0, 25))
    r += 1
    W(ws, r, 1, "Working capital balances", kind="t", bold=True); r += 1
    line("onhand", "Inventory on hand at landed cost", "$", lambda m, c, p: (f"={c}{R['closing']}*{c}{R['landed_u']}" if m >= 1 else "=0"), FMT_MONEY, months=range(0, 25))
    line("intransit", "Deposits on open orders plus inventory in transit (funded by the business)", "$",
         lambda m, c, p: f"=SUMPRODUCT({rng('fobv')}*(({rng('ordm')}<={I(c)})*({I(c)}<{rng('shipm')})*{AA('deposit_pct')}+({rng('shipm')}<={I(c)})*({I(c)}<$D${R['idx']}:$AC${R['idx']})))", FMT_MONEY, months=range(0, 25))
    line("net_income", "Net income, all channels (monthly tab)", "$", lambda m, c, p: (f"={P('tot_net', m)}" if m >= 1 else "=0"), FMT_MONEY, months=range(0, 25))
    line("opcf_gate", "Operating cash flow for the gates (net income less the change in receivable; inventory for growth is funded separately)", "$", lambda m, c, p: (f"={c}{R['net_income']}-({c}{R['recv']}-{p}{R['recv']})" if m >= 1 else "=0"), FMT_MONEY, months=range(0, 25))
    r += 1
    W(ws, r, 1, "Peak cash need, runway and funding gap", kind="s")
    for c in range(2, 34): ws.cell(row=r, column=c).fill = FILL_SEC
    r += 1
    def kv(key, label, f, fmt, kind="out"):
        nonlocal r
        R[key] = r
        W(ws, r, 1, label, kind="t"); W(ws, r, 3, f, fmt=fmt, kind=kind); r += 1
    BAL = f"$C${R['bal']}:$AA${R['bal']}"; IDX = f"$C${R['idx']}:$AA${R['idx']}"
    kv("min_monthly", "Minimum monthly cash balance before funding, Months 0 to 24", f"=MIN({BAL})", FMT_MONEY)
    kv("min_month_m", "Month of the monthly minimum", f"=INDEX({IDX},1,MATCH(C{r - 1},{BAL},0))", "0")
    # annual block written below; reserve rows for it first
    ann0 = r + 12
    kv("min_all", "Minimum cash balance before funding, Months 0 to 24 and the Year 3 to 5 year ends", f"=MIN(C{R['min_monthly']},$C${ann0 + 7}:$E${ann0 + 7})", FMT_MONEY)
    kv("min_label", "Peak cash need month (label)", f'=IF(C{R["min_monthly"]}<=MIN($C${ann0 + 7}:$E${ann0 + 7}),"Month "&C{R["min_month_m"]},CHOOSE(MATCH(C{R["min_all"]},$C${ann0 + 7}:$E${ann0 + 7},0),"Year 3","Year 4","Year 5"))', None)
    kv("min_month_num", "Peak cash need month (number; 36, 48 or 60 for a year end)", f"=IF(C{R['min_monthly']}<=MIN($C${ann0 + 7}:$E${ann0 + 7}),C{R['min_month_m']},24+12*MATCH(C{R['min_all']},$C${ann0 + 7}:$E${ann0 + 7},0))", "0")
    kv("peak_need", "Peak cash need (the negative minimum balance, 0 if never negative)", f"=MAX(0,-C{R['min_all']})", FMT_MONEY)
    kv("gap", "Funding gap (1.2 times the peak cash need)", f"=C{R['peak_need']}*{AA('cushion')}", FMT_MONEY)
    kv("need_by", "Funding needed by month (the month before the balance first goes negative; none if never)", f'=IFERROR(MAX(0,INDEX({IDX},1,MATCH(1,$C${R["negflag"]}:$AA${R["negflag"]},0))-1),"none")', "0")
    kv("burn", "Average monthly net cash flow, Months 4 to 6", f"=AVERAGE($G${R['ncf']}:$I${R['ncf']})", FMT_MONEY, kind="f")
    kv("runway_unfunded", "Runway at the Day 90 gate before funding, months (cash at end of Month 3 divided by the average burn of Months 4 to 6; funded if no burn)", f'=IF(C{R["burn"]}<0,MAX(0,$F${R["bal"]})/-C{R["burn"]},"funded")', FMT_NUM1)
    R["funding"] = r; R["bal_funded"] = r + 1
    # funding rows (monthly) placed here
    W(ws, r, 1, "Funding drawn (line of credit or investor amount, in the month it is needed)", kind="t"); W(ws, r, 2, "$", kind="t")
    for m in range(0, 25):
        c = ML(m)
        W(ws, r, mcol(m), f"=IF({I(c)}=$C${R['need_by']},$C${R['gap']},0)", fmt=FMT_MONEY)
    r += 1
    W(ws, r, 1, "Cash balance after funding", kind="t", bold=True); W(ws, r, 2, "$", kind="t")
    for m in range(0, 25):
        c = ML(m); p = ML(m - 1) if m > 0 else None
        W(ws, r, mcol(m), (f"={c}{R['bal']}+{c}{R['funding']}" if m == 0 else f"={p}{r}+{c}{R['ncf']}+{c}{R['funding']}"), fmt=FMT_MONEY, bold=True)
    r += 1
    kv("runway_funded", "Runway at the Day 90 gate after funding, months (used by the gate)", f'=IF(C{R["burn"]}<0,MAX(0,$F${R["bal_funded"]})/-C{R["burn"]},"funded")', FMT_NUM1)
    assert r <= ann0, (r, ann0)
    r = ann0
    W(ws, r, 1, "Years 3 to 5 (annual): inventory days of landed COGS, receivable days of revenue", kind="t", bold=True)
    W(ws, r, 3, "Year 3", kind="h"); W(ws, r, 4, "Year 4", kind="h"); W(ws, r, 5, "Year 5", kind="h"); r += 1
    ANr = lambda key, y: refto(ws, aws, 2 + y, AN[key])
    def arow(key, label, f_of_y, fmt, bold=False):
        nonlocal r
        R[key] = r
        W(ws, r, 1, label, kind="t")
        for y in (3, 4, 5):
            W(ws, r, y, f_of_y(y, L(y), L(y - 1)), fmt=fmt, bold=bold)
        r += 1
    arow("a_net", "Net income", lambda y, c, p: f"={ANr('net', y)}", FMT_MONEY)
    arow("a_inv", "Inventory (on hand plus in transit) at year end", lambda y, c, p: f"={ANr('cogs_t', y)}*{AA('inv_days')}/365", FMT_MONEY)
    arow("a_recv", "Receivable at year end", lambda y, c, p: f"={ANr('revenue', y)}*{AA('recv_days')}/365", FMT_MONEY)
    arow("a_dinv", "Increase in inventory (Year 3 vs Month 24 on hand plus in transit)", lambda y, c, p: (f"={c}{R['a_inv']}-($AA${R['onhand']}+$AA${R['intransit']})" if y == 3 else f"={c}{R['a_inv']}-{p}{R['a_inv']}"), FMT_MONEY)
    arow("a_drecv", "Increase in receivable", lambda y, c, p: (f"={c}{R['a_recv']}-$AA${R['recv']}" if y == 3 else f"={c}{R['a_recv']}-{p}{R['a_recv']}"), FMT_MONEY)
    arow("a_ncf", "Net cash flow", lambda y, c, p: f"={c}{R['a_net']}-{c}{R['a_dinv']}-{c}{R['a_drecv']}", FMT_MONEY)
    assert r == ann0 + 7, (r, ann0)
    arow("a_bal", "Cash balance at year end before new funding", lambda y, c, p: (f"=$AA${R['bal']}+{c}{R['a_ncf']}" if y == 3 else f"={p}{R['a_bal']}+{c}{R['a_ncf']}"), FMT_MONEY, bold=True)
    arow("a_bal_funded", "Cash balance at year end after funding", lambda y, c, p: (f"=$AA${R['bal_funded']}+{c}{R['a_ncf']}" if y == 3 else f"={p}{R['a_bal_funded']}+{c}{R['a_ncf']}"), FMT_MONEY, bold=True)
    r += 1
    W(ws, r, 1, "Funding plan: retained earnings first (the cash balance already includes them), then a line of credit or investor amount equal to the funding gap, in place by the month shown.", kind="note"); r += 1
    W(ws, r, 1, "Low case note: the order rule (closing stock = cover months times the next three months of sales) keeps buying through the lost quarter to rebuild cover for the recovery month, so the Low peak need at the end of Year 2 assumes purchases continue during the suppression. An operator would pause orders until the listing is restored; the modeled figure is conservative, not wrong.", kind="note"); r += 1
    R["end"] = r
    return R

def write_gates(ws, r0, col, title, PL, CASH, compact=False):
    """Day 90 and Month 12 gate tables with PASS, REVIEW or KILL."""
    R = {"ws": ws}
    pws = PL["ws"]; cws = CASH["ws"]
    AA = lambda k: A(k, col)
    P = lambda key, m: refto(ws, pws, mcol(m), PL[key], absolute=True)
    Cc = lambda key: refto(ws, cws, 3, CASH[key], absolute=True)
    W(ws, r0, 1, title, kind="s")
    for c in range(2, 9): ws.cell(row=r0, column=c).fill = FILL_SEC
    r = r0 + 1
    for i, h in enumerate(["Metric", "Model value", "Pass threshold", "Kill trigger", "Pass?", "Kill?", "Note"], 1): W(ws, r, i, h, kind="h")
    r += 1
    def grow(key, label, value, pass_txt, kill_txt, pass_f, kill_f, fmt, note=""):
        nonlocal r
        R[key] = r
        W(ws, r, 1, label, kind="t"); W(ws, r, 2, value, fmt=fmt, kind="f" if isinstance(value, str) and value.startswith("=") else "t")
        W(ws, r, 3, pass_txt, kind="f" if pass_txt.startswith("=") else "t"); W(ws, r, 4, kill_txt, kind="f" if kill_txt.startswith("=") else "t")
        W(ws, r, 5, pass_f, kind="f"); W(ws, r, 6, kill_f, kind="f"); W(ws, r, 7, note, kind="t"); r += 1
    G = lambda k: AA("g_" + k)      # gate threshold cells on Assumptions (one block, same in every scenario)
    TP = lambda k: f'TEXT({G(k)},"0%")'; TM = lambda k: f'TEXT({G(k)},"$#,##0;($#,##0)")'; TN = lambda k: f'TEXT({G(k)},"0")'
    W(ws, r, 1, "Day 90 gate (end of Month 3). Thresholds are read from the Gate thresholds block on Assumptions.", kind="t", bold=True); r += 1
    grow("d90_conv", "Unit session conversion on at least 3 of the 5 launch SKUs", "[PLACEHOLDER: observed at Day 90]", f'="at or above "&{TP("d90_conv_pass")}', f'="hero below "&{TP("d90_conv_kill")}', '="observed"', '="observed"', None, "Not modeled; read from Amazon Business Reports at the gate")
    grow("d90_tacos", "TACOS in Month 3 (and falling month over month)", f"={P('tacos', 3)}", f'="at or below "&{TP("d90_tacos_pass")}', f'="above "&{TP("d90_tacos_kill")}', f"=B{r}<={G('d90_tacos_pass')}", f"=B{r}>{G('d90_tacos_kill')}", FMT_PCT, "Model ad spend phases are flat inside a phase; the falling test is observed")
    grow("d90_ret", "Return rate in Month 3", f"={P('ret_pct', 3)}", f'="at or below "&{TP("d90_ret_pass")}', f'="above "&{TP("d90_ret_kill")}', f"=B{r}<={G('d90_ret_pass')}", f"=B{r}>{G('d90_ret_kill')}", FMT_PCT)
    grow("d90_cpu", "Contribution per unit before advertising (hero basis) on at least 3 of 5 SKUs", f"={P('cpu_b', 3)}", f'="at or above "&{TM("d90_cpu_pass")}', f'="hero below "&{TM("d90_cpu_kill")}', f"=B{r}>={G('d90_cpu_pass')}", f"=B{r}<{G('d90_cpu_kill')}", FMT_MONEY2, "Medium hero $12.14 (ASP $42)")
    grow("d90_runway", "Cash runway after funding the next purchase order, months", f"={Cc('runway_funded')}", f'="at or above "&{TN("d90_run_pass")}&" months"', f'="below "&{TN("d90_run_kill")}&" months"', f'=IF(ISNUMBER(B{r}),B{r}>={G("d90_run_pass")},TRUE)', f'=IF(ISNUMBER(B{r}),B{r}<{G("d90_run_kill")},FALSE)', FMT_NUM1,
         "funded means no burn in Months 4 to 6. Runway = cash after the funding draw at the end of Month 3 divided by the average burn of Months 4 to 6, so it depends on the size of the draw (1.2 times the peak need): a more profitable case with a smaller draw can show a shorter runway than a weaker case. Size the line at the larger of the funding gap and four months of burn if the test is marginal.")
    grow("d90_supp", "No suppressed listing and no IP complaint open more than 14 days", "[PLACEHOLDER: observed at Day 90]", "none open", f'="any open more than "&{TN("d90_ip_days")}&" days"', '="observed"', '="observed"', None, "Observed from Account Health")
    R["d90_result"] = r
    W(ws, r, 1, "Day 90 result", kind="t", bold=True)
    W(ws, r, 2, f'=IF(OR(F{R["d90_tacos"]},F{R["d90_ret"]},F{R["d90_cpu"]},F{R["d90_runway"]}),"KILL",IF(AND(E{R["d90_tacos"]},E{R["d90_ret"]},E{R["d90_cpu"]},E{R["d90_runway"]}),"PASS","REVIEW"))', kind="out")
    W(ws, r, 7, "KILL if any kill trigger fires; PASS if every modeled pass threshold is met; REVIEW otherwise (fix and retest within 30 days)", kind="t"); r += 2
    W(ws, r, 1, "Month 12 gate", kind="t", bold=True); r += 1
    grow("m12_rev", "Monthly revenue in Month 12", f"={P('rev', 12)}", f'="at or above "&{TM("m12_rev_pass")}', f'="below "&{TM("m12_rev_kill")}', f"=B{r}>={G('m12_rev_pass')}", f"=B{r}<{G('m12_rev_kill')}", FMT_MONEY, "Medium Month 12 is $86,100; Low is $46,740")
    grow("m12_cum", "Cumulative net at Month 12 (Amazon US, August basis)", f"={P('cum', 12)}", f'="at or above "&{TM("m12_cum_pass")}', f'="below "&{TM("m12_cum_kill")}', f"=B{r}>={G('m12_cum_pass')}", f"=B{r}<{G('m12_cum_kill')}", FMT_MONEY)
    grow("m12_tacos", "TACOS in Month 12", f"={P('tacos', 12)}", f'="at or below "&{TP("m12_tacos_pass")}', f'="above "&{TP("m12_tacos_kill")}', f"=B{r}<={G('m12_tacos_pass')}", f"=B{r}>{G('m12_tacos_kill')}", FMT_PCT)
    grow("m12_ret", "Return rate in Month 12", f"={P('ret_pct', 12)}", f'="at or below "&{TP("m12_ret_pass")}', f'="above "&{TP("m12_ret_kill")}', f"=B{r}<={G('m12_ret_pass')}", f"=B{r}>{G('m12_ret_kill')}", FMT_PCT)
    opc = lambda m: refto(ws, cws, mcol(m), CASH["opcf_gate"], absolute=True)
    grow("m12_opcf", "Operating cash flow positive two consecutive months (shown: the lower of Months 11 and 12)", f"=MIN({opc(11)},{opc(12)})", "positive in Months 11 and 12", f'="negative in Months 10, 11 and 12 with a funding gap above "&{TM("m12_gap_kill")}',
         f"=B{r}>0", f"=AND({opc(10)}<0,{opc(11)}<0,{opc(12)}<0,{Cc('gap')}>{G('m12_gap_kill')})", FMT_MONEY, f"Operating cash flow = net income less the change in receivable. It excludes growth inventory, which the line of credit funds. Full net cash flow including inventory is on the cash tab row {CASH['ncf']}")
    R["m12_result"] = r
    W(ws, r, 1, "Month 12 result", kind="t", bold=True)
    W(ws, r, 2, f'=IF(OR(F{R["m12_rev"]},F{R["m12_cum"]},F{R["m12_tacos"]},F{R["m12_ret"]},F{R["m12_opcf"]}),"KILL",IF(AND(E{R["m12_rev"]},E{R["m12_cum"]},E{R["m12_tacos"]},E{R["m12_ret"]},E{R["m12_opcf"]}),"PASS","REVIEW"))', kind="out")
    W(ws, r, 7, "Kill means: no further capital, sell down inventory, delist losers, move to Plan B", kind="t"); r += 1
    R["end"] = r + 1
    return R

# ============================================================================= MAIN TABS (active scenario)
pl = addsheet("Monthly PL Y1 Y2")
W(pl, 1, 1, "Monthly P&L, Years 1 and 2 (Months 1 to 24)", bold=True, size=14, kind="t")
W(pl, 2, 1, '="Active scenario: "&ScenarioName&". Duty: "&DutyScenarioName&". FX "&TEXT(FXRate,"#,##0")&" KRW per USD. Amazon US rows reproduce the August model through Month 18; Year 2 driver changes start in Month 19; eBay, Walmart and B2B rows start in Month 13."', kind="f")
PL = write_pl(pl, 3, "F", "Monthly P&L engine, active scenario", "LandedPctActive")
pl.freeze_panes = "D6"
widths(pl, {"A": 66, "B": 12}); 
for c in range(3, 30): pl.column_dimensions[L(c)].width = 11.5

an = addsheet("Annual PL Y1 to Y5")
W(an, 1, 1, "Annual P&L, Years 1 to 5 by channel", bold=True, size=14, kind="t")
W(an, 2, 1, '="Active scenario: "&ScenarioName&". Years 1 and 2 are sums of the monthly tab (check rows show 0). Years 3 to 5 come from the Year 3 to 5 drivers on Assumptions."', kind="f")
AN = write_annual(an, 3, "F", "Annual P&L engine, active scenario", PL, "LandedPctActive")
an.freeze_panes = "C5"
widths(an, {"A": 66, "B": 12, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 70})

cf = addsheet("Cash Flow and Working Capital")
W(cf, 1, 1, "Cash Flow and Working Capital", bold=True, size=14, kind="t")
W(cf, 2, 1, '="Active scenario: "&ScenarioName&". Inventory cover, PO deposits, air then ocean lead times, 14 day payout lag, in transit stock funded by the business. Peak cash need, funding gap and runway below."', kind="f")
CASH = write_cash(cf, 3, "F", "Cash flow engine, active scenario", PL, AN, "LandedPctActive")
cf.freeze_panes = "D6"
widths(cf, {"A": 78, "B": 10})
for c in range(3, 34): cf.column_dimensions[L(c)].width = 11.5

# ----------------------------------------------------------------------------- Funding and Breakeven
fb = addsheet("Funding and Breakeven")
W(fb, 1, 1, "Funding and Breakeven", bold=True, size=14, kind="t")
W(fb, 2, 1, '="Active scenario: "&ScenarioName&". Opening capital, breakeven, the funding plan, the Day 90 and Month 12 gates and the Plan B summary."', kind="f")
FB = {}
r = 4
W(fb, r, 1, "Breakeven and funding plan", kind="s"); 
for c in range(2, 8): fb.cell(row=r, column=c).fill = FILL_SEC
r += 1
def fbrow(key, label, f, fmt, note="", kind="out"):
    global r
    FB[key] = r
    W(fb, r, 1, label, kind="t"); W(fb, r, 2, f, fmt=fmt, kind=kind); W(fb, r, 3, note, kind="t"); r += 1
cfc = lambda key: f"{qs(cf.title)}!$C${CASH[key]}"
plc = lambda key: f"{qs(pl.title)}!$C${PL[key]}"
fbrow("opening", "Opening capital (self funded)", f"={A('opening_capital')}", FMT_MONEY, "DECISIONS.md: self funded to $39K, the August ask")
fbrow("be18", "Cumulative breakeven month, Amazon US, within 18 months", f"={plc('m18_be')}", "0", "none means no breakeven within 18 months")
fbrow("be24", "Cumulative breakeven month, all channels, within 24 months", f"={plc('be24')}", "0")
fbrow("cum12", "Cumulative net at Month 12 (retained earnings available at the Month 12 gate)", f"={qs(pl.title)}!$O${PL['cum']}", FMT_MONEY)
fbrow("cum24", "Cumulative net at Month 24, all channels", f"={qs(pl.title)}!$AA${PL['tot_cum']}", FMT_MONEY)
fbrow("peak_need", "Peak cash need (negative minimum balance before funding)", f"={cfc('peak_need')}", FMT_MONEY)
fbrow("peak_month", "Peak cash need month", f"={cfc('min_label')}", None)
fbrow("gap", "Funding gap (1.2 times the peak need)", f"={cfc('gap')}", FMT_MONEY)
fbrow("need_by", "Funding needed by month", f"={cfc('need_by')}", "0", "the month before the balance first goes negative")
fbrow("runway_unf", "Runway at the Day 90 gate before funding, months", f"={cfc('runway_unfunded')}", FMT_NUM1)
fbrow("runway_f", "Runway at the Day 90 gate after funding, months", f"={cfc('runway_funded')}", FMT_NUM1)
r += 1
W(fb, r, 1, "Funding plan (in order)", kind="t", bold=True); W(fb, r, 2, "Amount", kind="h"); W(fb, r, 3, "When", kind="h"); W(fb, r, 4, "Source and terms", kind="h"); r += 1
FB["plan0"] = r
W(fb, r, 1, "1. Opening capital from Lucy (self funded)", kind="t"); W(fb, r, 2, f"=B{FB['opening']}", fmt=FMT_MONEY); W(fb, r, 3, "Month 0, after the Korean overseas direct investment report", kind="t"); W(fb, r, 4, "Owner capital contribution to BRAND LLC", kind="t"); r += 1
W(fb, r, 1, "2. Retained earnings (cumulative net to Month 12, already inside the cash balance)", kind="t"); W(fb, r, 2, f"=B{FB['cum12']}", fmt=FMT_MONEY); W(fb, r, 3, "Months 1 to 12", kind="t"); W(fb, r, 4, "No distributions until the Month 12 gate passes", kind="t"); r += 1
W(fb, r, 1, "3. Line of credit or investor amount (the funding gap)", kind="t", bold=True); W(fb, r, 2, f"=B{FB['gap']}", fmt=FMT_MONEY, kind="out"); W(fb, r, 3, f'=IF(ISNUMBER(B{FB["need_by"]}),"In place by Month "&B{FB["need_by"]},"Not needed")', kind="f"); W(fb, r, 4, "[PLACEHOLDER: lender or investor and terms]. Sized at 1.2 times the peak cash need so one late payout does not break the plan.", kind="t"); r += 1
W(fb, r, 1, "4. Repayment capacity: cash balance after funding at Month 24", kind="t"); W(fb, r, 2, f"={qs(cf.title)}!$AA${CASH['bal_funded']}", fmt=FMT_MONEY); W(fb, r, 3, "Month 24", kind="t"); W(fb, r, 4, "Cash after funding, before any repayment or owner draw", kind="t"); r += 1
W(fb, r, 1, "5. Cash balance after funding at the end of Year 5", kind="t"); W(fb, r, 2, f"={qs(cf.title)}!$E${CASH['a_bal_funded']}", fmt=FMT_MONEY); W(fb, r, 3, "Year 5", kind="t"); W(fb, r, 4, "Before repayment, tax and owner draws", kind="t"); r += 2
GATE = write_gates(fb, r, "F", "Kill criteria: Day 90 gate and Month 12 gate (active scenario)", PL, CASH)
r = GATE["end"] + 1
W(fb, r, 1, "Plan B summary (if Amazon fails): B2B installer distribution plus continued Korea wholesale. Full P&L on the Plan B tab.", kind="s")
for c in range(2, 8): fb.cell(row=r, column=c).fill = FILL_SEC
r += 1
W(fb, r, 1, "Line", kind="h"); W(fb, r, 2, "Year 1", kind="h"); W(fb, r, 3, "Year 2", kind="h"); W(fb, r, 4, "Year 3", kind="h"); r += 1
FB["planb_rev"] = r; FB["planb_net"] = r + 1
W(fb, r, 1, "Plan B revenue", kind="t"); W(fb, r + 1, 1, "Plan B net", kind="t")
for y in (1, 2, 3):
    W(fb, r, 1 + y, f"=planb_Y{y}_revenue", fmt=FMT_MONEY); W(fb, r + 1, 1 + y, f"=planb_Y{y}_net", fmt=FMT_MONEY)
r += 3
fb.freeze_panes = "A4"
widths(fb, {"A": 80, "B": 20, "C": 30, "D": 44, "E": 10, "F": 10, "G": 70})

# ----------------------------------------------------------------------------- Sensitivity
se = addsheet("Sensitivity")
W(se, 1, 1, "Sensitivity (one way): 18 month net and Year 5 net, active scenario", bold=True, size=14, kind="t")
W(se, 2, 1, '="Base: "&ScenarioName&" scenario, "&DutyScenarioName&" duty, FX "&TEXT(FXRate,"#,##0")&". Every cell is a live formula: a compact one line P&L that applies the delta to the base from the Monthly and Annual tabs. No pasted values. The rows are approximations of the full model: use them for the tornado and quote duty and FX effects from the Scenario Summary or a full recalc (figures.json)."', kind="f")
SE = {}
r = 4
W(se, r, 1, "Base values", kind="s"); r += 1
def serow(key, label, f, fmt):
    global r
    SE[key] = r
    W(se, r, 1, label, kind="t"); W(se, r, 2, f, fmt=fmt); r += 1
plr = lambda key, a, b: f"SUM({qs(pl.title)}!${a}${PL[key]}:${b}${PL[key]})"
anr = lambda key: f"{qs(an.title)}!$G${AN[key]}"
serow("rev18", "18 month revenue (Amazon US)", f"={plc('m18_rev')}", FMT_MONEY)
serow("net18", "18 month net (Amazon US, August basis)", f"={plc('m18_net')}", FMT_MONEY)
serow("pct18", "18 month percentage based costs (COGS, referral, returns, storage, advertising)", f"={plr('cogs','D','U')}+{plr('ref','D','U')}+{plr('ret','D','U')}+{plr('sto','D','U')}+{plr('ads','D','U')}", FMT_MONEY)
serow("cogs18", "18 month landed COGS", f"={plr('cogs','D','U')}", FMT_MONEY)
serow("landed18", "Landed COGS % in Months 1 to 18 (active)", "=LandedPctActive", FMT_PCT)
serow("ret1", "Returns % Year 1 (active)", f"={A('returns_y1')}", FMT_PCT)
serow("rev5", "Year 5 revenue, all channels", f"={anr('revenue')}", FMT_MONEY)
serow("net5", "Year 5 net", f"={anr('net')}", FMT_MONEY)
serow("pct5", "Year 5 percentage based costs (COGS, channel fees, advertising, returns)", f"={anr('cogs_t')}+{anr('fees_t')}+{anr('ads_t')}+{anr('ret_t')}", FMT_MONEY)
serow("cogs5", "Year 5 landed COGS, all channels", f"={anr('cogs_t')}", FMT_MONEY)
serow("landed5", "Landed COGS % in Year 5 (Amazon US)", f"={anr('landed')}", FMT_PCT)
serow("ret5", "Returns % Year 5 (active)", f"={A('returns_y2')}", FMT_PCT)
serow("fobshare", "Hero FOB share of price", "=HeroFobShare", FMT_PCT2)
serow("duty", "Active duty rate", "=DutyRateActive", FMT_PCT)
serow("fxdelta", "FX delta already in the base landed %", "=LandedFxDelta", FMT_PCT2)
serow("fxbase", "FX embedded in the KRW cost base", f"={A('fx_base')}", "#,##0")
serow("alib", "Alibaba benchmark FOB midpoint ($17 to $22) relative to the hero FOB, as a change", f"=('SKU Unit Economics'!{C('alib_lo')}+'SKU Unit Economics'!{C('alib_hi')})/2/HeroFobUsd-1", FMT_PCT)
r += 1
W(se, r, 1, "One way sensitivity table", kind="s"); r += 1
for i, h in enumerate(["Driver", "Setting", "Setting value", "Landed % change (points)", "18 month net", "Year 5 net", "18 month net vs base", "Year 5 net vs base"], 1): W(se, r, i, h, kind="h", wrap=True)
se.row_dimensions[r].height = 30
r += 1
B = lambda k: f"$B${SE[k]}"
def landed_apply(dl):   # returns (net18 formula, net5 formula) for a landed % change expression dl
    return (f"={B('net18')}-({dl})*{B('rev18')}", f"={B('net5')}-{B('cogs5')}*({dl})/{B('landed5')}")
SENS = {}
def sens_block(drv, settings):
    """settings: list of (label, value, value_fmt, net18_f, net5_f, dl_f or None)."""
    global r
    SENS[drv] = {}
    for lbl, val, vfmt, n18, n5, dl in settings:
        SENS[drv][lbl] = r
        W(se, r, 1, drv, kind="t"); W(se, r, 2, lbl, kind="t"); W(se, r, 3, val, fmt=vfmt, kind="in" if not (isinstance(val, str) and val.startswith("=")) else "f")
        W(se, r, 4, dl if dl else "=0", fmt=FMT_PCT2)
        W(se, r, 5, n18.replace("{dl}", f"D{r}"), fmt=FMT_MONEY); W(se, r, 6, n5.replace("{dl}", f"D{r}"), fmt=FMT_MONEY)
        W(se, r, 7, f"=E{r}-{B('net18')}", fmt=FMT_MONEY); W(se, r, 8, f"=F{r}-{B('net5')}", fmt=FMT_MONEY)
        r += 1
LA18, LA5 = landed_apply("{dl}")
sens_block("Duty rate", [(lbl, f"={ref}", FMT_PCT, LA18, LA5, f"=(C{{r}}-{B('duty')})*{B('fobshare')}") for lbl, ref in [("Fiber 12.5%", "duty_fiber_rate"), ("Copper 40.1%", "duty_copper_rate"), ("Stress 50.1%", "duty_stress_rate")]])
# fix row self references written as {r}
def fix_self(drv):
    for lbl, rr in SENS[drv].items():
        for c in (4, 5, 6):
            v = se.cell(row=rr, column=c).value
            if isinstance(v, str): se.cell(row=rr, column=c).value = v.replace("{r}", str(rr))
fix_self("Duty rate")
sens_block("ASP", [(lbl, val, FMT_PCT, f"={B('net18')}+C{{r}}*({B('rev18')}-{B('pct18')})", f"={B('net5')}+C{{r}}*({B('rev5')}-{B('pct5')})", None) for lbl, val in [("minus 10%", -0.10), ("base", 0.0), ("plus 10%", 0.10)]])
fix_self("ASP")
sens_block("Ad spend", [(lbl, val, FMT_PCT, f"={B('net18')}-C{{r}}*{B('rev18')}", f"={B('net5')}-C{{r}}*{B('rev5')}", None) for lbl, val in [("minus 3 points", -0.03), ("base", 0.0), ("plus 3 points", 0.03)]])
fix_self("Ad spend")
sens_block("Return rate", [(lbl, val, FMT_PCT, f"={B('net18')}-(C{{r}}-{B('ret1')})*{B('rev18')}", f"={B('net5')}-(C{{r}}-{B('ret5')})*{B('rev5')}", None) for lbl, val in [("2%", 0.02), ("3%", 0.03), ("5%", 0.05)]])
fix_self("Return rate")
sens_block("FX (KRW per USD)", [(lbl, val, "#,##0", LA18, LA5, f"={B('fobshare')}*({B('fxbase')}/C{{r}}-1)-{B('fxdelta')}") for lbl, val in [("1,300", 1300), ("1,430", 1430), ("1,550", 1550)]])
fix_self("FX (KRW per USD)")
sens_block("Factory cost (FOB)", [(lbl, val, FMT_PCT, LA18, LA5, f"={B('fobshare')}*C{{r}}*(1+{B('duty')})") for lbl, val in [("minus 10%", -0.10), ("base", 0.0), ("plus 25%", 0.25), ("Alibaba benchmark", f"={B('alib')}")]])
fix_self("Factory cost (FOB)")
W(se, r, 1, "Duty, FX and factory cost move landed COGS % through the hero FOB share of price (duty on FOB included for the factory cost rows). ASP moves revenue and every percentage cost; FBA fees and fixed costs stay. The B2B, DTC, Canada and pro line COGS are scaled with the same landed % change (approximation noted). The FX rows follow the spec 2.4 formula literally (FX delta = hero FOB share times (1,430 divided by FX, minus 1), no duty on the FX delta), so a 10% FOB rise from FX moves landed COGS % less than the same rise from factory cost, which carries duty. Simplification recorded in the Decision Log.", kind="note"); r += 2
W(se, r, 1, "Tornado data (Year 5 net, active scenario; order by swing, largest first)", kind="s"); r += 1
for i, h in enumerate(["Driver", "Low label", "Low value", "High label", "High value", "Base", "Swing", "Low minus base", "High minus base"], 1): W(se, r, i, h, kind="h")
r += 1
TOR = {}
tor0 = r
TORNADO = [("Factory cost (FOB)", "Alibaba benchmark", "minus 10%"), ("Duty rate", "Stress 50.1%", "Fiber 12.5%"), ("ASP", "minus 10%", "plus 10%"), ("Ad spend", "plus 3 points", "minus 3 points"), ("FX (KRW per USD)", "1,300", "1,550"), ("Return rate", "5%", "2%")]
for drv, lo, hi in TORNADO:
    TOR[drv] = r
    W(se, r, 1, drv, kind="t"); W(se, r, 2, lo, kind="t"); W(se, r, 3, f"=F{SENS[drv][lo]}", fmt=FMT_MONEY); W(se, r, 4, hi, kind="t"); W(se, r, 5, f"=F{SENS[drv][hi]}", fmt=FMT_MONEY)
    W(se, r, 6, f"={B('net5')}", fmt=FMT_MONEY); W(se, r, 7, f"=E{r}-C{r}", fmt=FMT_MONEY); W(se, r, 8, f"=C{r}-F{r}", fmt=FMT_MONEY); W(se, r, 9, f"=E{r}-F{r}", fmt=FMT_MONEY); r += 1
TOR_ROWS = (tor0, r - 1)
defname("tornado_block", se, f"$A${tor0}:$I${r - 1}")
r += 1
W(se, r, 1, "Tornado data (18 month net, active scenario)", kind="s"); r += 1
for i, h in enumerate(["Driver", "Low label", "Low value", "High label", "High value", "Base", "Swing", "Low minus base", "High minus base"], 1): W(se, r, i, h, kind="h")
r += 1
tor18_0 = r
for drv, lo, hi in TORNADO:
    W(se, r, 1, drv, kind="t"); W(se, r, 2, lo, kind="t"); W(se, r, 3, f"=E{SENS[drv][lo]}", fmt=FMT_MONEY); W(se, r, 4, hi, kind="t"); W(se, r, 5, f"=E{SENS[drv][hi]}", fmt=FMT_MONEY)
    W(se, r, 6, f"={B('net18')}", fmt=FMT_MONEY); W(se, r, 7, f"=E{r}-C{r}", fmt=FMT_MONEY); W(se, r, 8, f"=C{r}-F{r}", fmt=FMT_MONEY); W(se, r, 9, f"=E{r}-F{r}", fmt=FMT_MONEY); r += 1
defname("tornado18_block", se, f"$A${tor18_0}:$I${r - 1}")
W(se, r + 1, 1, "Static row order (largest expected swing first). Values are live; if inputs change the order may no longer be sorted.", kind="note")
se.freeze_panes = "A5"
widths(se, {"A": 54, "B": 22, "C": 16, "D": 18, "E": 16, "F": 16, "G": 18, "H": 18, "I": 16})

# ----------------------------------------------------------------------------- Plan B
pb = addsheet("Plan B")
W(pb, 1, 1, "Plan B: B2B installer distribution in the US plus continued Korea wholesale (if Amazon fails)", bold=True, size=14, kind="t")
W(pb, 2, 1, "A side business, shown honestly. B2B at the B2B economics (price 60% of retail, 3PL only, no Amazon fees). Korea wholesale at the August deal pace. Reads the Medium column of Assumptions (not the Active column) because Plan B is a fixed side case per spec 2.7; the landed % follows the duty and FX selectors.", kind="note")
PB = {}
r = 4
W(pb, r, 1, "Line", kind="h"); W(pb, r, 2, "Unit", kind="h"); W(pb, r, 3, "Year 1", kind="h"); W(pb, r, 4, "Year 2", kind="h"); W(pb, r, 5, "Year 3", kind="h"); W(pb, r, 6, "Basis", kind="h"); r += 1
def pbrow(key, label, unit, vals_or_f, fmt, basis="", kind="f", bold=False):
    global r
    PB[key] = r
    W(pb, r, 1, label, kind="t", bold=bold); W(pb, r, 2, unit, kind="t")
    for y in (1, 2, 3):
        v = vals_or_f[y - 1] if isinstance(vals_or_f, list) else vals_or_f(y, L(2 + y), L(1 + y))
        if v is None: continue
        W(pb, r, 2 + y, v, fmt=fmt, kind=kind, bold=bold)
    W(pb, r, 6, basis, kind="t"); r += 1
pbrow("b2b_rev", "B2B installer revenue", "$", mc.PLANB["b2b"], FMT_MONEY, "Management estimate (spec 2.7)", kind="in", bold=True)
pbrow("asp", "Retail ASP reference (Medium)", "$ per unit", lambda y, c, p: f"={A('asp_y12', 'D')}", FMT_MONEY2, "August model")
pbrow("landed", "Landed COGS % (Medium, active duty and FX)", "%", lambda y, c, p: "=LandedPctMedium", FMT_PCT)
pbrow("units", "B2B units (price 60% of retail)", "units", lambda y, c, p: f"={c}{PB['b2b_rev']}/({c}{PB['asp']}*{A('b2b_price_factor', 'D')})", FMT_INT)
pbrow("cogs", "Landed COGS", "$", lambda y, c, p: f"={c}{PB['units']}*{c}{PB['asp']}*{c}{PB['landed']}", FMT_MONEY)
pbrow("3pl", "3PL pick pack ship", "$", lambda y, c, p: f"={c}{PB['units']}*{A('b2b_3pl', 'D')}", FMT_MONEY)
pbrow("com", "Sales commission", "$", lambda y, c, p: f"={c}{PB['b2b_rev']}*{A('b2b_commission', 'D')}", FMT_MONEY)
pbrow("mkt", "Marketing", "$", lambda y, c, p: f"={c}{PB['b2b_rev']}*{A('b2b_marketing', 'D')}", FMT_MONEY)
pbrow("ret", "Returns", "$", lambda y, c, p: f"={c}{PB['b2b_rev']}*{A('b2b_returns', 'D')}", FMT_MONEY)
pbrow("b2b_con", "B2B contribution", "$", lambda y, c, p: f"={c}{PB['b2b_rev']}-{c}{PB['cogs']}-{c}{PB['3pl']}-{c}{PB['com']}-{c}{PB['mkt']}-{c}{PB['ret']}", FMT_MONEY, bold=True)
r += 1
pbrow("deals", "Korea wholesale deals per year (August deal pace)", "deals", [mc.PLANB["korea_deals"]] * 3, "0", "Management estimate", kind="in")
pbrow("deal_units", "Units per deal", "units", lambda y, c, p: f"=FX!$B${FXI['units']}", FMT_INT, "August deal units (FX tab input)")
pbrow("gp_unit_y1", "Gross profit per unit, Year 1 (August deal: about $2.50 per unit)", "$ per unit", [mc.PLANB["korea_gp_unit"], None, None], FMT_MONEY2, "Management estimate", kind="in")
pbrow("gp_growth", "Gross profit per unit growth per year", "%", [None, mc.PLANB["korea_growth"], mc.PLANB["korea_growth"]], FMT_PCT, "Management estimate", kind="in")
pbrow("gp_unit", "Gross profit per unit (Year 1 input, then growing by the rate above)", "$ per unit", lambda y, c, p: (f"={c}{PB['gp_unit_y1']}" if y == 1 else f"={p}{r}*(1+{c}{PB['gp_growth']})"), FMT_MONEY2, "Management estimate")
pbrow("kr_rev_u", "August deal revenue per unit in USD (FX tab deal revenue divided by the units per deal)", "$ per unit", lambda y, c, p: f"=FX!$B${FXR['deal_rev_usd']}/{c}{PB['deal_units']}", FMT_MONEY2, "August deal revenue (FX tab input) on the units per deal, at FXRate")
pbrow("kr_rev", "Korea wholesale revenue (deals times units times the August deal revenue per unit)", "$", lambda y, c, p: f"={c}{PB['deals']}*{c}{PB['deal_units']}*{c}{PB['kr_rev_u']}", FMT_MONEY, "August deal")
pbrow("kr_gp", "Korea wholesale gross profit", "$", lambda y, c, p: f"={c}{PB['deals']}*{c}{PB['deal_units']}*{c}{PB['gp_unit']}", FMT_MONEY, bold=True)
r += 1
pbrow("fixed_m", "Fixed costs (software, insurance, 3PL account) per month", "$", list(mc.PLANB["fixed_m"]), FMT_MONEY, "Management estimate: 850 per month Year 1, 1,200 after", kind="in")
pbrow("fixed", "Fixed costs per year", "$", lambda y, c, p: f"={c}{PB['fixed_m']}*12", FMT_MONEY)
pbrow("comp", "Compliance (Form 5472, Wyoming report, bond, insurance) from Year 2", "$", lambda y, c, p: ("=0" if y == 1 else f"={A('form5472', 'D')}+{A('wyoming', 'D')}+{A('bond', 'D')}+{A('insurance', 'D')}"), FMT_MONEY)
pbrow("revenue", "Plan B revenue (B2B plus Korea wholesale)", "$", lambda y, c, p: f"={c}{PB['b2b_rev']}+{c}{PB['kr_rev']}", FMT_MONEY, bold=True)
pbrow("net", "Plan B net (B2B contribution plus Korea gross profit less fixed and compliance)", "$", lambda y, c, p: f"={c}{PB['b2b_con']}+{c}{PB['kr_gp']}-{c}{PB['fixed']}-{c}{PB['comp']}", FMT_MONEY, bold=True)
pbrow("margin", "Plan B net margin", "%", lambda y, c, p: f"=IF({c}{PB['revenue']}=0,0,{c}{PB['net']}/{c}{PB['revenue']})", FMT_PCT)
pbrow("cum", "Cumulative Plan B net", "$", lambda y, c, p: (f"={c}{PB['net']}" if y == 1 else f"={p}{r}+{c}{PB['net']}"), FMT_MONEY)
for y in (1, 2, 3):
    defname(f"planb_Y{y}_revenue", pb, f"${L(2 + y)}${PB['revenue']}"); defname(f"planb_Y{y}_net", pb, f"${L(2 + y)}${PB['net']}")
    defname(f"planb_Y{y}_b2b_revenue", pb, f"${L(2 + y)}${PB['b2b_rev']}"); defname(f"planb_Y{y}_korea_gross_profit", pb, f"${L(2 + y)}${PB['kr_gp']}")
W(pb, r + 1, 1, "Plan B keeps the LLC, the bond and the insurance. It does not need Amazon fees, advertising at Amazon rates or the inventory cover of the Amazon plan.", kind="note")
pb.freeze_panes = "C5"
widths(pb, {"A": 80, "B": 12, "C": 14, "D": 14, "E": 14, "F": 44})

# ----------------------------------------------------------------------------- Levers
lv = addsheet("Levers")
W(lv, 1, 1, "Twelve levers, each mapped to the assumption it moves", bold=True, size=14, kind="t")
W(lv, 2, 1, "From and To values quote the Medium columns of the model (links where the model holds the number; yellow cells are typed targets). They are pinned to Medium on purpose, so the lever table reads the same whatever the Scenario selector shows (recorded in the Decision Log). Column I quantifies each lever from the Medium engine on the Scenario Summary tab. Owner and evidence per the build spec 2.8.", kind="note")
LEV_HDR = ["#", "Lever", "Assumption it moves", "From (Medium values)", "To (Medium values)", "Owner", "Evidence", "Model cell or driver", "Model effect on net, Medium ($)", "How the effect is computed (approximate; Scenario Summary Medium engine)"]
for i, h in enumerate(LEV_HDR, 1): W(lv, 4, i, h, kind="h")
LEVERS = [
    (1, "CBP binding ruling: fiber classification 8544.70.00", "Duty rate, landed COGS %", f"={DQ}!${'BCD'[0]}${LS['Medium']}", f"={DQ}!${'BCD'[1]}${LS['Medium']}", "Broker, Mark", "August deck lever, brief", "Landed Cost & Tariff Scenarios: Medium row, Copper vs Fiber", FMT_PCT),
    (2, "Air to ocean freight after the Day 90 gate", "Freight per unit, landed COGS %", f"='SKU Unit Economics'!{C('air')}", f"='SKU Unit Economics'!{C('ocean')}", "Lucy, 3PL", "August deck lever", "SKU Unit Economics constants: air and ocean per kg; Assumptions: landed COGS change Year 2", FMT_MONEY2),
    (3, "Direction proof connector design", "Return rate", f"={A('returns_y1', 'D')}", f"={A('returns_y2', 'D')}", "Lucy, factory", "August deck lever", "Assumptions: returns % Year 1 and Year 2 onward (Medium)", FMT_PCT),
    (4, "Keyword pruning and negative keywords", "Ad spend %", "phase rate", "minus 2 points", "Mark, VA", "August deck lever", "Assumptions: ad spend % by phase; Sensitivity: ad spend minus 3 points", None),
    (5, "Review velocity (Vine, insert card, Brand Registry follow up)", "Conversion, unit ramp", "base ramp", "plus 10% units", "Mark", "August deck lever", "Assumptions: unit ramp multiplier", None),
    (6, "Length variations under one parent ASIN", "Conversion, review pooling, ad efficiency", "separate ASINs", "one parent", "Mark", "August deck lever", "Not a numeric driver; supports levers 4 and 5", None),
    (7, "Factory volume pricing tiers", "FOB cost", "H2 base", "minus 5% at 5,000 units per quarter", "Lucy", "Management estimate", "SKU Unit Economics: H2 cost KRW; Sensitivity: factory cost minus 10%", None),
    (8, "Packaging to a smaller FBA size tier", "FBA fee per unit", f"={A('fba', 'D')}", 4.75, "Lucy, factory", "Management estimate", "Assumptions: FBA fee per unit (hero basis); SKU tab FBA fee column", FMT_MONEY2),
    (9, "3PL feeding FBA, avoid aged inventory surcharges", "Storage and inbound %", f"={A('storage', 'D')}", 0.015, "3PL, Lucy", "Management estimate", "Assumptions: storage and inbound %", FMT_PCT),
    (10, "Price ladder and coupons (list vs sale price)", "ASP", f"={A('asp_y12', 'D')}", "plus or minus 5% test", "Mark", "Management estimate", "Assumptions: ASP Years 1 and 2; Sensitivity: ASP plus or minus 10%", FMT_MONEY2),
    (11, "Off Amazon email list (Spec Match finder, trade up registration)", "DTC revenue, ad %", "=0", f"={A('dtc_y3', 'D')}", "Mark", "Management estimate", "Assumptions: DTC site revenue Year 3 (channel from Year 3)", FMT_MONEY),
    (12, "B2B installer program (AV integrators, custom installers)", "B2B channel revenue", "=0", f"={A('b2b_y2', 'D')}", "Mark, Lucy", "Management estimate", "Assumptions: B2B revenue Year 2", FMT_MONEY),
]
LEV_ROW0 = 5
for i, (n, lever, moves, frm, to, owner, ev, cell, fmt) in enumerate(LEVERS):
    rr = LEV_ROW0 + i
    W(lv, rr, 1, n, kind="t", fmt="0"); W(lv, rr, 2, lever, kind="t"); W(lv, rr, 3, moves, kind="t")
    lk = lambda v: "f" if isinstance(v, str) and v.startswith("=") else ("in" if isinstance(v, (int, float)) and v != 0 else "t")
    W(lv, rr, 4, frm, fmt=fmt, kind=lk(frm)); W(lv, rr, 5, to, fmt=fmt, kind=lk(to))
    W(lv, rr, 6, owner, kind="t"); W(lv, rr, 7, ev, kind="t"); W(lv, rr, 8, cell, kind="t")
defname("levers_block", lv, f"$A${LEV_ROW0}:$J${LEV_ROW0 + 11}")
lv.freeze_panes = "A5"
widths(lv, {"A": 4, "B": 56, "C": 36, "D": 18, "E": 30, "F": 14, "G": 22, "H": 60, "I": 20, "J": 90})

# ----------------------------------------------------------------------------- Scenario Summary (three parallel engines)
ss = addsheet("Scenario Summary")
W(ss, 1, 1, "Scenario Summary: Low, Medium and High side by side", bold=True, size=14, kind="t")
W(ss, 2, 1, "Computed live by three parallel copies of the monthly, annual, cash and gate engines (rows 45 onward), each wired to its own Assumptions column, so the table does not depend on the Scenario selector. Duty scenario and FX follow the selectors.", kind="note")
W(ss, 3, 1, "The active scenario column on the main tabs must equal the matching column here (the build script checks this after recalculation).", kind="note")
SSR = {}
r = 5
for i, h in enumerate(["Metric", "Low", "Medium", "High", "Unit"], 1): W(ss, r, i, h, kind="h")
r += 1
SS_ROWS = [
    ("m18_revenue", "18 month revenue (Amazon US)", FMT_MONEY), ("m18_net", "18 month net (Amazon US, August basis)", FMT_MONEY), ("m18_net_margin", "18 month net margin", FMT_PCT),
    ("m18_breakeven_month", "Cumulative breakeven month (18 months; none if never)", "0"), ("m18_run_rate", "Month 18 run rate", FMT_MONEY), ("m18_exit_margin", "Exit margin (Month 18)", FMT_PCT), ("m18_units", "18 month units", FMT_INT),
    ("annual_Y1_revenue", "Year 1 revenue", FMT_MONEY), ("annual_Y1_net", "Year 1 net", FMT_MONEY), ("annual_Y2_revenue", "Year 2 revenue", FMT_MONEY), ("annual_Y2_net", "Year 2 net", FMT_MONEY),
    ("annual_Y3_revenue", "Year 3 revenue", FMT_MONEY), ("annual_Y3_net", "Year 3 net", FMT_MONEY), ("annual_Y4_revenue", "Year 4 revenue", FMT_MONEY), ("annual_Y4_net", "Year 4 net", FMT_MONEY),
    ("annual_Y5_revenue", "Year 5 revenue", FMT_MONEY), ("annual_Y5_net", "Year 5 net", FMT_MONEY), ("annual_Y5_net_margin", "Year 5 net margin", FMT_PCT), ("annual_Y5_amazon_us", "Year 5 Amazon US revenue", FMT_MONEY),
    ("cash_peak_need", "Peak cash need", FMT_MONEY), ("cash_peak_need_month", "Peak cash need month (number; 36, 48 or 60 for a year end)", "0"), ("cash_peak_need_month_label", "Peak cash need month (label)", None), ("cash_min_balance", "Minimum cash balance before funding", FMT_MONEY),
    ("cash_funding_gap", "Funding gap (1.2 times peak need)", FMT_MONEY), ("cash_funding_needed_by_month", "Funding needed by month", "0"), ("cash_runway_months_day90", "Runway at Day 90 after funding, months", FMT_NUM1),
    ("gate_day90", "Day 90 gate result", None), ("gate_month12", "Month 12 gate result", None),
]
for key, label, fmt in SS_ROWS:
    SSR[key] = r
    W(ss, r, 1, label, kind="t"); W(ss, r, 5, {"0": "month", FMT_PCT: "%", FMT_MONEY: "$", FMT_INT: "units", FMT_NUM1: "months", None: ""}[fmt], kind="t")
    r += 1
ENG = {}
r = 45
for i, sname in enumerate(mc.SCEN_NAMES):
    col = "CDE"[i]
    PLs = write_pl(ss, r, col, f"{sname} scenario: monthly P&L engine (Assumptions column {col})", f"LandedPct{sname}")
    ANs = write_annual(ss, PLs["end"] + 1, col, f"{sname} scenario: annual P&L engine", PLs, f"LandedPct{sname}")
    CAs = write_cash(ss, ANs["end"] + 1, col, f"{sname} scenario: cash flow engine", PLs, ANs, f"LandedPct{sname}")
    GAs = write_gates(ss, CAs["end"] + 1, col, f"{sname} scenario: gates", PLs, CAs)
    ENG[sname] = dict(pl=PLs, an=ANs, cash=CAs, gate=GAs)
    r = GAs["end"] + 2
    c = 2 + i
    def put(key, f, fmt):
        W(ss, SSR[key], c, f, fmt=fmt, kind="out")
        defname(f"ss_{sname}_{key}", ss, f"${L(c)}${SSR[key]}")
    put("m18_revenue", f"=C{PLs['m18_rev']}", FMT_MONEY); put("m18_net", f"=C{PLs['m18_net']}", FMT_MONEY); put("m18_net_margin", f"=C{PLs['m18_margin']}", FMT_PCT)
    put("m18_breakeven_month", f"=C{PLs['m18_be']}", "0"); put("m18_run_rate", f"=C{PLs['m18_runrate']}", FMT_MONEY); put("m18_exit_margin", f"=C{PLs['m18_exit']}", FMT_PCT); put("m18_units", f"=C{PLs['m18_units']}", FMT_INT)
    for y in range(1, 6):
        put(f"annual_Y{y}_revenue", f"={Y_COLS[y]}{ANs['revenue']}", FMT_MONEY); put(f"annual_Y{y}_net", f"={Y_COLS[y]}{ANs['net']}", FMT_MONEY)
    put("annual_Y5_net_margin", f"=G{ANs['margin']}", FMT_PCT); put("annual_Y5_amazon_us", f"=G{ANs['amz_rev']}", FMT_MONEY)
    put("cash_peak_need", f"=C{CAs['peak_need']}", FMT_MONEY); put("cash_peak_need_month", f"=C{CAs['min_month_num']}", "0"); put("cash_peak_need_month_label", f"=C{CAs['min_label']}", None); put("cash_min_balance", f"=C{CAs['min_all']}", FMT_MONEY)
    put("cash_funding_gap", f"=C{CAs['gap']}", FMT_MONEY); put("cash_funding_needed_by_month", f"=C{CAs['need_by']}", "0"); put("cash_runway_months_day90", f"=C{CAs['runway_funded']}", FMT_NUM1)
    put("gate_day90", f"=B{GAs['d90_result']}", None); put("gate_month12", f"=B{GAs['m12_result']}", None)
SS_TABLE_ROWS = (SSR["m18_revenue"], SSR["gate_month12"])
defname("scenario_summary_table", ss, f"$A${SS_TABLE_ROWS[0]}:$D${SS_TABLE_ROWS[1]}")

# Levers: model effect column (columns I and J), from the Medium engine on the Scenario Summary tab so it does not follow the selector
MP = ENG["Medium"]["pl"]; MA = ENG["Medium"]["an"]; SSQ = qs(ss.title)
m18rev = f"{SSQ}!$C${MP['m18_rev']}"; m18units = f"{SSQ}!$C${MP['m18_units']}"
m18con = f"SUM({SSQ}!$D${MP['contrib']}:$U${MP['contrib']})"
m18pct = "+".join(f"SUM({SSQ}!$D${MP[k]}:$U${MP[k]})" for k in ("cogs", "ref", "ret", "sto", "ads"))
LEV_EFFECT = {
    1: (f"=({DQ}!$B${LS['Medium']}-{DQ}!$C${LS['Medium']})*{m18rev}", "18 month net: (copper landed % minus fiber landed %) times Medium 18 month revenue"),
    2: (f"=('SKU Unit Economics'!$J${HR}-'SKU Unit Economics'!$K${HR})*{m18units}", "18 month net, upper bound: (air minus ocean freight per hero unit) times 18 month units; the model phases the switch in from the ocean start month through the Year 2 landed change"),
    3: (f"=({A('returns_y1', 'D')}-{A('returns_y2', 'D')})*{m18rev}", "18 month net: the returns point saved times 18 month revenue"),
    4: (f"=0.02*{m18rev}", "18 month net: 2 points of ad spend times 18 month revenue"),
    5: (f"=0.10*{m18con}", "18 month net: 10% more units at the same contribution per unit after advertising; fixed costs unchanged"),
    6: ("not numeric", "Supports levers 4 and 5 (shared reviews and ad budget across the lengths)"),
    7: (f"=0.05*HeroFobShare*(1+DutyRateCopper)*{m18rev}", "18 month net: FOB minus 5% with the duty on that FOB saved too, through the hero FOB share of price"),
    8: (f"=({A('fba', 'D')}-E{LEV_ROW0 + 7})*{m18units}", "18 month net, upper bound: the FBA fee saving per unit times 18 month units (applies to the 7M and 10M lengths only)"),
    9: (f"=({A('storage', 'D')}-E{LEV_ROW0 + 8})*{m18rev}", "18 month net: half a point of storage and inbound times 18 month revenue"),
    10: (f"=0.05*({m18rev}-({m18pct}))", "18 month net: plus 5% ASP with every percentage cost scaling with revenue (minus 5% costs the same amount)"),
    11: (f"={SSQ}!$E${MA['dtc_con']}", "Year 3 contribution of the DTC channel (Medium)"),
    12: (f"={SSQ}!$D${MA['b2b_con']}", "Year 2 contribution of the B2B channel (Medium)"),
}
for i in range(12):
    rr = LEV_ROW0 + i
    f, how = LEV_EFFECT[i + 1]
    W(lv, rr, 9, f, fmt=FMT_MONEY, kind="f" if f.startswith("=") else "t", bold=f.startswith("=")); W(lv, rr, 10, how, kind="t")
    defname(f"lever_{i + 1}_effect", lv, f"$I${rr}")
ss.freeze_panes = "B6"
widths(ss, {"A": 70, "B": 16, "C": 16, "D": 16, "E": 12})
for c in range(6, 34): ss.column_dimensions[L(c)].width = 11

# ----------------------------------------------------------------------------- KPI Dashboard
kp = addsheet("KPI Dashboard")
W(kp, 1, 1, "KPI Dashboard", bold=True, size=14, kind="t")
W(kp, 2, 1, '="Active scenario: "&ScenarioName&". Duty: "&DutyScenarioName&". FX "&TEXT(FXRate,"#,##0")&"."', kind="f")
KP = {}
r = 4
W(kp, r, 1, "KPI", kind="h"); W(kp, r, 2, "Value", kind="h"); W(kp, r, 3, "Note", kind="h"); r += 1
def kpi(key, label, f, fmt, note=""):
    global r
    KP[key] = r
    W(kp, r, 1, label, kind="t"); W(kp, r, 2, f, fmt=fmt, kind="out"); W(kp, r, 3, note, kind="t"); r += 1
kpi("m18_rev", "18 month revenue (Amazon US)", f"={plc('m18_rev')}", FMT_MONEY); kpi("m18_net", "18 month net", f"={plc('m18_net')}", FMT_MONEY); kpi("m18_margin", "18 month net margin", f"={plc('m18_margin')}", FMT_PCT)
kpi("be", "Cumulative breakeven month (18 months)", f"={plc('m18_be')}", "0"); kpi("runrate", "Month 18 run rate", f"={plc('m18_runrate')}", FMT_MONEY); kpi("exit", "Exit margin (Month 18)", f"={plc('m18_exit')}", FMT_PCT)
kpi("y5rev", "Year 5 revenue, all channels", f"={anr('revenue')}", FMT_MONEY); kpi("y5net", "Year 5 net", f"={anr('net')}", FMT_MONEY); kpi("y5margin", "Year 5 net margin", f"={anr('margin')}", FMT_PCT)
kpi("peak", "Peak cash need", f"={cfc('peak_need')}", FMT_MONEY); kpi("peakm", "Peak cash need month", f"={cfc('min_label')}", None); kpi("gap", "Funding gap (1.2 times peak need)", f"={cfc('gap')}", FMT_MONEY); kpi("needby", "Funding needed by month", f"={cfc('need_by')}", "0")
kpi("d90", "Day 90 gate", f"={qs(fb.title)}!$B${GATE['d90_result']}", None); kpi("m12", "Month 12 gate", f"={qs(fb.title)}!$B${GATE['m12_result']}", None)
r += 1
# chart data helper block: months, revenue, net, cash balance, marker
W(kp, r, 1, "Chart data (links)", kind="s"); r += 1
CD = {"hdr": r}
W(kp, r, 1, "Month", kind="h"); W(kp, r, 2, "Revenue, all channels", kind="h"); W(kp, r, 3, "Net, all channels", kind="h"); W(kp, r, 4, "Cash balance before funding", kind="h"); W(kp, r, 5, "Minimum marker", kind="h"); W(kp, r, 6, "Minimum balance line", kind="h"); r += 1
CD["m0"] = r
for m in range(0, 25):
    W(kp, r, 1, f"Month {m}", kind="t")
    if m >= 1:
        W(kp, r, 2, f"={qs(pl.title)}!{ML(m)}{PL['tot_rev']}", fmt=FMT_MONEY); W(kp, r, 3, f"={qs(pl.title)}!{ML(m)}{PL['tot_net']}", fmt=FMT_MONEY)
    W(kp, r, 4, f"={qs(cf.title)}!{ML(m)}{CASH['bal']}", fmt=FMT_MONEY)
    W(kp, r, 5, f"=IF({m}={cfc('min_month_num')},D{r},0)", fmt=FMT_MONEY); W(kp, r, 6, f"={cfc('min_monthly')}", fmt=FMT_MONEY)
    r += 1
CD["m24"] = r - 1
r += 1
W(kp, r, 1, "Annual revenue by channel (links)", kind="h")
for y in range(1, 6): W(kp, r, 1 + y, f"Year {y}", kind="h")
r += 1
CD["ch0"] = r
for ch, label in [("amz", "Amazon US"), ("ebay", "eBay and Walmart"), ("b2b", "B2B installer and AV integrator"), ("dtc", "DTC site"), ("ca", "Amazon Canada"), ("pro", "Pro line")]:
    W(kp, r, 1, label, kind="t")
    for y in range(1, 6): W(kp, r, 1 + y, f"={qs(an.title)}!{Y_COLS[y]}{AN[f'{ch}_rev']}", fmt=FMT_MONEY)
    r += 1
CD["ch1"] = r - 1
r += 1
W(kp, r, 1, "Tornado (Year 5 net vs base)", kind="h"); W(kp, r, 2, "Low minus base", kind="h"); W(kp, r, 3, "High minus base", kind="h"); r += 1
CD["t0"] = r
for i in range(TOR_ROWS[0], TOR_ROWS[1] + 1):
    W(kp, r, 1, f"={qs(se.title)}!A{i}", kind="f"); W(kp, r, 2, f"={qs(se.title)}!H{i}", fmt=FMT_MONEY); W(kp, r, 3, f"={qs(se.title)}!I{i}", fmt=FMT_MONEY); r += 1
CD["t1"] = r - 1
r += 1
W(kp, r, 1, "Duty waterfall (hero net per unit)", kind="h"); W(kp, r, 2, "Invisible base", kind="h"); W(kp, r, 3, "Bar", kind="h"); r += 1
CD["w0"] = r
for i in range(WF_ROWS[0], WF_ROWS[1] + 1):
    W(kp, r, 1, f"={qs(lan.title)}!E{i}", kind="f"); W(kp, r, 2, f"={qs(lan.title)}!F{i}", fmt=FMT_MONEY2); W(kp, r, 3, f"={qs(lan.title)}!G{i}", fmt=FMT_MONEY2); r += 1
CD["w1"] = r - 1
widths(kp, {"A": 44, "B": 20, "C": 20, "D": 22, "E": 16, "F": 20, "G": 14})
# charts
def style_series(s, color, width=19050):
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = width
c1 = LineChart(); c1.title = "Revenue and net by month, all channels (Months 1 to 24)"; c1.style = 2; c1.height = 8; c1.width = 18
c1.y_axis.title = "USD"; c1.x_axis.title = "Model month"; c1.y_axis.number_format = '$#,##0;($#,##0);0'
for colidx, title, colr in ((2, "Revenue", "2DE2FF"), (3, "Net", "8B5CF6")):
    s = Series(Reference(kp, min_col=colidx, min_row=CD["m0"] + 1, max_row=CD["m24"]), title=title); style_series(s, colr); c1.series.append(s)
c1.set_categories(Reference(kp, min_col=1, min_row=CD["m0"] + 1, max_row=CD["m24"]))
kp.add_chart(c1, "I4")
c2 = BarChart(); c2.type = "col"; c2.grouping = "stacked"; c2.overlap = 100; c2.title = "Annual revenue by channel, Years 1 to 5"; c2.height = 8; c2.width = 18; c2.y_axis.number_format = '$#,##0;($#,##0);0'
PAL = ["2DE2FF", "8B5CF6", "3DDC97", "FFB547", "FF4FD8", "9AA5C4"]
for i, rr in enumerate(range(CD["ch0"], CD["ch1"] + 1)):
    s = Series(Reference(kp, min_col=2, max_col=6, min_row=rr, max_row=rr), title=kp.cell(row=rr, column=1).value); s.graphicalProperties.solidFill = PAL[i]; s.graphicalProperties.line.solidFill = PAL[i]; c2.series.append(s)
c2.set_categories(Reference(kp, min_col=2, max_col=6, min_row=CD["ch0"] - 1, max_row=CD["ch0"] - 1))
kp.add_chart(c2, "I22")
c3 = BarChart(); c3.type = "col"; c3.title = "Cash balance by month before funding, minimum marked (Months 0 to 24)"; c3.height = 8; c3.width = 18; c3.y_axis.number_format = '$#,##0;($#,##0);0'
sm = Series(Reference(kp, min_col=5, min_row=CD["m0"], max_row=CD["m24"]), title="Minimum (peak cash need)"); sm.graphicalProperties.solidFill = "FF4FD8"; sm.graphicalProperties.line.solidFill = "FF4FD8"; c3.series.append(sm)
c3.set_categories(Reference(kp, min_col=1, min_row=CD["m0"], max_row=CD["m24"]))
c3l = LineChart()
sb = Series(Reference(kp, min_col=4, min_row=CD["m0"], max_row=CD["m24"]), title="Cash balance before funding"); style_series(sb, "2DE2FF"); c3l.series.append(sb)
sl = Series(Reference(kp, min_col=6, min_row=CD["m0"], max_row=CD["m24"]), title="Minimum balance"); style_series(sl, "FFB547", 9525); sl.graphicalProperties.line.dashStyle = "dash"; c3l.series.append(sl)
c3 += c3l
kp.add_chart(c3, "I40")
c4 = BarChart(); c4.type = "bar"; c4.grouping = "clustered"; c4.overlap = 100; c4.title = "Sensitivity tornado: Year 5 net vs base"; c4.height = 8; c4.width = 18; c4.x_axis.number_format = '$#,##0;($#,##0);0'
s_lo = Series(Reference(kp, min_col=2, min_row=CD["t0"], max_row=CD["t1"]), title="Downside setting"); s_lo.graphicalProperties.solidFill = "FFB547"; s_lo.graphicalProperties.line.solidFill = "FFB547"; c4.series.append(s_lo)
s_hi = Series(Reference(kp, min_col=3, min_row=CD["t0"], max_row=CD["t1"]), title="Upside setting"); s_hi.graphicalProperties.solidFill = "3DDC97"; s_hi.graphicalProperties.line.solidFill = "3DDC97"; c4.series.append(s_hi)
c4.set_categories(Reference(kp, min_col=1, min_row=CD["t0"], max_row=CD["t1"]))
c4.y_axis.tickLblPos = "low"
kp.add_chart(c4, "I58")
c5 = BarChart(); c5.type = "col"; c5.grouping = "stacked"; c5.overlap = 100; c5.title = "Duty scenario waterfall: hero net per unit (copper, fiber, stress)"; c5.height = 8; c5.width = 18; c5.y_axis.number_format = '$#,##0.00;($#,##0.00);0'
s_base = Series(Reference(kp, min_col=2, min_row=CD["w0"], max_row=CD["w1"]), title="base"); s_base.graphicalProperties.noFill = True; s_base.graphicalProperties.line.noFill = True; c5.series.append(s_base)
s_val = Series(Reference(kp, min_col=3, min_row=CD["w0"], max_row=CD["w1"]), title="Net per unit"); s_val.graphicalProperties.solidFill = "2DE2FF"; s_val.graphicalProperties.line.solidFill = "2DE2FF"; c5.series.append(s_val)
c5.set_categories(Reference(kp, min_col=1, min_row=CD["w0"], max_row=CD["w1"]))
c5.legend = None
kp.add_chart(c5, "I76")
# chart fonts: Arial on titles and axis labels (house font)
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font as DFont, RichTextProperties, RegularTextRun
def _arial(sz=900, bold=False):
    return CharacterProperties(sz=sz, b=bold, latin=DFont(typeface="Arial"))
def _txpr(sz=900):
    return RichText(bodyPr=RichTextProperties(), p=[Paragraph(pPr=ParagraphProperties(defRPr=_arial(sz)), endParaRPr=_arial(sz))])
def _arialize(ch):
    text = ch.title.tx.rich.p[0].r[0].t
    ch.title = Title(tx=Text(rich=RichText(bodyPr=RichTextProperties(), p=[Paragraph(pPr=ParagraphProperties(defRPr=_arial(1200, True)), r=[RegularTextRun(rPr=_arial(1200, True), t=text)])])), overlay=False)
    for ax in (ch.x_axis, ch.y_axis): ax.txPr = _txpr()
    if ch.legend is not None: ch.legend.txPr = _txpr()
for ch in (c1, c2, c3, c4, c5): _arialize(ch)
kp.freeze_panes = "A5"

# ----------------------------------------------------------------------------- FX repatriation rows (need the annual tab)
rr = FX_REPAT_ROW
W(fx, rr, 1, "Year 1 net (active scenario) in KRW", kind="t")
for j, col in enumerate("BCDE"): W(fx, rr, 2 + j, f"={qs(an.title)}!$C${AN['net']}*{col}$7", fmt=FMT_KRW)
W(fx, rr, 6, "Model", kind="t"); rr += 1
W(fx, rr, 1, "Year 5 net (active scenario) in KRW", kind="t")
for j, col in enumerate("BCDE"): W(fx, rr, 2 + j, f"={qs(an.title)}!$G${AN['net']}*{col}$7", fmt=FMT_KRW)
W(fx, rr, 6, "Model", kind="t"); rr += 1
W(fx, rr, 1, "Cumulative 5 year net (active scenario) in KRW", kind="t")
for j, col in enumerate("BCDE"): W(fx, rr, 2 + j, f"={qs(an.title)}!$G${AN['cum_net']}*{col}$7", fmt=FMT_KRW)
W(fx, rr, 6, "Model; before US and Korean tax", kind="t")

# ----------------------------------------------------------------------------- Defined names for figures.json
for key, nm in [("m18_rev", "m18_revenue"), ("m18_net", "m18_net"), ("m18_margin", "m18_net_margin"), ("m18_be", "m18_breakeven_month"), ("m18_runrate", "m18_run_rate"), ("m18_exit", "m18_exit_margin"), ("m18_units", "m18_units"), ("be24", "m24_breakeven_month"), ("m24_rev", "m24_revenue"), ("m24_net", "m24_net")]:
    defname(nm, pl, f"$C${PL[key]}")
for key, nm in [("units", "monthly_units"), ("rev", "monthly_amazon_revenue"), ("net", "monthly_amazon_net"), ("cum", "monthly_amazon_cum_net"), ("tot_rev", "monthly_revenue"), ("tot_net", "monthly_net"), ("tot_cum", "monthly_cum_net"), ("tacos", "monthly_tacos"), ("cpu_b", "monthly_contribution_before_ads"), ("cpu_a", "monthly_contribution_after_ads"), ("margin", "monthly_amazon_net_margin"), ("asp", "monthly_asp"), ("landed_pct", "monthly_landed_pct"), ("ads_pct", "monthly_ads_pct"), ("ret_pct", "monthly_returns_pct"), ("eb_rev", "monthly_ebay_walmart_revenue"), ("b2b_rev", "monthly_b2b_revenue")]:
    defname(nm, pl, f"$D${PL[key]}:$AA${PL[key]}")
for y in range(1, 6):
    yc = Y_COLS[y]
    for key, nm in [("revenue", "revenue"), ("net", "net"), ("margin", "net_margin"), ("cogs_t", "cogs"), ("fees_t", "fees"), ("ful_t", "fulfillment"), ("ads_t", "advertising"), ("ret_t", "returns"), ("con_t", "contribution"), ("fixed", "fixed"), ("fee", "seller_fee"), ("launch", "launch"), ("comp", "compliance"), ("cum_net", "cum_net"), ("drv_units", "amazon_units"), ("asp", "asp"), ("landed", "landed_pct"), ("ads_pct", "ads_pct"), ("amz_con_pct", "amazon_contribution_pct")]:
        defname(f"annual_Y{y}_{nm}", an, f"${yc}${AN[key]}")
    for ch, nm in [("amz", "amazon_us"), ("ebay", "ebay_walmart"), ("b2b", "b2b"), ("dtc", "dtc"), ("ca", "amazon_canada"), ("pro", "pro_line")]:
        defname(f"annual_Y{y}_revenue_by_channel_{nm}", an, f"${yc}${AN[f'{ch}_rev']}")
        defname(f"annual_Y{y}_contribution_by_channel_{nm}", an, f"${yc}${AN[f'{ch}_con']}")
defname("annual_check_revenue_Y1", an, f"$C${AN['chk_rev']}"); defname("annual_check_revenue_Y2", an, f"$D${AN['chk_rev']}"); defname("annual_check_net_Y1", an, f"$C${AN['chk_net']}"); defname("annual_check_net_Y2", an, f"$D${AN['chk_net']}")
for key, nm in [("peak_need", "cash_peak_need"), ("min_label", "cash_peak_need_month_label"), ("min_month_num", "cash_peak_need_month"), ("min_all", "cash_min_balance"), ("gap", "cash_funding_gap"), ("need_by", "cash_funding_needed_by_month"), ("runway_funded", "cash_runway_months_day90"), ("runway_unfunded", "cash_runway_months_day90_unfunded"), ("burn", "cash_burn_months_4_to_6")]:
    defname(nm, cf, f"$C${CASH[key]}")
for key, nm in [("bal", "monthly_cash_balance"), ("bal_funded", "monthly_cash_balance_after_funding"), ("onhand", "monthly_inventory_on_hand"), ("intransit", "monthly_inventory_in_transit"), ("inv_cash", "monthly_inventory_cash_out"), ("arrivals", "monthly_arrivals"), ("closing", "monthly_closing_stock"), ("opcf_gate", "monthly_operating_cash_flow"), ("ncf", "monthly_net_cash_flow"), ("recv", "monthly_receivable")]:
    defname(nm, cf, f"$C${CASH[key]}:$AA${CASH[key]}")
defname("cash_opening_capital", asm, f"$F${ROW['opening_capital']}")
for y in (3, 4, 5):
    defname(f"cash_Y{y}_balance", cf, f"${L(y)}${CASH['a_bal']}"); defname(f"cash_Y{y}_balance_after_funding", cf, f"${L(y)}${CASH['a_bal_funded']}"); defname(f"cash_Y{y}_inventory", cf, f"${L(y)}${CASH['a_inv']}")
defname("gate_day90_result", fb, f"$B${GATE['d90_result']}"); defname("gate_month12_result", fb, f"$B${GATE['m12_result']}")
for key in ["d90_tacos", "d90_ret", "d90_cpu", "d90_runway", "m12_rev", "m12_cum", "m12_tacos", "m12_ret", "m12_opcf"]:
    defname(f"gate_{key}_value", fb, f"$B${GATE[key]}"); defname(f"gate_{key}_pass", fb, f"$E${GATE[key]}"); defname(f"gate_{key}_kill", fb, f"$F${GATE[key]}")
defname("gates_day90_block", fb, f"$A${GATE['d90_conv']}:$G${GATE['d90_result']}"); defname("gates_month12_block", fb, f"$A${GATE['m12_rev']}:$G${GATE['m12_result']}")
defname("funding_plan_block", fb, f"$A${FB['plan0']}:$D${FB['plan0'] + 4}")
for key, nm in [("net18", "sens_base_m18_net"), ("net5", "sens_base_y5_net"), ("rev5", "sens_base_y5_revenue")]:
    defname(nm, se, f"$B${SE[key]}")
defname("sensitivity_table", se, f"$A${SENS['Duty rate']['Fiber 12.5%']}:$H${SENS['Factory cost (FOB)']['Alibaba benchmark'] }")
for drv, settings in SENS.items():
    for lbl, rr in settings.items():
        tag = (drv + "_" + lbl).lower().replace(" (krw per usd)", "_krw").replace(" (fob)", "_fob").replace("%", "pct").replace(",", "").replace(" ", "_").replace(".", "_")
        defname(f"sens_{tag}_m18_net", se, f"$E${rr}"); defname(f"sens_{tag}_y5_net", se, f"$F${rr}")
defname("fx_table", fx, f"$A$7:$E${FXR['hero_cost_krw']}")

# sheet order and save
wb._sheets = [wb[n] for n in TABS]
for ws in wb.worksheets:
    ws.sheet_properties.tabColor = None
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
with open(os.path.join(HERE, "figures_map.json"), "w") as f:
    json.dump(NAMES, f, indent=1)
print(f"saved {OUT} with {len(NAMES)} defined names")


# ----------------------------------------------------------------------------- recalculate and verify
if __name__ == "__main__" and "--no-recalc" not in sys.argv:
    import verify_model
    res = verify_model.recalc(OUT)
    print(json.dumps(res))
    if res.get("status") != "success" or res.get("total_errors", 1) != 0:
        print("RECALC FAILED"); sys.exit(1)
    ok = verify_model.verify(OUT)
    sys.exit(0 if ok else 1)
